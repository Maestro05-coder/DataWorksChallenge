"""
    1. Can a fast ML stand-in for the slow reservoir simulator be trusted?
    2. Which reservoir settings reproduce the real field's history and which settings does the history NOT pin down?
    3. What will the field produce from 2008 to 2028, and how sure are we?

HOW TO RUN
    pip install numpy pandas matplotlib openpyxl scipy scikit-learn
    python HistoryMatchingJudges.py path/to/folder      (or edit DATA_DIR below)

OUTPUTS 
    1_the_problem.png                    simulated cases vs the real field
    2_can_we_trust_the_proxy.png         proxy vs simulator on cases it never saw
    3_what_history_tells_us.png          which settings the history pins down
    4_history_match.png                  best settings vs the real field
    5_forecast_2008_2028.png             the forecast, its spread, and the cases submitted
    6_how_reliable_is_the_forecast.png   forecast tested on the past (hindcast)
    Results_Summary.xlsx                 every table, each sheet with a plain explanation
    <template>_FILLED.xlsx               the deliverable (template filled in)

    The simulated curves and the real field both stop in 2008. The 2008-2028
    forecast is therefore an EXTRAPOLATION: it is tested only by a hindcast, and the
    injection / well plan after 2008 is not in the data.
"""
import re
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.optimize import curve_fit, differential_evolution
from scipy.stats import qmc
from sklearn.decomposition import PCA
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, Matern, WhiteKernel


DATA_DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("data")
OUTPUT_DIR = Path("Judges_Output")
SHOW_PLOTS = True                 # False = only save the PNGs

FILES = {                         # file names without .xlsx
    "intro":    "01 Introduction",
    "field":    "02 Field Production",
    "template": "09 Template Deliverable",
}
PARAM_SHEET = "Uncertainty Values"
PRODUCTION_SHEETS = {
    "gas_cum":    "Gas Production Cumulative",
    "gas_rate":   "Gas Production Rate",
    "oil_cum":    "Oil production cumulative",
    "oil_rate":   "Oil Production Rate",
    "water_cum":  "Water Production Cumulative",
    "water_rate": "Water Production Rate",
}
HISTORY_COL = "RAMP_History"      # the real observed field
CASE_PREFIX = "Case_"

# Which Case_N is train / validation / test (ASSUMED from the sheet counts: confirm!)
TRAIN_IDS = range(1, 71)
VAL_IDS = range(71, 86)
TEST_IDS = range(86, 101)

# Streams used to SCORE the match and how each is transformed first.
# Gas is left out (it follows oil, so it would count oil twice); oil and gas RATES are
# left out (every case starts equal to the history, so they say little).
MATCH_STREAMS = {"oil_cum": "linear", "water_cum": "linear", "water_rate": "log"}
SIM_START = None                  # None = first date found in RAMP_History
LOG_PARAMS = ["Permeability Multiplier"]   # settings handled on a log scale
PCA_VAR, PCA_MAX = 0.999, 6       # curve compression
N_SOBOL_POW2 = 14                 # 2**14 trial settings for the search
ACCEPT_MARGIN_MIN = 0.01          # a trial is "as good as the best" if its misfit is within the
ACCEPT_MARGIN_REL = 0.10          # LARGEST of: this amount, this share of the best misfit,
                                  # or 2x the proxy's own error
FINAL_REFIT_ON_VAL = True         # after reporting accuracy, refit on train+validation
RANDOM_SEED = 0

# --- forecast 2008 -> 2028 and the submission template
FORECAST_START_IF_NO_TEMPLATE = "2008-04-01"
FORECAST_END_IF_NO_TEMPLATE = "2028-01-01"
FIT_WINDOW_YEARS = None           # None = choose automatically with the hindcast
FIT_WINDOW_CANDIDATES = [2.0, 3.0, 5.0]
HINDCAST_YEARS = 3.0              # hide this many years of history to test the forecast
WOR_MAX = 50.0                    # cap on water-oil ratio (50 = 98% water cut)
ANCHOR_TO_HISTORY = True          # make forecasts join the real 2008 history exactly
FORECAST_REFIT_ON_ALL = False     # False = test cases are never used for fitting
MIN_FORECAST_SETS, MAX_FORECAST_SETS = 50, 500
N_SELECT = 10                     # cases to submit (= sheets in the template)
# "diverse_params": best match first, then the most different settings that still fit
# "best":           the lowest-misfit settings (often near duplicates)
# "span_forecast":  settings spread from low to high final oil forecast
SELECTION_MODE = "diverse_params"

NICE = {"gas_cum": "Gas cumulative", "gas_rate": "Gas rate",
        "oil_cum": "Oil cumulative", "oil_rate": "Oil rate",
        "water_cum": "Water cumulative", "water_rate": "Water rate"}
PLAIN = {"oil_cum": "Oil produced so far (cumulative)",
         "water_cum": "Water produced so far (cumulative)",
         "water_rate": "Water production rate"}
UNITS = {"oil_cum": "STB", "water_cum": "STB", "water_rate": "STB/day",
         "gas": "MSCF", "oil": "STB", "water": "STB"}
STREAMS = ["gas", "oil", "water"]                  # order of the template columns
GREEN, AMBER, RED, BLUE = "#1a7f37", "#b8860b", "#b22222", "#1f77b4"


# =============================================================================
# - FIND THE EXCEL FILES AND READ THE SHEETS
# =============================================================================
class SheetNotFound(Exception):
    pass


_NAMES_CACHE = {}


def listWorkbooks():
    return [p for p in sorted(DATA_DIR.glob("*.xlsx")) if not p.name.startswith("~$")]


def findFile(stem):
    """The .xlsx in DATA_DIR whose name matches `stem` (exact, then 'starts with')."""
    target = stem.strip().lower()
    for p in listWorkbooks():
        if p.stem.strip().lower() == target:
            return p
    for p in listWorkbooks():
        if p.stem.strip().lower().startswith(target):
            return p
    return None


def sheetNames(path):
    """Sheet names of a workbook without loading its cells (fast)."""
    key = str(path)
    if key not in _NAMES_CACHE:
        with pd.ExcelFile(path) as xl:
            _NAMES_CACHE[key] = xl.sheet_names
    return _NAMES_CACHE[key]


def findWorkbookWithSheet(sheetName):
    """Fallback for ONE combined workbook: the first .xlsx that has this sheet."""
    target = sheetName.strip().lower()
    for p in listWorkbooks():
        if target in [n.strip().lower() for n in sheetNames(p)]:
            return p
    return None


def readSheet(fileStem, sheetName, header=None):
    path = findFile(fileStem)
    if path is None:
        path = findWorkbookWithSheet(sheetName)
    if path is None:
        raise FileNotFoundError(
            f"Could not find a file named like '{fileStem}.xlsx' in {DATA_DIR.resolve()}.\n"
            f"Excel files I can see there: {[p.name for p in listWorkbooks()]}")
    lookup = {n.strip().lower(): n for n in sheetNames(path)}
    real = lookup.get(sheetName.strip().lower())
    if real is None:
        raise SheetNotFound(f"Sheet '{sheetName}' not found in '{path.name}'. "
                            f"Sheets there: {sheetNames(path)}")
    return pd.read_excel(path, sheet_name=real, header=header)


def loadParameters():
    """Uncertainty Values -> one row per case ('Case_1'...), one column per setting."""
    df = readSheet(FILES["intro"], PARAM_SHEET, header=0)
    df = df.rename(columns={df.columns[0]: "Item"})
    df.columns = [str(c).strip() for c in df.columns]
    df["Item"] = df["Item"].astype(str).str.strip().str.replace(" ", "_", regex=False)
    return df.set_index("Item").apply(pd.to_numeric, errors="coerce").dropna(how="all")


def loadProduction(sheetName):
    """One production sheet -> clean (dates x cases) table. Tolerates the different
    header layouts by locating the row whose first cell says 'Date'."""
    raw = readSheet(FILES["field"], sheetName, header=None)
    headerRow = raw.index[raw.iloc[:, 0].astype(str).str.strip().eq("Date")][0]
    body = raw.iloc[headerRow + 2:, :].copy()                  # skip header + units rows
    body.columns = ["Date"] + [str(c).strip() for c in raw.iloc[headerRow, 1:]]
    body["Date"] = pd.to_datetime(body["Date"], errors="coerce")
    body = body.dropna(subset=["Date"]).set_index("Date").sort_index()
    body = body.apply(pd.to_numeric, errors="coerce")
    return body.loc[:, ~body.columns.isin(["nan", ""])]


def caseCols(df):
    return [c for c in df.columns if str(c).startswith(CASE_PREFIX)]


def lastValid(series):
    s = series.dropna()
    return s.iloc[-1] if len(s) else np.nan


def caseId(name):
    return int(str(name).split("_")[-1])


def caseNumber(name):
    nums = re.findall(r"\d+", str(name))
    return int(nums[-1]) if nums else -1


def splitOf(name):
    i = caseId(name)
    if i in TRAIN_IDS:
        return "train"
    if i in VAL_IDS:
        return "validation"
    if i in TEST_IDS:
        return "test"
    return "unassigned"


def printDataChecks(data, params):
    """A short check of the data before it is trusted."""
    print("\n=== DATA CHECK ===")
    print(f"settings: {len(params)} cases x {list(params.columns)}")
    for key, df in data.items():
        cc = caseCols(df)
        empty = [c for c in cc if df[c].notna().sum() == 0]
        print(f"  {key:10s} cases={len(cc)}  history points={df[HISTORY_COL].notna().sum()}  "
              f"dates {df.index.min().date()} -> {df.index.max().date()}"
              + (f"  EMPTY cases: {empty[:5]}" if empty else ""))
    histEnd = data["oil_cum"][HISTORY_COL].dropna().index.max()
    if any(df.index.max() > histEnd + pd.Timedelta(days=1) for df in data.values()):
        print("  NOTE: some sheets hold dates AFTER the history end - check for future data.")
    ids = [c for c in caseCols(data["oil_cum"]) if c in params.index]
    print(f"  cases with both settings and curves: {len(ids)}  "
          f"splits: {pd.Series([splitOf(c) for c in ids]).value_counts().to_dict()}")


# =============================================================================
# STEP 2 - PREPARE CURVES: every case on the SAME dates as the real history
# =============================================================================
def toT(x, kind):
    """Optional transform before scoring/modelling ('log' keeps small water visible)."""
    x = np.asarray(x, float)
    return np.log10(np.clip(x, 0, None) + 1.0) if kind == "log" else x


def fromT(y, kind):
    y = np.asarray(y, float)
    return 10.0 ** y - 1.0 if kind == "log" else y


def evalCase(series, dates, isCum, t0):
    """Value of ONE case's curve at `dates`. Cumulative: linear between reported
    points, 0 at the start. Rate: nearest reported value. Outside the case's
    range the result is NaN (no extrapolation)."""
    s = series.dropna()
    if len(s) < 2:
        return np.full(len(dates), np.nan)

    def ns(idx):
        return pd.DatetimeIndex(idx).values.astype("datetime64[ns]").astype("int64").astype(float)

    x, y, xi = ns(s.index), s.values.astype(float), ns(dates)
    if isCum:
        if x[0] > t0:
            x, y = np.r_[t0, x], np.r_[0.0, y]
        return np.interp(xi, x, y, left=np.nan, right=np.nan)
    hi = np.clip(np.searchsorted(x, xi), 0, len(x) - 1)
    lo = np.clip(hi - 1, 0, len(x) - 1)
    pick = np.where(np.abs(xi - x[lo]) <= np.abs(x[hi] - xi), lo, hi)
    out = y[pick]
    out[(xi < x[0]) | (xi > x[-1])] = np.nan
    return out


def valueAt(series, dates, key, t0):
    return evalCase(series, pd.DatetimeIndex(dates), key.endswith("_cum"), t0)


def buildTargets(data, cases, t0):
    """For each scored stream: the real history and every case's curve on the SAME
    dates. Returns (targets, usable cases)."""
    T = {}
    for key, kind in MATCH_STREAMS.items():
        df = data[key]
        hist = df[HISTORY_COL].dropna()
        hist = hist[~hist.index.duplicated()]
        Y = np.array([toT(evalCase(df[c], hist.index, key.endswith("_cum"), t0), kind)
                      for c in cases])
        T[key] = dict(kind=kind, dates=hist.index, obs=toT(hist.values, kind), Y=Y)
    for t in T.values():                              # keep dates nearly every case covers
        keep = np.isfinite(t["Y"]).mean(axis=0) >= 0.95
        t["Y"], t["obs"], t["dates"] = t["Y"][:, keep], t["obs"][keep], t["dates"][keep]
    ok = np.ones(len(cases), bool)
    for t in T.values():
        ok &= np.isfinite(t["Y"]).all(axis=1)         # drop cases with holes
    for t in T.values():
        t["Y"] = t["Y"][ok]
        t["scale"] = max(float(np.abs(t["obs"]).max()), 1e-9)
    dropped = [c for c, o in zip(cases, ok) if not o]
    if dropped:
        print(f"   {len(dropped)} cases dropped (holes on the history dates): {dropped[:8]}")
    return T, [c for c, o in zip(cases, ok) if o]


def streamRmse(pred, t):
    """Error of predicted curves vs the real history, as a share of its largest value."""
    return np.sqrt(np.mean((pred - t["obs"][None, :]) ** 2, axis=1)) / t["scale"]


def totalMisfit(preds, T):
    """Sum over the scored streams (lower = closer to the real field)."""
    return sum(streamRmse(preds[k], T[k]) for k in T)


# =============================================================================
# STEP 3 - THE PROXY: PCA (compress each curve) + Gaussian process (settings -> curve)
# =============================================================================
class CurveProxy:
    """4 settings in -> the scored curves out. Per stream: PCA squeezes the curves into
    a few shape numbers, then one Gaussian process predicts each number. A Gaussian
    process suits ~70-100 runs: it needs few samples and reports its own uncertainty."""

    def __init__(self, paramCols, bounds):
        self.cols, self.bounds = list(paramCols), bounds

    def toUnit(self, P):
        """Physical settings -> [0,1] box (log scale for LOG_PARAMS)."""
        P = np.array(P, float, ndmin=2)
        U = np.empty_like(P)
        for j, c in enumerate(self.cols):
            lo, hi = self.bounds[c]
            if c in LOG_PARAMS:
                U[:, j] = (np.log(P[:, j]) - np.log(lo)) / (np.log(hi) - np.log(lo))
            else:
                U[:, j] = (P[:, j] - lo) / (hi - lo)
        return U

    def fromUnit(self, U):
        U = np.array(U, float, ndmin=2)
        P = np.empty_like(U)
        for j, c in enumerate(self.cols):
            lo, hi = self.bounds[c]
            if c in LOG_PARAMS:
                P[:, j] = np.exp(np.log(lo) + U[:, j] * (np.log(hi) - np.log(lo)))
            else:
                P[:, j] = lo + U[:, j] * (hi - lo)
        return P

    def fit(self, P, Y):
        U = self.toUnit(P)
        self.models = {}
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            for key, y in Y.items():
                mu = y.mean(axis=0)
                npc = int(min(PCA_MAX, y.shape[0] - 1, y.shape[1]))
                pca = PCA(n_components=npc).fit(y - mu)
                cum = np.nan_to_num(np.cumsum(pca.explained_variance_ratio_))
                k = int(min(npc, np.searchsorted(cum, PCA_VAR) + 1))
                Z = pca.transform(y - mu)[:, :k]
                gps = []
                for j in range(k):
                    kern = (ConstantKernel(1.0, (1e-2, 1e2))
                            * Matern(np.ones(U.shape[1]), (0.05, 50.0), nu=2.5)
                            + WhiteKernel(1e-4, (1e-8, 1e-1)))
                    gps.append(GaussianProcessRegressor(
                        kern, normalize_y=True, n_restarts_optimizer=2,
                        random_state=RANDOM_SEED).fit(U, Z[:, j]))
                self.models[key] = dict(mu=mu, comp=pca.components_[:k], gps=gps)
        return self

    def predict(self, P):
        U = self.toUnit(P)
        out = {}
        for key, m in self.models.items():
            Z = np.column_stack([gp.predict(U) for gp in m["gps"]])
            out[key] = m["mu"] + Z @ m["comp"]
        return out


def cvPredict(P, Y, paramCols, bounds, k=5):
    """Out-of-fold predictions: every case is predicted by a proxy that never saw it."""
    folds = np.arange(len(P)) % k
    np.random.RandomState(RANDOM_SEED).shuffle(folds)
    out = {key: np.zeros_like(y) for key, y in Y.items()}
    for f in range(k):
        tr, te = folds != f, folds == f
        pred = CurveProxy(paramCols, bounds).fit(P[tr], {key: y[tr] for key, y in Y.items()}).predict(P[te])
        for key in Y:
            out[key][te] = pred[key]
    return out


def verdictWord(errPct):
    """Plain-language grade for an error given as % of the largest history value."""
    return "excellent" if errPct < 2 else "good" if errPct < 5 else "fair" if errPct < 10 else "poor"


def proxyReport(pred, Y, T, label):
    """Accuracy of predicted curves vs the real simulated curves for these cases."""
    rows = []
    for key in T:
        err = np.sqrt(np.mean((pred[key] - Y[key]) ** 2, axis=1)) / T[key]["scale"]
        lt, lp = Y[key][:, -1], pred[key][:, -1]
        sst = np.sum((lt - lt.mean()) ** 2)
        rows.append({"dataset": label, "measure": key, "cases": len(err),
                     "typical_error_pct": 100 * err.mean(), "worst_error_pct": 100 * err.max(),
                     "R2_final": 1 - np.sum((lt - lp) ** 2) / sst if sst > 0 else np.nan})
    return pd.DataFrame(rows)


# =============================================================================
# STEP 4 - HISTORY MATCH: search ALL possible settings through the proxy
# =============================================================================
def runHistoryMatch(data, params):
    print("\n=== PROXY + HISTORY MATCH ===")
    pcols = list(params.columns)
    first = next(iter(MATCH_STREAMS))
    cases = [c for c in caseCols(data[first]) if c in params.index and params.loc[c].notna().all()]
    t0 = pd.Timestamp(SIM_START or min(data[k][HISTORY_COL].dropna().index.min()
                                       for k in MATCH_STREAMS)).value
    T, cases = buildTargets(data, cases, t0)
    P = params.loc[cases, pcols].values.astype(float)
    split = np.array([splitOf(c) for c in cases])
    bounds = {c: (params[c].min(), params[c].max()) for c in pcols}      # the tested ranges
    Yall = {k: t["Y"] for k, t in T.items()}
    print("cases usable: " + ", ".join(f"{s}={int((split == s).sum())}"
                                       for s in ["train", "validation", "test", "unassigned"]))
    tr = split == "train"
    if tr.sum() < 15:
        print("   WARNING: fewer than 15 train cases under TRAIN_IDS - all cases treated as train.")
        tr = np.ones(len(cases), bool)

    def sub(m):
        return {k: y[m] for k, y in Yall.items()}

    # ---- how accurate is the proxy? (cases it never saw) ----
    cvp = cvPredict(P[tr], sub(tr), pcols, bounds)
    reports = [proxyReport(cvp, sub(tr), T, "train (cross-validation)")]
    proxy = CurveProxy(pcols, bounds).fit(P[tr], sub(tr))
    held = {}
    for name in ["validation", "test"]:
        m = split == name
        if m.any():
            held[name] = (m, proxy.predict(P[m]))
            reports.append(proxyReport(held[name][1], sub(m), T, name))
    rep = pd.concat(reports, ignore_index=True)
    rep["verdict"] = rep["typical_error_pct"].map(verdictWord)
    print(rep.round(3).to_string(index=False))
    refM, refPred = held["validation"] if "validation" in held else (tr, cvp)
    refTrue = {k: y[refM] for k, y in Yall.items()}
    proxyErr = float(np.mean(np.abs(totalMisfit(refTrue, T) - totalMisfit(refPred, T))))
    print(f"proxy's own error in misfit units: {proxyErr:.4f}")
    use = tr | (split == "validation") if FINAL_REFIT_ON_VAL else tr
    if FINAL_REFIT_ON_VAL and (split == "validation").any():
        proxy = CurveProxy(pcols, bounds).fit(P[use], sub(use))
        print(f"final proxy refit on train+validation ({int(use.sum())} cases); test untouched")

    # ---- search every setting, not just the existing runs ----
    d = len(pcols)
    U = qmc.Sobol(d=d, scramble=True, seed=RANDOM_SEED).random_base2(N_SOBOL_POW2)
    samples = proxy.fromUnit(U)
    mis = totalMisfit(proxy.predict(samples), T)

    def objective(x):
        x = np.asarray(x, float)
        one = x.ndim == 1
        m = totalMisfit(proxy.predict(proxy.fromUnit(x.reshape(d, -1).T)), T)
        return float(m[0]) if one else m

    res = differential_evolution(objective, [(0.0, 1.0)] * d, seed=RANDOM_SEED, maxiter=150,
                                 popsize=20, tol=1e-8, vectorized=True,
                                 updating="deferred", polish=True)
    bestP = proxy.fromUnit(np.clip(res.x, 0, 1))[0]
    bestM = float(objective(np.clip(res.x, 0, 1)))
    if mis.min() < bestM:
        bestP, bestM = samples[int(mis.argmin())], float(mis.min())
    bestPred = proxy.predict(bestP[None, :])

    # ---- every setting that fits "about as well as the best" ----
    margin = max(ACCEPT_MARGIN_MIN, ACCEPT_MARGIN_REL * bestM, 2 * proxyErr)
    acc = mis <= bestM + margin
    A = samples[acc]
    existing = pd.Series(totalMisfit(Yall, T), index=cases).sort_values()
    print(f"\nBEST settings (misfit {bestM:.4f}); acceptable = within {margin:.4f} of the best: "
          f"{int(acc.sum())} of {len(mis)} trial sets")
    print(f"best EXISTING simulated case: {existing.index[0]} (misfit {existing.iloc[0]:.4f})")

    Ua, rows = proxy.toUnit(A), []
    for j, c in enumerate(pcols):
        w = float(Ua[:, j].max() - Ua[:, j].min())
        verdict = "PINNED DOWN" if w < 0.5 else "partly constrained" if w < 0.8 else "NOT constrained"
        rows.append({"setting": c, "best": bestP[j], "acceptable_min": A[:, j].min(),
                     "acceptable_max": A[:, j].max(), "tested_min": bounds[c][0],
                     "tested_max": bounds[c][1], "share_of_range": w, "verdict": verdict})
        print(f"  {verdict:18s} {c}: best {bestP[j]:.4g}, acceptable {A[:, j].min():.3g} to "
              f"{A[:, j].max():.3g} (tested {bounds[c][0]:.3g} to {bounds[c][1]:.3g})")
    ident = pd.DataFrame(rows)

    resid = {}
    for k, t in T.items():
        r = (bestPred[k][0] - t["obs"]) / t["scale"]
        resid[k] = float(np.sqrt(np.mean(r ** 2)))
    print("misfit at the best settings (RMSE as % of the largest history value): "
          + ", ".join(f"{k} {100 * v:.1f}%" for k, v in resid.items()))
    if max(resid.values()) > 0.10:
        print("   CHECK: the proxy cannot reproduce history in at least one stream; the 4 settings "
              "may not span the truth.")
    return dict(proxy=proxy, bestParams=bestP, bestMisfit=bestM, acceptable=A, accMis=mis[acc],
                targets=T, Yall=Yall, held=held, rep=rep, ident=ident, bestPred=bestPred,
                resid=resid, existing=existing, t0=t0, pcols=pcols, bounds=bounds,
                samples=samples, mis=mis)


# =============================================================================
# STEP 5 - FORECAST 2008 -> 2028 (EXTRAPOLATION beyond the data)
# =============================================================================
def arps(t, qi, Di, b):
    """Hyperbolic decline: the rate falls fast at first, then more slowly."""
    return qi / np.power(1.0 + b * Di * t, 1.0 / b)


def fitArps(t, q):
    """Fit (qi, Di, b) to a rate history; t = years from the start of the window."""
    m = np.isfinite(t) & np.isfinite(q) & (q > 0)
    if m.sum() < 5:
        return None
    t, q = t[m], q[m]
    lq = np.log(q)

    def f(tt, lqi, lDi, b):
        return np.log(arps(tt, np.exp(lqi), np.exp(lDi), b))

    best = None
    for b0 in (0.1, 0.5, 0.9):
        try:
            popt, _ = curve_fit(f, t, lq, p0=[lq[0], np.log(0.15), b0],
                                bounds=([lq.min() - 3, -9, 0.01], [lq.max() + 3, 2.0, 1.0]),
                                maxfev=4000)
            sse = np.sum((f(t, *popt) - lq) ** 2)
            if best is None or sse < best[0]:
                best = (sse, popt)
        except Exception:
            continue
    if best is None:                                    # fallback: plain exponential
        slope, icpt = np.polyfit(t, lq, 1)
        return (float(np.exp(icpt)), float(max(-slope, 1e-6)), 0.01)
    return (float(np.exp(best[1][0])), float(np.exp(best[1][1])), float(best[1][2]))


def fitShapes(data, case, cut, windowYears, t0):
    """Everything the forecast needs from one case using history up to `cut`:
    its state at `cut` and decline shapes fitted on the last `windowYears`."""
    w0 = cut - pd.Timedelta(days=int(365.25 * windowYears))
    out = {"w0": w0, "cut": cut}
    for s in ("oil", "gas"):
        ser = data[f"{s}_rate"][case].dropna()
        ser = ser[(ser.index >= w0) & (ser.index <= cut)]
        out[s] = fitArps((ser.index - w0).days.values / 365.25, ser.values)
        if out[s] is None:
            return None
    st = {k: float(valueAt(data[k][case], [cut], k, t0)[0])
          for k in ["oil_cum", "gas_cum", "water_cum", "oil_rate", "gas_rate", "water_rate"]}
    if not np.all(np.isfinite(list(st.values()))):
        return None
    out["state"] = st
    grid = pd.date_range(w0, cut, periods=24)
    qo = valueAt(data["oil_rate"][case], grid, "oil_rate", t0)
    qw = valueAt(data["water_rate"][case], grid, "water_rate", t0)
    npv = valueAt(data["oil_cum"][case], grid, "oil_cum", t0)
    with np.errstate(divide="ignore", invalid="ignore"):
        wor = np.where(qo > 0, qw / qo, np.nan)
    ok = np.isfinite(wor) & np.isfinite(npv) & (wor > 1e-4)
    slopeW = 0.0
    if ok.sum() >= 5:
        slope, _ = np.polyfit(npv[ok] / max(st["oil_cum"], 1.0), np.log(wor[ok]), 1)
        slopeW = max(float(slope), 0.0)                 # water cut does not fall with production
    out["worSlope"] = slopeW
    return out


def project(shp, outDates):
    """Cumulative gas / oil / water at `outDates` (all on/after shp['cut'])."""
    cut, st = shp["cut"], shp["state"]
    outDates = pd.DatetimeIndex(outDates)
    d = np.arange(0, (outDates.max() - cut).days + 8, 7.0)           # weekly steps
    off = (cut - shp["w0"]).days / 365.25
    ty, dd = off + d / 365.25, np.diff(d)

    def integrate(q):
        return np.concatenate([[0.0], np.cumsum(0.5 * (q[1:] + q[:-1]) * dd)])

    q = {}
    for s in ("oil", "gas"):                            # rescaled to join the last real rate
        qi, Di, b = shp[s]
        q[s] = arps(ty, qi, Di, b) * (st[f"{s}_rate"] / max(arps(off, qi, Di, b), 1e-12))
    cumO = st["oil_cum"] + integrate(q["oil"])
    worEnd = max(st["water_rate"] / max(st["oil_rate"], 1e-9), 1e-4)
    wor = np.minimum(worEnd * np.exp(shp["worSlope"] * (cumO - st["oil_cum"])
                                     / max(st["oil_cum"], 1.0)), max(WOR_MAX, worEnd))
    x = (outDates - cut).days.values.astype(float)
    return {"gas": np.interp(x, d, st["gas_cum"] + integrate(q["gas"])),
            "oil": np.interp(x, d, cumO),
            "water": np.interp(x, d, st["water_cum"] + integrate(wor * q["oil"]))}


def hindcast(data, cases, windowYears, histEnd, t0):
    """Hide the last HINDCAST_YEARS of history, forecast them, and compare the predicted
    production INCREASE with what the simulator really did."""
    cut = histEnd - pd.Timedelta(days=int(365.25 * HINDCAST_YEARS))
    rows = {}
    for c in cases:
        shp = fitShapes(data, c, cut, windowYears, t0)
        if shp is None:
            continue
        pred, r = project(shp, [histEnd]), {}
        for s in STREAMS:
            trueEnd = float(valueAt(data[f"{s}_cum"][c], [histEnd], f"{s}_cum", t0)[0])
            inc = trueEnd - shp["state"][f"{s}_cum"]
            r[s] = (pred[s][0] - shp["state"][f"{s}_cum"]) / inc - 1 if inc > 0 else np.nan
        rows[c] = r
    return pd.DataFrame(rows).T


def selectCases(U, oilEnd, n, mode):
    """Indices into the acceptable settings (sorted best-misfit first)."""
    n = min(n, len(U))
    if mode == "best":
        return list(range(n))
    if mode == "span_forecast":
        order = np.argsort(oilEnd)
        pick = np.unique(np.round(np.linspace(0.03, 0.97, n) * (len(order) - 1)).astype(int))
        return sorted(order[pick].tolist())
    chosen = [0]                                        # best match, then the most different
    while len(chosen) < n:
        dist = np.min(np.linalg.norm(U[:, None, :] - U[chosen][None, :, :], axis=2), axis=1)
        dist[chosen] = -1
        chosen.append(int(np.argmax(dist)))
    return chosen


# ---- the submission template ------------------------------------------------
def findTemplate():
    path = findFile(FILES["template"])
    if path is None:
        for p in listWorkbooks():
            if "template" in p.stem.lower():
                return p
    return path


def readTemplateLayout(path):
    """Find, in every 'Case N' sheet: header row, date column, the gas / oil / water
    cumulative columns and the (row, date) of each data row. Nothing is hard-coded."""
    wb = load_workbook(path)
    layout = {}
    for name in wb.sheetnames:
        if not re.search(r"case\s*_?\s*\d+", name, re.I):
            continue
        ws, hdr = wb[name], None
        for r in range(1, min(ws.max_row, 15) + 1):
            cells = [str(c.value).lower() if c.value is not None else "" for c in ws[r]]
            if sum("cumulative" in x for x in cells) >= 2:
                hdr = r
                break
        if hdr is None:
            continue
        cols = {}
        for c in ws[hdr]:
            v = str(c.value).lower() if c.value is not None else ""
            if "cumulative" in v:
                for s in STREAMS:
                    if s in v:
                        cols[s] = c.column
                        break
        counts = {col: sum(hasattr(ws.cell(row=r, column=col).value, "year")
                           for r in range(hdr + 1, ws.max_row + 1))
                  for col in range(1, ws.max_column + 1)}
        dateCol = max(counts, key=counts.get)
        rows = [(r, pd.Timestamp(ws.cell(row=r, column=dateCol).value))
                for r in range(hdr + 1, ws.max_row + 1)
                if hasattr(ws.cell(row=r, column=dateCol).value, "year")]
        layout[name] = {"hdr": hdr, "cols": cols, "dateCol": dateCol, "rows": rows}
    return wb, dict(sorted(layout.items(), key=lambda kv: caseNumber(kv[0])))


def makeOwnTemplate(dates):
    """Fallback when no template file exists: same layout, our own dates."""
    path = OUTPUT_DIR / "own_template.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    for k in range(1, N_SELECT + 1):
        ws = wb.create_sheet(f"Case {k}")
        ws.append(["Date", "Gas production cumulative [MSCF]",
                   "Oil production cumulative [STB]", "Water production cumulative [STB]"])
        for d in dates:
            ws.append([d.to_pydatetime()])
    wb.save(path)
    return path


def checkDeliverable(path, obsEnd):
    """Read the FILLED template back from disk and verify the deliverable's rules.
    Returns a table (one row per sheet) so it can also go into the summary Excel."""
    wb, layout = readTemplateLayout(path)
    print("\n=== DELIVERABLE CHECK (reading the saved file back) ===")
    rows = []
    for name, lay in layout.items():
        ws, problems = wb[name], []
        for s in STREAMS:
            vals = np.array([ws.cell(row=r, column=lay["cols"][s]).value
                             if ws.cell(row=r, column=lay["cols"][s]).value is not None else np.nan
                             for r, _ in lay["rows"]], float)
            if np.isnan(vals).any():
                problems.append(f"{s}: blank cells")
                continue
            if (vals < 0).any():
                problems.append(f"{s}: negative values")
            if (np.diff(vals) < -1e-6 * max(abs(vals).max(), 1)).any():
                problems.append(f"{s}: cumulative decreases")
            if np.isfinite(obsEnd.get(s, np.nan)) and vals[0] < obsEnd[s] * (1 - 1e-9):
                problems.append(f"{s}: starts below the 2008 history value")
        status = "PASS" if not problems else "FAIL: " + "; ".join(problems)
        print(f"  {name:8s} rows={len(lay['rows'])}  {status}")
        rows.append({"sheet": name, "rows": len(lay["rows"]), "result": status})
    return pd.DataFrame(rows)


def runForecast(data, params, hm):
    print("\n=== FORECAST 2008 -> 2028 ===")
    proxy, t0, pcols, bounds = hm["proxy"], hm["t0"], hm["pcols"], hm["bounds"]
    histEnd = data["oil_cum"][HISTORY_COL].dropna().index.max()
    cases = [c for c in caseCols(data["oil_cum"]) if c in params.index and params.loc[c].notna().all()]
    print(f"history ends {histEnd.date()}; the simulator curves stop here too, so everything "
          "after is extrapolation")

    # ---- the template defines the dates to forecast ----
    tpath = findTemplate()
    layout = None
    if tpath is not None:
        wb, layout = readTemplateLayout(tpath)
        fdates = pd.DatetimeIndex([d for _, d in next(iter(layout.values()))["rows"]])
        print(f"template: {tpath.name} ({len(layout)} case sheets), {len(fdates)} dates "
              f"{fdates.min().date()} -> {fdates.max().date()}")
        if (fdates <= histEnd).any():
            print("  WARNING: template has dates on/before the history end; those rows are skipped.")
            fdates = fdates[fdates > histEnd]
    else:
        print("no template file found -> writing our own 10-sheet workbook")
        fdates = pd.date_range(FORECAST_START_IF_NO_TEMPLATE, FORECAST_END_IF_NO_TEMPLATE, freq="QS")

    # ---- hindcast: test the method on the past, choose the fitting window ----
    errs = {w: hindcast(data, cases, w, histEnd, t0) for w in FIT_WINDOW_CANDIDATES}
    for w, e in errs.items():
        print(f"  hindcast, decline curves fitted on last {w:g} yr -> typical error of the "
              f"{HINDCAST_YEARS:g}-yr increase: "
              + "  ".join(f"{s} {e[s].abs().median():.1%}" for s in STREAMS))
    chosen = FIT_WINDOW_YEARS or min(
        errs, key=lambda w: np.nanmean([errs[w][s].abs().median() for s in STREAMS]))
    if chosen not in errs:
        errs[chosen] = hindcast(data, cases, chosen, histEnd, t0)
    print(f"  chosen fitting window: last {chosen:g} years")

    # ---- one extrapolated forecast per case, then a proxy that forecasts for ANY settings ----
    fdExt = pd.DatetimeIndex([histEnd]).append(fdates)        # column 0 = state at history end
    F, okCases = {s: [] for s in STREAMS}, []
    for c in cases:
        shp = fitShapes(data, c, histEnd, chosen, t0)
        if shp is None:
            continue
        pr = project(shp, fdExt)
        okCases.append(c)
        for s in STREAMS:
            F[s].append(pr[s])
    print(f"  cases with a usable forecast fit: {len(okCases)} of {len(cases)}")
    P = params.loc[okCases, pcols].values.astype(float)
    split = np.array([splitOf(c) for c in okCases])
    use = np.ones(len(okCases), bool) if FORECAST_REFIT_ON_ALL else np.isin(split, ["train", "validation"])
    if use.sum() < 15:
        print("  WARNING: fewer than 15 train/validation cases - using ALL cases for the forecast proxy.")
        use = np.ones(len(okCases), bool)
    kinds = {f"fc_{s}": ("log" if s == "water" else "linear") for s in STREAMS}
    Yf = {f"fc_{s}": toT(np.array(F[s]), kinds[f"fc_{s}"]) for s in STREAMS}
    fproxy = CurveProxy(pcols, bounds).fit(P[use], {k: y[use] for k, y in Yf.items()})

    # ---- forecast for every ACCEPTABLE setting -> spread between equally good matches ----
    A, am = hm["acceptable"], hm["accMis"]
    if len(A) < max(MIN_FORECAST_SETS, N_SELECT):
        k = max(MIN_FORECAST_SETS, N_SELECT * 5)
        idx = np.argsort(hm["mis"])[:k]
        print(f"  only {len(A)} acceptable sets; using the {k} best-fitting trial sets instead")
        A, am = hm["samples"][idx], hm["mis"][idx]
    if len(A) > MAX_FORECAST_SETS:
        idx = np.random.RandomState(RANDOM_SEED).choice(len(A), MAX_FORECAST_SETS, replace=False)
        A, am = A[idx], am[idx]
    order = np.argsort(am)
    A, am = A[order], am[order]
    pred = fproxy.predict(A)
    obsEnd = {s: lastValid(data[f"{s}_cum"][HISTORY_COL]) for s in STREAMS}
    ens = {}
    for s in STREAMS:
        arr = fromT(pred[f"fc_{s}"], kinds[f"fc_{s}"])
        fc = arr[:, 1:]
        if ANCHOR_TO_HISTORY and np.isfinite(obsEnd[s]):
            fc = fc + (obsEnd[s] - arr[:, 0])[:, None]          # join the real history exactly
        fc = np.maximum.accumulate(np.maximum(fc, 0), axis=1)
        if np.isfinite(obsEnd[s]):
            fc = np.maximum(fc, obsEnd[s])
        ens[s] = fc
    bands = {s: tuple(np.percentile(ens[s], q, axis=0) for q in (10, 50, 90)) for s in STREAMS}

    # ---- choose the cases to submit ----
    pick = selectCases(proxy.toUnit(A), ens["oil"][:, -1], N_SELECT, SELECTION_MODE)
    pick = sorted(pick, key=lambda i: am[i])
    rows = []
    for rank, i in enumerate(pick, start=1):
        row = {"case": rank, "misfit": am[i], **dict(zip(pcols, A[i]))}
        for s in STREAMS:
            row[f"{s}_cum_{fdates[-1].year}"] = ens[s][i, -1]
        rows.append(row)
    sel = pd.DataFrame(rows)
    print(f"\nSELECTED CASES (mode '{SELECTION_MODE}', {len(pick)} of {len(A)} acceptable sets):")
    print(sel.round(4).to_string(index=False))

    # ---- fill a COPY of the template and check it ----
    if tpath is None:
        tpath = makeOwnTemplate(fdates)
        wb, layout = readTemplateLayout(tpath)
    outPath = OUTPUT_DIR / f"{tpath.stem}_FILLED.xlsx"
    sheets = list(layout.items())
    if len(sheets) < len(pick):
        print(f"WARNING: template has {len(sheets)} case sheets but {len(pick)} cases were selected.")
    dateK = {d: k for k, d in enumerate(fdates)}
    for (name, lay), i in zip(sheets, pick):
        for r, d in lay["rows"]:
            if d in dateK:
                for s in STREAMS:
                    wb[name].cell(row=r, column=lay["cols"][s]).value = float(ens[s][i, dateK[d]])
    wb.save(outPath)
    print(f"\nsaved filled template: {outPath}   (original template untouched)")
    check = checkDeliverable(outPath, obsEnd)
    return dict(fdates=fdates, bands=bands, histEnd=histEnd, obsEnd=obsEnd, errs=errs, chosen=chosen,
                selected=sel, selCurves={s: [ens[s][i] for i in pick] for s in STREAMS},
                check=check, templatePath=outPath, nAcceptable=len(A))


# =============================================================================
# STEP 6 - THE GRAPHS (each has a headline that states its conclusion)
# =============================================================================
def savePlot(fig, name):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    fig.savefig(OUTPUT_DIR / name, dpi=140, bbox_inches="tight")
    print(f"saved {OUTPUT_DIR / name}")


def millions(ax):
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v / 1e6:.0f}M"))


def plotProblem(hm):
    """1: the problem - 100 simulated reservoirs, one real field."""
    T = hm["targets"]
    fig, axes = plt.subplots(1, len(T), figsize=(6 * len(T), 5), constrained_layout=True)
    for ax, (key, t) in zip(np.atleast_1d(axes), T.items()):
        for y in t["Y"]:
            ax.plot(t["dates"], fromT(y, t["kind"]), color="#bbbbbb", lw=0.7)
        ax.plot(t["dates"], fromT(t["obs"], t["kind"]), "ko", ms=4, label="real field")
        ax.plot([], [], color="#bbbbbb", label=f"{len(t['Y'])} simulated cases")
        ax.set_title(PLAIN[key], weight="bold")
        ax.set_ylabel(UNITS[key])
        ax.tick_params(axis="x", rotation=30)
        ax.grid(alpha=0.3)
        if key.endswith("_cum"):
            millions(ax)
    np.atleast_1d(axes)[0].legend(loc="upper left")
    fig.suptitle("The problem: many possible reservoirs (grey), one real field (black dots)",
                 fontsize=15, weight="bold")
    fig.supxlabel("Each grey line is one simulation run with a different combination of 4 uncertain "
                  "reservoir settings. History matching = finding the settings whose curve lands on the dots.",
                  fontsize=11)
    savePlot(fig, "1_the_problem.png")


def plotProxyTrust(hm):
    """2: can we trust the fast ML stand-in (proxy)?"""
    T, held, Yall, rep = hm["targets"], hm["held"], hm["Yall"], hm["rep"]
    heldRep = rep[rep["dataset"].isin(["validation", "test"])]
    worst = heldRep.groupby("measure")["typical_error_pct"].mean().max() if len(heldRep) else np.nan
    head = ("Yes - the proxy predicts unseen cases accurately" if worst < 2 else
            "Mostly - the proxy predicts unseen cases well, with some error" if worst < 5 else
            "With caution - the proxy has noticeable error on unseen cases")
    fig, axes = plt.subplots(1, len(T), figsize=(6 * len(T), 5.4), constrained_layout=True)
    for ax, (key, t) in zip(np.atleast_1d(axes), T.items()):
        kind, vals = t["kind"], []
        for name, col in [("validation", BLUE), ("test", "#ff7f0e")]:
            if name in held:
                m, pred = held[name]
                tru, prd = fromT(Yall[key][m][:, -1], kind), fromT(pred[key][:, -1], kind)
                ax.scatter(tru, prd, s=50, color=col, edgecolor="k", label=f"{name} cases", zorder=3)
                vals += list(tru) + list(prd)
        if vals:
            ax.plot([min(vals), max(vals)], [min(vals), max(vals)], "k--", lw=1.2, label="perfect prediction")
        r = heldRep[heldRep["measure"] == key]
        err = float(r["typical_error_pct"].mean()) if len(r) else np.nan
        word = verdictWord(err) if np.isfinite(err) else "n/a"
        ax.set_title(f"{PLAIN[key]}\ntypical error {err:.1f}% ({word})", weight="bold",
                     color={"excellent": GREEN, "good": GREEN, "fair": AMBER, "poor": RED}.get(word, "k"))
        ax.set_xlabel(f"real simulator value at the end of history [{UNITS[key]}]")
        ax.set_ylabel("proxy prediction")
        ax.grid(alpha=0.3)
        if key.endswith("_cum"):
            millions(ax)
            ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v / 1e6:.0f}M"))
    np.atleast_1d(axes)[0].legend(fontsize=9)
    fig.suptitle(f"Can we trust the proxy?  {head}", fontsize=15, weight="bold")
    fig.supxlabel("Each dot is a simulated case the proxy never saw during training. Dots on the dashed line "
                  "= prediction equals the real simulator. Error = % of the largest history value.",
                  fontsize=11)
    savePlot(fig, "2_can_we_trust_the_proxy.png")


def plotWhatHistoryTells(hm):
    """3: which settings does the real history pin down?"""
    ident, proxy = hm["ident"], hm["proxy"]
    pinned = list(ident.loc[ident["verdict"] == "PINNED DOWN", "setting"])
    free = list(ident.loc[ident["verdict"] == "NOT constrained", "setting"])
    head = ("History pins down: " + (", ".join(pinned) if pinned else "none of the settings")
            + (".  It cannot tell: " + ", ".join(free) if free else "."))
    col = {"PINNED DOWN": GREEN, "partly constrained": AMBER, "NOT constrained": RED}
    fig, ax = plt.subplots(figsize=(14, 1.6 * len(ident) + 2.2), constrained_layout=True)
    labels = []

    def position(setting, value):
        """Where `value` sits inside the tested range (0 = lowest, 1 = highest; log for LOG_PARAMS)."""
        row = np.array(hm["bestParams"], float)
        j = hm["pcols"].index(setting)
        row[j] = value
        return float(proxy.toUnit(row[None, :])[0, j])

    for i, r in enumerate(ident.iloc[::-1].itertuples()):
        lo, hi, b = (position(r.setting, v) for v in (r.acceptable_min, r.acceptable_max, r.best))
        ax.barh(i, 1, color="#e6e6e6", edgecolor="k", height=0.55)
        ax.barh(i, hi - lo, left=lo, color=col[r.verdict], edgecolor="k", height=0.55)
        ax.plot(b, i, "k*", ms=18, zorder=5)
        labels.append(f"{r.setting}\n{r.verdict}\nbest {r.best:.3g}  |  plausible {r.acceptable_min:.3g} "
                      f"to {r.acceptable_max:.3g}  |  tested {r.tested_min:.3g} to {r.tested_max:.3g}")
    ax.set_yticks(range(len(ident)))
    ax.set_yticklabels(labels, fontsize=10)
    ax.set_xlim(0, 1)
    ax.set_xticks([0, 1])
    ax.set_xticklabels(["lowest value tested", "highest value tested"])
    fig.suptitle(head, fontsize=14, weight="bold")
    fig.supxlabel("Grey bar = every value we tested.  Coloured bar = values that still match the real history "
                  "about as well as the best\n(the narrower, the better the history determines it).  "
                  "Star = single best value.", fontsize=11)
    savePlot(fig, "3_what_history_tells_us.png")


def plotHistoryMatch(hm):
    """4: the best settings against the real field."""
    T, proxy, A, bestPred, resid = hm["targets"], hm["proxy"], hm["acceptable"], hm["bestPred"], hm["resid"]
    rng = np.random.RandomState(RANDOM_SEED)
    ens = proxy.predict(A[rng.choice(len(A), min(len(A), 300), replace=False)])
    avg = 100 * np.mean(list(resid.values()))
    fig, axes = plt.subplots(1, len(T), figsize=(6 * len(T), 5), constrained_layout=True)
    for ax, (key, t) in zip(np.atleast_1d(axes), T.items()):
        kind, x = t["kind"], t["dates"]
        lo, hi = np.percentile(ens[key], [10, 90], axis=0)
        ax.fill_between(x, fromT(lo, kind), fromT(hi, kind), color="#d62728", alpha=0.25,
                        label="other settings that fit about as well")
        ax.plot(x, fromT(bestPred[key][0], kind), color="#d62728", lw=2.5, label="best match found")
        ax.plot(x, fromT(t["obs"], kind), "ko", ms=4, label="real field")
        ax.set_title(f"{PLAIN[key]}\nmismatch {100 * resid[key]:.1f}% of peak", weight="bold")
        ax.set_ylabel(UNITS[key])
        ax.tick_params(axis="x", rotation=30)
        ax.grid(alpha=0.3)
        if key.endswith("_cum"):
            millions(ax)
    np.atleast_1d(axes)[0].legend(fontsize=9, loc="upper left")
    fig.suptitle(f"History match: the best settings reproduce the real field to within about {avg:.1f}%",
                 fontsize=15, weight="bold")
    fig.supxlabel("Red line = proxy at the best settings.  Red band = the range of curves from all settings "
                  "that match nearly as well (the history cannot tell them apart).", fontsize=11)
    savePlot(fig, "4_history_match.png")


def plotForecast(data, fc):
    """5: the forecast to 2028 - what we submit and how spread out the answers are."""
    fdates, bands, histEnd, obsEnd = fc["fdates"], fc["bands"], fc["histEnd"], fc["obsEnd"]
    fig, axes = plt.subplots(2, 3, figsize=(19, 10), constrained_layout=True, sharex=True)
    cmap = plt.get_cmap("viridis")
    days = np.diff(np.concatenate([[histEnd.value],
                                   fdates.values.astype("datetime64[ns]").astype("int64")])) / 8.64e13
    for j, s in enumerate(STREAMS):
        ax, ax2 = axes[0, j], axes[1, j]
        lo, mid, hi = bands[s]
        h = data[f"{s}_cum"][HISTORY_COL].dropna()
        ax.plot(h.index, h.values, "ko", ms=3, label="real field (history)")
        ax.fill_between(fdates, lo, hi, color=BLUE, alpha=0.3, label="spread between equally good matches")
        ax.plot(fdates, mid, color=BLUE, lw=2.2, label="middle forecast")
        hr = data[f"{s}_rate"][HISTORY_COL].dropna()
        ax2.plot(hr.index, hr.values, "ko", ms=3, label="real field (history)")
        for arr, ls, lw in [(lo, "--", 1), (mid, "-", 2.2), (hi, "--", 1)]:
            prev = np.concatenate([[obsEnd[s]], arr[:-1]])
            ax2.plot(fdates, (arr - prev) / days, color=BLUE, ls=ls, lw=lw)
        n = len(fc["selCurves"][s])
        for rank, c in enumerate(fc["selCurves"][s]):
            col = cmap(rank / max(n - 1, 1))
            ax.plot(fdates, c, color=col, lw=1.3, alpha=0.95)
            prev = np.concatenate([[obsEnd[s]], c[:-1]])
            ax2.plot(fdates, (c - prev) / days, color=col, lw=1.0, alpha=0.85)
        for a in (ax, ax2):
            a.axvline(histEnd, color="k", ls=":")
            a.grid(alpha=0.3)
        millions(ax)
        ax.set_title(f"{s.capitalize()} produced so far [{UNITS[s]}]", weight="bold")
        ax2.set_title(f"{s.capitalize()} production rate [{UNITS[s]}/day]", weight="bold")
        ax.text(histEnd, ax.get_ylim()[0], " forecast starts", va="bottom", fontsize=9)
    axes[0, 0].legend(fontsize=9, loc="upper left")
    fig.suptitle("Forecast 2008 to 2028: what the field is expected to produce", fontsize=15, weight="bold")
    fig.supxlabel("Coloured lines = the 10 cases submitted (dark = best match to history).  Blue band = spread "
                  "between settings that match history equally well.\nRight of the dotted line is an "
                  "EXTRAPOLATION: the real uncertainty is larger than the band.", fontsize=11)
    savePlot(fig, "5_forecast_2008_2028.png")


def plotReliability(fc):
    """6: the forecast method tested on the past."""
    e, chosen = fc["errs"][fc["chosen"]], fc["chosen"]
    med = [100 * e[s].abs().median() for s in STREAMS]
    q75 = [100 * e[s].abs().quantile(0.75) for s in STREAMS]
    fig, ax = plt.subplots(figsize=(10, 6), constrained_layout=True)
    cols = [GREEN if m < 5 else AMBER if m < 10 else RED for m in med]
    ax.bar([s.capitalize() for s in STREAMS], med, color=cols, edgecolor="k",
           yerr=[np.zeros(len(med)), np.array(q75) - np.array(med)], capsize=8)
    for i, m in enumerate(med):
        ax.text(i, m + 0.3, f"{m:.1f}%", ha="center", weight="bold", fontsize=13)
    ax.set_ylabel(f"typical error of the {HINDCAST_YEARS:g}-year production increase [%]")
    ax.grid(alpha=0.3, axis="y")
    fig.suptitle(f"How reliable is the forecast?  Tested on the past: typical error {np.mean(med):.1f}% "
                 f"over {HINDCAST_YEARS:g} years", fontsize=14, weight="bold")
    fig.supxlabel(f"We hid the last {HINDCAST_YEARS:g} years of history, forecast them from the earlier data, and "
                  "compared with what the simulator really did (bar = median case, whisker = 75th percentile). "
                  "A 20-year forecast will be LESS accurate than this.", fontsize=11)
    savePlot(fig, "6_how_reliable_is_the_forecast.png")


# =============================================================================
# STEP 7 - THE SUMMARY EXCEL (every sheet starts with a plain-language explanation)
# =============================================================================
def writeSheet(wb, title, explanation, df, units=None):
    """One sheet: bold title, explanation text, then the table."""
    ws = wb.create_sheet(title[:31])
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = explanation
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ncol = max(len(df.columns), 4)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws.row_dimensions[2].height = 15 * (1 + len(explanation) // (14 * ncol + 40))
    r0 = 4
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(row=r0, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start=r0 + 1):
        for j, v in enumerate(row, start=1):
            if isinstance(v, (np.floating, float)):
                v = None if not np.isfinite(v) else float(v)
            elif isinstance(v, np.integer):
                v = int(v)
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            cell = ws.cell(row=i, column=j, value=v)
            if isinstance(v, float):
                cell.number_format = "#,##0" if abs(v) >= 1000 else "0.000"
            if hasattr(v, "year"):
                cell.number_format = "yyyy-mm-dd"
    for j, c in enumerate(df.columns, start=1):
        width = max([len(str(c))] + [len(f"{v:,.0f}" if isinstance(v, (int, float)) else str(v))
                                     for v in df[c].head(60)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 3, 12), 48)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    return ws


def writeSummary(hm, fc):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Read me")
    lines = [
        ("Results summary - what each sheet shows", True),
        ("", False),
        ("Goal: build a fast ML stand-in (proxy) for the reservoir simulator, use it to find the reservoir "
         "settings that reproduce the real field's history, and forecast production from 2008 to 2028.", False),
        ("", False),
        ("Sheet  ->  what it shows", True),
        ("Selected cases  ->  the 10 cases submitted in the template: their 4 settings and the production they forecast for 2028.", False),
        ("Parameter findings  ->  for each of the 4 reservoir settings: the best value and whether the history can determine it.", False),
        ("Proxy accuracy  ->  how well the proxy predicts simulator results for cases it never saw (the trust test).", False),
        ("Forecast reliability  ->  the forecast method tested on the past: hide recent history, forecast it, compare.", False),
        ("Forecast bands  ->  the 10th / middle / 90th percentile forecast for each date (the blue band in graph 5).", False),
        ("Deliverable check  ->  automatic check of the filled template (no blanks, no negatives, cumulative never decreases).", False),
        ("", False),
        ("Important limits", True),
        ("1. The simulated curves and the real field both end in 2008, so 2008-2028 is an extrapolation, tested only by the hindcast.", False),
        ("2. The water injection / well plan after 2008 is not in the data; forecasts assume past behaviour continues.", False),
        ("3. Settings that match history equally well can forecast very differently - the band shows this, but real uncertainty is larger.", False),
        ("4. Which cases are train / validation / test is assumed from the sheet counts (cases 1-70 / 71-85 / 86-100).", False),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(bold=bold, size=14 if r == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130

    sel = fc["selected"].copy()
    yr = fc["fdates"][-1].year
    sel = sel.rename(columns={"case": "Template sheet (Case N)", "misfit": "Mismatch with history (lower = closer)",
                              f"gas_cum_{yr}": f"Gas produced by {yr} [MSCF]",
                              f"oil_cum_{yr}": f"Oil produced by {yr} [STB]",
                              f"water_cum_{yr}": f"Water produced by {yr} [STB]"})
    writeSheet(wb, "Selected cases",
               "These are the cases submitted in the template. Case 1 matches the real history best. Each row is one "
               "combination of the 4 reservoir settings (columns 3-6) with the cumulative production it forecasts by "
               f"{yr}. The cases fit the history about equally well but are chosen to differ from each other "
               f"(selection mode: {SELECTION_MODE}), so their forecasts show the range of plausible outcomes.", sel)

    ident = hm["ident"].copy()
    meaning = {"PINNED DOWN": "The history determines this setting: only a narrow range matches.",
               "partly constrained": "The history narrows this setting somewhat but not fully.",
               "NOT constrained": "The history cannot tell this setting's value: any value in the range fits."}
    ident["what it means"] = ident["verdict"].map(meaning)
    ident["share_of_range"] = 100 * ident["share_of_range"]
    ident.columns = ["Setting", "Best value", "Plausible low", "Plausible high", "Tested low", "Tested high",
                     "Share of tested range still plausible [%]", "Verdict", "What it means"]
    writeSheet(wb, "Parameter findings",
               "Each of the 4 uncertain reservoir settings was searched across its tested range. 'Plausible' values "
               "are those whose simulated production matches the real history about as well as the best value. A "
               "narrow plausible range means history determines the setting; a wide one means it does not.", ident)

    rep = hm["rep"].copy()
    rep["measure"] = rep["measure"].map(lambda k: PLAIN.get(k, k))
    rep.columns = ["Dataset", "What was predicted", "Cases", "Typical error [% of largest history value]",
                   "Worst error [%]", "R2 at end of history (1 = perfect)", "Verdict"]
    writeSheet(wb, "Proxy accuracy",
               "The proxy is the fast ML stand-in for the simulator. This table scores it on simulated cases it never "
               "saw. 'validation' and 'test' are held-out cases; 'train (cross-validation)' scores each training case "
               "using a proxy that did not include it. Error is measured against the real simulator result.", rep)

    rows = []
    for w, e in fc["errs"].items():
        for s in STREAMS:
            rows.append({"Fitting window [years]": w, "Chosen": "yes" if w == fc["chosen"] else "",
                         "Stream": s.capitalize(),
                         "Typical error of the hidden production increase [%]": 100 * e[s].abs().median()})
    writeSheet(wb, "Forecast reliability",
               f"The forecast method was tested on the past: the last {HINDCAST_YEARS:g} years of history were hidden, "
               "forecast from earlier data, and compared with what the simulator actually produced. Several fitting "
               "windows were tried and the best was chosen. A 20-year forecast will be less accurate than this test.",
               pd.DataFrame(rows))

    band = {"Date": fc["fdates"]}
    for s in STREAMS:
        for nm, arr in zip(["10th percentile", "middle", "90th percentile"], fc["bands"][s]):
            band[f"{s.capitalize()} cumulative {nm} [{UNITS[s]}]"] = arr
    writeSheet(wb, "Forecast bands",
               "Cumulative production forecast for each template date. 'Middle' is the median forecast across all "
               f"{fc['nAcceptable']} settings that match history about equally well; the 10th and 90th percentiles "
               "show the spread between them. This is the blue band in graph 5.", pd.DataFrame(band))

    writeSheet(wb, "Deliverable check",
               "Automatic check of the filled-in template, read back from the saved file: every sheet must have all "
               "rows filled, no negative values, cumulative production that never decreases, and a start at or above "
               "the real 2008 history value.", fc["check"])
    path = OUTPUT_DIR / "Results_Summary.xlsx"
    wb.save(path)
    print(f"saved {path}")


# =============================================================================
# MAIN
# =============================================================================
def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"Reading from: {DATA_DIR.resolve()}")
    params = loadParameters()
    data = {key: loadProduction(sheet) for key, sheet in PRODUCTION_SHEETS.items()}
    printDataChecks(data, params)

    hm = runHistoryMatch(data, params)
    plotProblem(hm)
    plotProxyTrust(hm)
    plotWhatHistoryTells(hm)
    plotHistoryMatch(hm)

    fc = runForecast(data, params, hm)
    plotForecast(data, fc)
    plotReliability(fc)
    writeSummary(hm, fc)
    print("\nREMINDERS: (1) 2008-2028 is an extrapolation; the injection / well plan after 2008 is not in "
          "the data.\n(2) Confirm what 'Case 1..10' mean and how they are scored; change SELECTION_MODE if "
          "needed.")
    if SHOW_PLOTS:
        plt.show()


if __name__ == "__main__":
    main()