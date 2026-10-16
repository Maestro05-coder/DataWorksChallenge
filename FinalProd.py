"""
ML-assisted history matching, forecast and submission (oil & gas hackathon)  -  V6
===================================================================================

V6 = V5 (exploration + proxy history match)  +  the forecast / submission layer.

WHAT THIS SCRIPT ANSWERS
    1. "How do the 4 uncertain reservoir settings change the production curves,
        and which simulated cases look most like the real field?"        (steps 1-4)
    2. "Which settings (even between the 100 runs) match the history, and which
        settings does the history really pin down?"                      (step 5)
    3. "What does the field produce 2008 -> 2028, and what do we submit?" (step 6)

FILES (found automatically in DATA_DIR by name; case-insensitive, 'starts with')
    01 Introduction        USED     sheet 'Uncertainty Values' = the 4 settings per
                                    case (the proxy INPUTS)
    02 Field Production    USED     6 production sheets (Case_N + RAMP_History = the
                                    real field) and the 'Field Production' sheet
    09 Template Deliverable USED    a COPY is filled in; the original is never
                                    modified. If missing, the script writes its own
                                    10-sheet workbook.
    03-08 Train/Validation/Test     only the SHEET NAMES are listed (inventory);
                                    the heavy grids are not read.
    If the files are not separate, any single .xlsx in DATA_DIR containing the
    needed sheet names is used instead (e.g. one big trainingData.xlsx).

HOW THE PIECES FIT  (read once, it explains the plots)
    Steps 1-5 work INSIDE the data (1998-2008): the proxy is validated on cases it
    never saw, and the history match is a search over the 4 settings.
    Step 6 is an EXTRAPOLATION: the simulator curves and the real field both stop at
    2008. Its only honest test is the hindcast (plot 12). Expect the 20-year result
    to be less certain than the hindcast suggests.

HOW TO RUN
    pip install numpy pandas matplotlib openpyxl scipy scikit-learn
    python HistoryMatchingV6.py path/to/folder      (or edit DATA_DIR below)

PLOTS (OUTPUT_DIR) - each plot states its own verdict
    01  every simulated case vs the real field
    02  did the runs cover the 4 settings evenly?   GREEN / YELLOW / RED per pair
    03  for each result, the setting that controls it most (and the weakest links)
    03b every setting vs every result (full detail grid)
    04  who correlates with whom + printed list of STRONG links
    05  the cases closest to the real history
    06  real field production and water injection
    07  WHEN each setting matters (correlation over time)
    08  funnel + what the history says about each setting's range
    09  proxy accuracy on cases it never saw
    10  history match: best settings and the acceptable band vs the real field
    11  where the acceptable settings sit inside the tested ranges
    12  hindcast: how good is the decline-curve forecast on the past?     (NEW)
    13  forecast 2008-2028: band + the cases submitted                    (NEW)
    + <template>_FILLED.xlsx, selected_cases.csv, forecast_selected_cases.csv,
      hindcast_errors.csv, proxy_validation.csv, best_parameters.csv, ...

Needs: pandas, numpy, matplotlib (3.4+), openpyxl, scipy, scikit-learn
"""

import re
import sys
from pathlib import Path

import warnings

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.cm import ScalarMappable
from matplotlib.colors import Normalize

# NEW (V4): needed for the proxy model + continuous history match.
#   pip install scikit-learn scipy
from openpyxl import Workbook, load_workbook
from scipy.optimize import curve_fit, differential_evolution
from scipy.stats import qmc
from sklearn.decomposition import PCA
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, Matern, WhiteKernel

# =============================================================================
# CONFIG - the only part you should need to edit
# =============================================================================
DATA_DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("data")
OUTPUT_DIR = Path("History_MatchingV6")

# File names WITHOUT the .xlsx extension. Matching ignores upper/lower case, and
# also accepts a name that merely STARTS with this text (e.g. "02 Field
# Production (1)"), in case a download added a suffix.
FILES = {
    "intro":      "01 Introduction",        # Nomenclature + Uncertainty Values
    "field":      "02 Field Production",    # production sheets + Field Production
    "train":      ["03 Train Cases 1", "04 Train Cases 2",
                   "05 Train Cases 3", "06 Train Cases 4"],
    "validation": "07 Validation Cases",
    "test":       "08 Test Cases",
    "template":   "09 Template Deliverable",   # NEW (V6): the file you fill in
}

# Colour the curves by one of the 4 settings so its effect is visible.
# Set to None for plain grey curves. Must match a column in Uncertainty Values.
COLOR_BY = "Aquifer Pore Volume"

TOP_N = 5            # how many best-matching cases to highlight
SHOW_PLOTS = True    # False = only save PNGs, do not open windows

# Sheet names (matching is case-insensitive, so 'Oil production cumulative'
# and 'Oil Production Cumulative' are treated as the same sheet).
PARAM_SHEET = "Uncertainty Values"
FIELD_SHEET = "Field Production"
PRODUCTION_SHEETS = {                      # short key -> sheet name
    "gas_cum":    "Gas Production Cumulative",
    "gas_rate":   "Gas Production Rate",
    "oil_cum":    "Oil production cumulative",
    "oil_rate":   "Oil Production Rate",
    "water_cum":  "Water Production Cumulative",
    "water_rate": "Water Production Rate",
}
RATE_KEYS = ["oil_rate", "gas_rate", "water_rate"]  # used to score history match

HISTORY_COL = "RAMP_History"   # column holding the real observed field
CASE_PREFIX = "Case_"          # simulated-run columns are Case_1, Case_2, ...

# --- thresholds that turn numbers into plain-English verdicts ----------------
STRONG_RHO = 0.6     # |rank correlation| at/above this  = "STRONG link"
MODERATE_RHO = 0.3   # at/above this = "moderate"; below = "weak / no link"
MATCH_TOL = 0.05     # a case "hits" the real value if within +-5% of it
MATCH_FRACTION = 0.3 # history-match funnel: best 30% of cases count as "good"
RANGE_TOP_N = 10     # best N cases used to show what history says per setting

# --- NEW (V4): proxy model + continuous history match ------------------------
# Which Case_N belongs to train / validation / test. ASSUMED from the sheet
# counts (70 train + 15 validation + 15 test) - CONFIRM against files 03-08.
TRAIN_IDS = range(1, 71)
VAL_IDS = range(71, 86)
TEST_IDS = range(86, 101)

# Streams used to SCORE the match, and how each is transformed first.
#   "linear" = as is;  "log" = log10(x + 1) so small water rates still count.
# Gas is left out on purpose: gas follows oil (corr +1.00), so it would just
# count oil twice. Oil/gas RATES are left out because every case equals the
# history at the start (rate-controlled), so they say little.
MATCH_STREAMS = {"oil_cum": "linear", "water_cum": "linear", "water_rate": "log"}

SIM_START = None            # None = first date found in RAMP_History
LOG_PARAMS = ["Permeability Multiplier"]   # params sampled/modelled on a log scale
PCA_VAR = 0.999             # keep PCA components explaining this much variance
PCA_MAX = 6                 # ... but never more than this many
N_SOBOL_POW2 = 14           # 2**14 = 16384 trial parameter sets for the search
ACCEPT_MARGIN_MIN = 0.01    # a set is "as good as the best" if its misfit is within
ACCEPT_MARGIN_REL = 0.10    # the LARGEST of: this absolute amount, this share of the
                            # best misfit, or 2x the proxy's own error. Raise a value
                            # to accept more sets (wider ranges), lower it for fewer.
FINAL_REFIT_ON_VAL = True   # after reporting proxy quality, refit on train+val
RANDOM_SEED = 0
FIT_PROXY_QUIET = True      # silence sklearn convergence warnings

# --- NEW (V6): forecast 2008 -> 2028 and the submission template --------------
RUN_FORECAST = True
FORECAST_START_IF_NO_TEMPLATE = "2008-04-01"   # only used if no template file is found
FORECAST_END_IF_NO_TEMPLATE = "2028-01-01"
FIT_WINDOW_YEARS = None     # None = pick automatically with the hindcast
FIT_WINDOW_CANDIDATES = [2.0, 3.0, 5.0]   # years of late history the decline curves are fitted on
HINDCAST_YEARS = 3.0        # hindcast: pretend history ended this long before its real end
WOR_MAX = 50.0              # cap on the water-oil ratio (50 = 98% water cut)
ANCHOR_TO_HISTORY = True    # shift forecasts so they join the real 2008 history exactly
FORECAST_REFIT_ON_ALL = False   # False = like the history proxy: test cases never used for fitting
MIN_FORECAST_SETS = 50      # forecast band needs at least this many settings
MAX_FORECAST_SETS = 500     # ... and uses at most this many (random subsample)
N_SELECT = 10               # number of cases to submit (= sheets in the template)
# "diverse_params": best match first, then the most different settings that still fit
# "best":           simply the lowest-misfit settings (often near duplicates)
# "span_forecast":  acceptable settings spread from low to high final oil forecast
SELECTION_MODE = "diverse_params"

NICE = {"gas_cum": "Gas cumulative", "gas_rate": "Gas rate",
        "oil_cum": "Oil cumulative", "oil_rate": "Oil rate",
        "water_cum": "Water cumulative", "water_rate": "Water rate"}


# =============================================================================
# STEP 1 - LOCATE THE EXCEL FILES AND READ SHEETS FROM THEM
# =============================================================================
class SheetNotFound(Exception):
    pass


def find_file(stem):
    """Return the .xlsx in DATA_DIR whose name matches `stem`, else None."""
    files = [p for p in sorted(DATA_DIR.glob("*.xlsx")) if not p.name.startswith("~$")]
    target = stem.strip().lower()
    for p in files:                                   # exact name first
        if p.stem.strip().lower() == target:
            return p
    for p in files:                                   # then 'starts with'
        if p.stem.strip().lower().startswith(target):
            return p
    return None


_NAMES_CACHE = {}


def sheet_names(path):
    """List the sheet names of a workbook WITHOUT loading its cells.

    openpyxl only reads the workbook's table of contents here, so this is fast
    even for the big train/validation/test grid files.
    """
    key = str(path)
    if key not in _NAMES_CACHE:
        with pd.ExcelFile(path) as xl:
            _NAMES_CACHE[key] = xl.sheet_names
    return _NAMES_CACHE[key]



def find_workbook_with_sheet(sheet_name):
    """Fallback for ONE combined workbook (e.g. trainingData.xlsx): return the
    first .xlsx in DATA_DIR that contains a sheet called `sheet_name`."""
    target = sheet_name.strip().lower()
    for p in sorted(DATA_DIR.glob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        if target in [n.strip().lower() for n in sheet_names(p)]:
            return p
    return None


def read_sheet(file_stem, sheet_name, header=None):
    """Read one sheet (matched ignoring upper/lower case) from one named file."""
    path = find_file(file_stem)
    if path is None:                      # maybe everything is in ONE workbook
        path = find_workbook_with_sheet(sheet_name)
    if path is None:
        have = [p.name for p in sorted(DATA_DIR.glob("*.xlsx"))]
        raise FileNotFoundError(
            f"Could not find a file named like '{file_stem}.xlsx' in "
            f"{DATA_DIR.resolve()}.\nExcel files I can see there: {have}"
        )
    names = sheet_names(path)
    lookup = {n.strip().lower(): n for n in names}
    real = lookup.get(sheet_name.strip().lower())
    if real is None:
        raise SheetNotFound(
            f"Sheet '{sheet_name}' not found in '{path.name}'.\n"
            f"Sheets in that file: {names}"
        )
    return pd.read_excel(path, sheet_name=real, header=header)


def inventory_case_files(n_case_columns):
    """List how many sheets (= cases) each train/validation/test file holds.

    Only sheet NAMES are read, never the grids, so this is cheap. It helps pin
    down how the case sheets relate to the Case_1..Case_N columns.
    """
    groups = {"train": FILES["train"],
              "validation": [FILES["validation"]],
              "test": [FILES["test"]]}
    totals = {}
    print("\n=== CASE-FILE INVENTORY (sheet names only, grids not read) ===")
    for group, stems in groups.items():
        count = 0
        for stem in stems:
            path = find_file(stem)
            if path is None:
                print(f"  {stem:22s} not found in {DATA_DIR} -> skipped")
                continue
            names = sheet_names(path)
            count += len(names)
            print(f"  {stem:22s} {len(names):3d} sheets  e.g. {names[:3]}")
        totals[group] = count
    total = sum(totals.values())
    print(f"  case sheets found: {totals}  (total {total})")
    print(f"  Case_ columns in the production sheets: {n_case_columns}")
    if total and total != n_case_columns:
        print("  NOTE: these two numbers differ. That can be fine, but confirm how a "
              "case sheet maps to a Case_N column / a row in Uncertainty Values.")
    print("==============================================================\n")


# =============================================================================
# STEP 2 - LOAD + CLEAN THE DATA
# =============================================================================
def load_parameters():
    """Uncertainty Values -> table with one row per case, one column per dial.

    Index is renamed 'Case 1' -> 'Case_1' so it matches the production sheets.
    """
    df = read_sheet(FILES["intro"], PARAM_SHEET, header=0)
    df = df.rename(columns={df.columns[0]: "Item"})
    df.columns = [str(c).strip() for c in df.columns]
    df["Item"] = df["Item"].astype(str).str.strip().str.replace(" ", "_", regex=False)
    df = df.set_index("Item").apply(pd.to_numeric, errors="coerce")
    return df.dropna(how="all")


def load_production(sheet_name):
    """Load one production sheet into a clean (dates x cases) table.

    The sheets have slightly different header layouts, e.g.
        row 0: Date | Case_1 | Case_2 ... | RAMP_History
        row 1: (blank) | units text ...
        row 2+: data
    but Water Production Rate has an EXTRA header row on top (X, Y1, Y2 ...).
    So instead of hard-coding row numbers we locate the row whose first cell
    says 'Date' and read relative to it.

    Returns (table, unit_label) e.g. (..., 'Oil production rate [STB/d]').
    """
    raw = read_sheet(FILES["field"], sheet_name, header=None)

    header_row = raw.index[raw.iloc[:, 0].astype(str).str.strip().eq("Date")][0]
    unit_label = str(raw.iloc[header_row + 1, 1]).strip()       # the units row

    body = raw.iloc[header_row + 2:, :].copy()                   # real data rows
    body.columns = ["Date"] + [str(c).strip() for c in raw.iloc[header_row, 1:]]
    body["Date"] = pd.to_datetime(body["Date"], errors="coerce")
    body = body.dropna(subset=["Date"]).set_index("Date").sort_index()
    body = body.apply(pd.to_numeric, errors="coerce")            # text -> numbers
    # NOTE: we deliberately keep columns that are entirely empty. A Case_N
    # column with no data may be a failed simulation, and the sanity checks
    # below report it instead of hiding it.
    body = body.loc[:, ~body.columns.isin(["nan", ""])]          # junk header cells
    return body, unit_label


def case_cols(df):
    """Names of the simulated-run columns (Case_1 ... Case_N)."""
    return [c for c in df.columns if str(c).startswith(CASE_PREFIX)]


def last_valid(series):
    """Last non-empty value of a curve (its value at the final date)."""
    s = series.dropna()
    return s.iloc[-1] if len(s) else np.nan


def interpolate_to(series, dates):
    """Evaluate one case's curve at the requested dates.

    Each case may report only on ITS OWN dates (the rest of its column is
    empty), so cases cannot be compared row by row. We keep the case's valid
    points and linearly interpolate between them. Dates outside the range the
    case covers become NaN (no extrapolation).
    """
    s = series.dropna()
    if len(s) < 2:
        return pd.Series(np.nan, index=dates)
    x = s.index.values.astype("datetime64[ns]").astype("int64")
    xi = pd.DatetimeIndex(dates).values.astype("datetime64[ns]").astype("int64")
    return pd.Series(np.interp(xi, x, s.values, left=np.nan, right=np.nan), index=dates)


# =============================================================================
# STEP 3 - SANITY CHECKS (catch data problems BEFORE modelling)
# =============================================================================
def run_sanity_checks(data, params):
    print("\n=== SANITY CHECKS ===")
    reference_cases = None
    for key, df in data.items():
        cc = case_cols(df)
        n_hist = df[HISTORY_COL].notna().sum() if HISTORY_COL in df else "MISSING"
        span = f"{df.index.min().date()} -> {df.index.max().date()}" if len(df) else "no dates"
        nan_pct = df[cc].isna().mean().mean() * 100 if cc else float("nan")
        print(f"{key:11s} dates={len(df):4d} ({span}) cases={len(cc):3d} "
              f"history_points={n_hist} empty_cells={nan_pct:.1f}%")

        # Cases with NO data at all (e.g. a failed simulation run).
        empty = [c for c in cc if df[c].isna().all()]
        if empty:
            print(f"   WARNING: case columns with no data at all: {empty}")
        if HISTORY_COL in df and df[HISTORY_COL].isna().all():
            print("   WARNING: RAMP_History is completely empty in this sheet")

        # How many dates does each case actually report on?
        if cc:
            pts = df[cc].notna().sum()
            print(f"   points per case: min={pts.min()} median={int(pts.median())} "
                  f"max={pts.max()} (out of {len(df)} dates)")
            if pts.median() < 0.5 * len(df):
                print("   NOTE: cases are SPARSE (each reports on its own dates). "
                      "Plots draw lines through each case's own points, and "
                      "scoring interpolates onto the history dates.")

        # Cumulative production can never go down; if it does, the data is odd.
        # (checked on each case's own valid points, ignoring the gaps)
        if key.endswith("_cum") and cc:
            n_drops = int(sum((df[c].dropna().diff() < -1e-6).sum() for c in cc))
            if n_drops:
                print(f"   WARNING: {n_drops} places where cumulative DECREASES")

        # Every sheet should contain the same set of cases.
        if reference_cases is None:
            reference_cases = set(cc)
        elif set(cc) != reference_cases:
            print(f"   WARNING: case columns differ from the first sheet "
                  f"(difference: {len(set(cc) ^ reference_cases)} columns)")

    # Every case should have its 4 settings in Uncertainty Values.
    missing = sorted(reference_cases - set(params.index)) if reference_cases else []
    print(f"parameters: {params.shape[0]} cases x {params.shape[1]} settings "
          f"-> {list(params.columns)}")
    if missing:
        print(f"   WARNING: {len(missing)} cases have no parameter row "
              f"(e.g. {missing[:3]}). Is the Uncertainty Values sheet complete?")
    print("=====================\n")


# =============================================================================
# STEP 4 - PLOTS
# =============================================================================
def save(fig, name):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    fig.savefig(OUTPUT_DIR / name, dpi=130, bbox_inches="tight")
    print(f"saved {OUTPUT_DIR / name}")


def plot_ensemble(data, labels, params):
    """Plot 1: reproduce the competition figure (6 panels).

    Every thin line is one simulated case. Black dots are the real field.
    If COLOR_BY is set, lines are coloured by that setting's value, so you can
    SEE which curves it pushes up or down.
    """
    order = ["gas_cum", "gas_rate", "oil_cum", "oil_rate", "water_cum", "water_rate"]
    fig, axes = plt.subplots(2, 3, figsize=(18, 9), constrained_layout=True)

    use_color = COLOR_BY is not None and COLOR_BY in params.columns
    if use_color:
        vals = params[COLOR_BY]
        norm = Normalize(vals.min(), vals.max())
        cmap = plt.get_cmap("viridis")

    for ax, key in zip(axes.ravel(), order):
        df = data[key]
        # rates are step-like in the simulator, cumulatives are smooth
        style = "steps-post" if key.endswith("_rate") else "default"
        for c in case_cols(df):
            if use_color and c in params.index and pd.notna(params.loc[c, COLOR_BY]):
                color = cmap(norm(params.loc[c, COLOR_BY]))
            else:
                color = "lightgrey"
            s_c = df[c].dropna()        # drop gaps so the line connects the points
            ax.plot(s_c.index, s_c.values, color=color, lw=0.8, alpha=0.7,
                    drawstyle=style)
        if HISTORY_COL in df:
            ax.plot(df.index, df[HISTORY_COL], "ko", ms=3, label="History")
        ax.set_title(labels[key], fontsize=10)
        ax.tick_params(axis="x", rotation=30)
        ax.grid(alpha=0.3)

    axes[0, 0].legend(loc="upper left")
    if use_color:
        fig.colorbar(ScalarMappable(norm=norm, cmap=cmap),
                     ax=axes.ravel().tolist(), shrink=0.6, label=COLOR_BY)
    fig.suptitle("Simulated cases vs real field history"
                 + (f"  (colour = {COLOR_BY})" if use_color else ""))
    save(fig, "01_ensemble_curves.png")


def pair_design_stats(x, y, bins):
    """Quality of the sampling of TWO settings: do they move together, and are
    there empty regions on the map of (x, y)?"""
    rho = x.rank().corr(y.rank())
    h, _, _ = np.histogram2d(x, y, bins=bins)
    empty = float((h == 0).mean())
    reasons = []
    if abs(rho) >= 0.3:
        verdict = "PROBLEM"
        reasons.append("the two settings move together")
    elif abs(rho) >= 0.15:
        verdict = "CHECK"
        reasons.append("slight tendency to move together")
    else:
        verdict = "OK"
    if empty > 0.25:
        verdict = "PROBLEM"
        reasons.append("large empty regions")
    elif empty > 0.12 and verdict == "OK":
        verdict = "CHECK"
        reasons.append("a few empty regions")
    return {"rho": rho, "empty": empty, "verdict": verdict,
            "why": "; ".join(reasons) if reasons else "evenly spread, independent"}


def plot_parameter_coverage(params):
    """Plot 2: did the runs cover the settings evenly?

    Lower-left : the real map of where the simulator was run (one dot = one
                 case), background colour = verdict for that pair.
    Upper-right: the same pair in numbers + a plain-English verdict.
    Diagonal   : how the values of ONE setting are spread (should be flat).
    GREEN = fine, YELLOW = look closer, RED = the ML model will struggle there.
    """
    p = params.dropna()
    cols = list(p.columns)
    n = len(cols)
    bins = 5 if len(p) >= 50 else 3
    tint = {"OK": "#d4edda", "CHECK": "#fff3cd", "PROBLEM": "#f8d7da"}
    edge = {"OK": "#1a7f37", "CHECK": "#b8860b", "PROBLEM": "#b22222"}
    fig, axes = plt.subplots(n, n, figsize=(3.2 * n, 3.0 * n), constrained_layout=True)
    rows = []
    for i, yc in enumerate(cols):
        for j, xc in enumerate(cols):
            ax = axes[i, j]
            if i == j:
                ax.hist(p[xc], bins=10, color="#9ecae1", edgecolor="k", lw=0.4)
                ax.set_title(xc, fontsize=10, weight="bold")
                ax.set_yticks([])
            else:
                st = pair_design_stats(p[xc], p[yc], bins)
                if i > j:                                   # lower: the map
                    ax.set_facecolor(tint[st["verdict"]])
                    ax.scatter(p[xc], p[yc], s=10, alpha=0.8, color="#1f77b4")
                    for sp in ax.spines.values():
                        sp.set_edgecolor(edge[st["verdict"]]); sp.set_linewidth(2)
                else:                                       # upper: the verdict
                    rows.append({"setting_x": xc, "setting_y": yc, "corr": st["rho"],
                                 "empty_share": st["empty"], "verdict": st["verdict"],
                                 "why": st["why"]})
                    ax.set_facecolor(tint[st["verdict"]])
                    ax.set_xticks([]); ax.set_yticks([])
                    ax.text(0.5, 0.68, st["verdict"], ha="center", va="center",
                            fontsize=20, weight="bold", color=edge[st["verdict"]],
                            transform=ax.transAxes)
                    ax.text(0.5, 0.33,
                            f"{xc}\nvs {yc}\ncorrelation {st['rho']:+.2f}\n"
                            f"empty map area {st['empty']:.0%}\n{st['why']}",
                            ha="center", va="center", fontsize=8, transform=ax.transAxes)
            if i == n - 1 and i != j:
                ax.set_xlabel(xc)
            if j == 0 and i != j:
                ax.set_ylabel(yc)
            ax.tick_params(labelsize=7)
    fig.suptitle("Plot 2 - Did the runs cover the settings evenly?\n"
                 "GREEN = good  |  YELLOW = look closer  |  RED = ML model will struggle "
                 "(gaps = it must guess, correlated = cannot separate the two effects)",
                 fontsize=12)
    save(fig, "02_parameter_coverage.png")
    table = pd.DataFrame(rows)
    table.to_csv(OUTPUT_DIR / "parameter_design_check.csv", index=False)
    print("\nParameter design check (are the runs well spread?):")
    for r in rows:
        print(f"  {r['verdict']:8s} {r['setting_x']} vs {r['setting_y']}: "
              f"corr {r['corr']:+.2f}, empty map {r['empty_share']:.0%} -> {r['why']}")


def build_summary(data, params):
    """One row per case: its 4 settings + its FINAL value for each output."""
    # "Final" = value at the last date the real field has data, so every case
    # is compared at the same moment (cases that stop earlier become NaN).
    finals = {}
    for key, df in data.items():
        ref_date = df[HISTORY_COL].dropna().index.max()
        finals[f"final_{key}"] = pd.Series({
            c: interpolate_to(df[c], [ref_date]).iloc[0] for c in case_cols(df)
        })
    finals = pd.DataFrame(finals)
    summary = params.join(finals, how="inner")
    summary.to_csv(OUTPUT_DIR / "case_summary.csv")
    return summary


def plot_param_vs_outcome(data, labels, summary, param_cols):
    """Plot 3b (detail view): EVERY setting vs EVERY final result.

    A clear trend means that setting controls that result. A shapeless cloud
    means it barely matters. The dashed line is the REAL field's final value,
    so you can see which settings land near it.
    """
    keys = list(data)
    fig, axes = plt.subplots(len(keys), len(param_cols),
                             figsize=(16, 2.6 * len(keys)), constrained_layout=True,
                             squeeze=False)
    for i, key in enumerate(keys):
        hist_final = last_valid(data[key][HISTORY_COL]) if HISTORY_COL in data[key] else np.nan
        for j, p in enumerate(param_cols):
            ax = axes[i, j]
            ax.scatter(summary[p], summary[f"final_{key}"], s=14, alpha=0.7)
            ax.axhline(hist_final, color="k", ls="--", lw=1)
            if i == len(keys) - 1:
                ax.set_xlabel(p)
            if j == 0:
                ax.set_ylabel(labels[key], fontsize=7)
    fig.suptitle("Each setting vs each final result (dashed = real field)")
    save(fig, "03b_all_pairs_grid.png")


def strength_word(r):
    a = abs(r)
    return "STRONG" if a >= STRONG_RHO else ("moderate" if a >= MODERATE_RHO else "weak")


def nice_out(col):
    """'final_water_cum' -> 'Water cumulative'."""
    return NICE.get(col.replace("final_", ""), col)


STRENGTH_COLOR = {"STRONG": "#1a7f37", "moderate": "#b8860b", "weak": "#b22222"}


def _link_panel(ax, data, summary, p, o, r, show_legend=False):
    """One big setting-vs-result panel with a plain-English title."""
    key = o.replace("final_", "")
    x, y = summary[p], summary[o]
    m = x.notna() & y.notna()
    ax.scatter(x[m], y[m], s=18, alpha=0.7, color="#1f77b4", label="one simulated case")
    # typical trend: median of y inside 6 slices of x
    edges = np.quantile(x[m], np.linspace(0, 1, 7))
    cx, cy = [], []
    for a, b in zip(edges[:-1], edges[1:]):
        sel = m & (x >= a) & (x <= b)
        if sel.sum() >= 2:
            cx.append(x[sel].median()); cy.append(y[sel].median())
    ax.plot(cx, cy, color="#ff7f0e", lw=2.8, label="typical trend (median)")
    H = last_valid(data[key][HISTORY_COL]) if HISTORY_COL in data[key] else np.nan
    if np.isfinite(H):
        ax.axhline(H, color="k", ls="--", lw=1.3, label="real field")
        tol = MATCH_TOL * max(abs(H), 0.1 * (y[m].max() - y[m].min()))
        near = m & ((y - H).abs() <= tol)
        if near.any():
            ax.scatter(x[near], y[near], s=50, color="#d62728", edgecolor="k",
                       zorder=5, label=f"within {MATCH_TOL:.0%} of real")
            ax.axvspan(x[near].min(), x[near].max(), color="#d62728", alpha=0.08)
    w = strength_word(r)
    effect = ("no clear effect" if w == "weak" else
              f"higher setting gives {'HIGHER' if r > 0 else 'LOWER'} result")
    ax.set_title(f"{p} \u2192 {nice_out(o)}\n{w} link ({effect}, corr {r:+.2f})",
                 fontsize=10, weight="bold", color=STRENGTH_COLOR[w])
    ax.set_xlabel(f"{p}  (setting)")
    ax.set_ylabel(f"{nice_out(o)} at end of history")
    ax.grid(alpha=0.3)
    if show_legend:
        ax.legend(fontsize=7.5, loc="best")


def plot_top_links(data, labels, summary, param_cols):
    """Plot 3: for EVERY result, the setting that controls it most, drawn big.

    Top block   : strongest setting for each result (sorted strongest first).
    Bottom row  : the 3 weakest setting/result pairs, for contrast.
    Orange line = typical trend. Dashed = real field. Red = cases that land on
    the real value, and the pink band = the setting range those cases occupy.
    """
    out_cols = [c for c in summary.columns if c.startswith("final_")]
    corr = summary[param_cols + out_cols].rank().corr()
    best, allp = [], []
    for o in out_cols:
        r = corr.loc[param_cols, o].dropna()
        if len(r):
            best.append((r.abs().idxmax(), o, r[r.abs().idxmax()]))
        allp += [(p, o, corr.loc[p, o]) for p in param_cols if pd.notna(corr.loc[p, o])]
    best.sort(key=lambda t: -abs(t[2]))
    weak = sorted(allp, key=lambda t: abs(t[2]))[:3]

    n_top_rows = int(np.ceil(len(best) / 3))
    fig = plt.figure(figsize=(17, 4.6 * (n_top_rows + 1)), constrained_layout=True)
    top, bot = fig.subfigures(2, 1, height_ratios=[n_top_rows, 1])
    top.suptitle("STRONGEST LINKS - for each result, the setting that controls it most "
                 "(these are what the history can pin down)", fontsize=12, weight="bold")
    axes = top.subplots(n_top_rows, 3, squeeze=False)
    for k, ax in enumerate(axes.ravel()):
        if k < len(best):
            _link_panel(ax, data, summary, *best[k], show_legend=(k == 0))
        else:
            ax.axis("off")
    bot.suptitle("WEAKEST LINKS - changing this setting barely changes this result "
                 "(history cannot tell you its value)", fontsize=12, weight="bold")
    for ax, t in zip(bot.subplots(1, 3, squeeze=False).ravel(), weak):
        _link_panel(ax, data, summary, *t)
    fig.suptitle("Plot 3 - Which settings control which results?", fontsize=14, weight="bold")
    save(fig, "03_top_links.png")


def plot_correlation(summary, param_cols):
    """Plot 4: who correlates with whom - EVERYTHING against EVERYTHING.

    Rank correlation: +1 = both rise together, -1 = one rises as the other
    falls, 0 = unrelated. Ranks are used so curved but steadily rising
    relationships still count. Bold boxed numbers are STRONG links.
    """
    out_cols = [c for c in summary.columns if c.startswith("final_")]
    cols = param_cols + out_cols
    corr = summary[cols].rank().corr()
    corr.to_csv(OUTPUT_DIR / "correlation_all.csv")
    corr.loc[param_cols, out_cols].to_csv(OUTPUT_DIR / "param_outcome_correlation.csv")
    n, k = len(cols), len(param_cols)
    names = param_cols + [nice_out(c) for c in out_cols]

    fig, ax = plt.subplots(figsize=(1.0 * n + 4, 0.9 * n + 1.5), constrained_layout=True)
    masked = np.ma.array(corr.values, mask=np.triu(np.ones((n, n), bool)))
    im = ax.imshow(masked, cmap="RdBu_r", vmin=-1, vmax=1)
    ax.set_xticks(range(n)); ax.set_xticklabels(names, rotation=40, ha="right")
    ax.set_yticks(range(n)); ax.set_yticklabels(names)
    for (i, j), v in np.ndenumerate(corr.values):
        if i <= j or np.isnan(v):
            continue
        strong = abs(v) >= STRONG_RHO
        ax.text(j, i, f"{v:+.2f}", ha="center", va="center", fontsize=10,
                weight="bold" if strong else "normal",
                color="white" if abs(v) > 0.75 else ("0.6" if abs(v) < MODERATE_RHO else "black"))
        if strong:
            ax.add_patch(plt.Rectangle((j - .5, i - .5), 1, 1, fill=False, ec="k", lw=2))
    # coloured outlines around the three blocks
    blocks = [((-.5, -.5), k, k, "#e67e22"),            # settings vs settings
              ((-.5, k - .5), k, n - k, "#2c3e50"),     # settings vs results
              ((k - .5, k - .5), n - k, n - k, "#27ae60")]  # results vs results
    for (x0, y0), w, h, c in blocks:
        ax.add_patch(plt.Rectangle((x0, y0), w, h, fill=False, ec=c, lw=3.5, ls="--"))
    guide = ("HOW TO READ\n"
             "Each number = how two columns move together.\n"
             "+1 (red): both go up together.  -1 (blue): one up, other down.\n"
             "0 (white / grey text): unrelated.   Boxed bold = STRONG (|corr| >= "
             f"{STRONG_RHO}).\n\n"
             "ORANGE block: setting vs setting  -> should be ~0\n"
             "   (runs were sampled independently; big = bad design)\n"
             "DARK block: setting vs result  -> what you can tune and\n"
             "   what it controls. This is the block that matters.\n"
             "GREEN block: result vs result  -> big numbers = two results\n"
             "   carry the same information (redundant).")
    ax.text(0.98, 0.98, guide, transform=ax.transAxes, ha="right", va="top", fontsize=9.5,
            bbox=dict(fc="#f8f9fa", ec="k"))
    fig.colorbar(im, ax=ax, label="rank correlation", shrink=0.7)
    ax.set_title("Plot 4 - Who correlates with whom?", fontsize=13, weight="bold")
    save(fig, "04_param_outcome_correlation.png")
    report_links(corr, param_cols, out_cols)


def report_links(corr, param_cols, out_cols):
    """Print (and save) every link in plain English, strongest first."""
    cols = param_cols + out_cols
    rows = []
    for a_i, a in enumerate(cols):
        for b in cols[a_i + 1:]:
            r = corr.loc[a, b]
            if pd.isna(r):
                continue
            if a in param_cols and b in param_cols:
                kind = "setting-setting"
            elif a in param_cols:
                kind = "setting->result"
            else:
                kind = "result-result"
            rows.append({"kind": kind, "a": a, "b": b, "corr": r,
                         "strength": strength_word(r)})
    links = pd.DataFrame(rows)
    links["abs"] = links["corr"].abs()
    links = links.sort_values("abs", ascending=False).drop(columns="abs")
    links.to_csv(OUTPUT_DIR / "links_ranked.csv", index=False)

    def nm(c):
        return nice_out(c) if c.startswith("final_") else c

    def sentence(r):
        if r["kind"] == "setting->result":
            d = "UP" if r["corr"] > 0 else "DOWN"
            return (f"  {r['strength']:8s} {nm(r['a'])} UP  ->  {nm(r['b'])} {d}"
                    f"   (corr {r['corr']:+.2f})")
        how = "rise and fall together" if r["corr"] > 0 else "move in OPPOSITE directions"
        return (f"  {r['strength']:8s} {nm(r['a'])} and {nm(r['b'])} {how}"
                f"   (corr {r['corr']:+.2f})")

    print("\n=== STRONG LINKS (|corr| >= %.1f) ===" % STRONG_RHO)
    for kind, title in [("setting->result", "Settings that drive results"),
                        ("result-result", "Results that carry the same information"),
                        ("setting-setting", "Settings that move together (design problem if any)")]:
        sub = links[(links["kind"] == kind) & (links["strength"] == "STRONG")]
        print(f"\n{title}:" + ("  none" if sub.empty else ""))
        for _, r in sub.iterrows():
            print(sentence(r))
    print("\nSettings with no strong effect on ANY result "
          "(history is unlikely to pin these down):")
    found = False
    for p in param_cols:
        mx = links[(links["kind"] == "setting->result") &
                   ((links["a"] == p) | (links["b"] == p))]["corr"].abs().max()
        if pd.notna(mx) and mx < STRONG_RHO:
            found = True
            print(f"  {p}  (strongest effect on any result: corr {mx:.2f})")
    if not found:
        print("  none - every setting strongly drives at least one result")


def misfit_by_stream(data):
    """Misfit of every case against history, SEPARATELY for oil, gas, water rate.

    For each stream: RMSE between the case and the history over the dates where
    history exists, divided by the largest history value of that stream (so
    oil, gas and water count equally despite having different units/sizes).
    Returns a table: one row per case, one column per stream (lower = better).
    """
    scores = {}
    for key in RATE_KEYS:
        df = data[key]
        if HISTORY_COL not in df or df[HISTORY_COL].notna().sum() == 0:
            raise ValueError(f"No usable '{HISTORY_COL}' in the {key} sheet, "
                             "cannot score cases against history.")
        hist = df[HISTORY_COL]
        h = hist[hist.notna()]
        rmse = pd.Series({
            c: np.sqrt(np.nanmean((interpolate_to(df[c], h.index).values - h.values) ** 2))
            if interpolate_to(df[c], h.index).notna().any() else np.nan
            for c in case_cols(df)
        })
        scores[key] = rmse / max(h.abs().max(), 1e-9)
    return pd.DataFrame(scores)


def rank_cases_by_misfit(data):
    """Total misfit per case = oil + gas + water scores added up (lower = better).

    A case that is empty in any stream gets NaN overall and is dropped instead
    of looking artificially good.
    """
    return misfit_by_stream(data).sum(axis=1, skipna=False).dropna().sort_values()


def sensitivity_over_time(df, params, n_grid=40):
    """Rank correlation between each setting and the result, date by date."""
    cc = [c for c in case_cols(df) if c in params.index]
    hist = df[HISTORY_COL].dropna().index if HISTORY_COL in df else df.index
    grid = pd.date_range(hist.min(), hist.max(), periods=n_grid)
    mat = pd.DataFrame({c: interpolate_to(df[c], grid).values for c in cc}, index=grid)
    out = pd.DataFrame(np.nan, index=grid, columns=params.columns)
    for p in params.columns:
        for d in grid:
            y = mat.loc[d]
            ok = y.notna() & params.loc[cc, p].notna()
            if ok.sum() >= 10 and y[ok].nunique() > 1:
                out.loc[d, p] = params.loc[cc, p][ok].rank().corr(y[ok].rank())
    return out


def plot_sensitivity_over_time(data, params):
    """Plot 7: WHEN does each setting matter?

    One line per setting. Line near 0 = no influence at that date. Line in the
    green band = strong influence (up = bigger setting gives bigger result,
    down = bigger setting gives smaller result). Lines that start at 0 and fan
    out show an effect that builds up over time - important for forecasting.
    """
    order = [k for k in ["oil_rate", "gas_rate", "water_rate",
                         "oil_cum", "gas_cum", "water_cum"] if k in data]
    fig, axes = plt.subplots(2, 3, figsize=(18, 9), constrained_layout=True, sharey=True)
    colors = dict(zip(params.columns, plt.get_cmap("tab10")(np.arange(len(params.columns)))))
    for ax, key in zip(axes.ravel(), order):
        sens = sensitivity_over_time(data[key], params)
        ax.axhspan(STRONG_RHO, 1.05, color="green", alpha=0.10)
        ax.axhspan(-1.05, -STRONG_RHO, color="green", alpha=0.10)
        ax.axhline(0, color="k", lw=0.8)
        for p in sens:
            ax.plot(sens.index, sens[p], lw=2.4, color=colors[p], label=p)
        ax.set_ylim(-1.05, 1.05)
        ax.set_title(NICE[key], fontsize=12, weight="bold")
        ax.tick_params(axis="x", rotation=30)
        ax.grid(alpha=0.3)
        ax.text(0.01, 0.97, "STRONG: bigger setting -> bigger result", transform=ax.transAxes,
                fontsize=8, va="top", color="green")
        ax.text(0.01, 0.03, "STRONG: bigger setting -> smaller result", transform=ax.transAxes,
                fontsize=8, va="bottom", color="green")
    h, l = axes[0, 0].get_legend_handles_labels()
    fig.legend(h, l, loc="outside lower center", ncol=len(l), fontsize=10)
    for ax in axes[:, 0]:
        ax.set_ylabel("correlation of setting with result")
    fig.suptitle("Plot 7 - WHEN does each setting matter?  (line far from 0 = that setting "
                 "controls the result at that date)", fontsize=13, weight="bold")
    save(fig, "07_sensitivity_over_time.png")


def plot_history_match_summary(data, params):
    """Plot 8: what does matching the history actually tell us?

    Left : funnel - how many cases stay 'good' as you demand a match on more
           streams (oil, then gas, then water).
    Right: for each setting, the full range tested (grey) vs the range used by
           the best RANGE_TOP_N cases (colour). Narrow colour bar = history pins
           that setting down; full-width bar = history cannot tell.
    """
    ms = misfit_by_stream(data).dropna()
    thresh = ms.quantile(MATCH_FRACTION)
    run = pd.Series(True, index=ms.index)
    labels_f, counts = [], []
    for key in ms.columns:
        run = run & (ms[key] <= thresh[key])
        labels_f.append(("good on " if not labels_f else "... AND ") + NICE[key].lower())
        counts.append(int(run.sum()))
    total = ms.sum(axis=1).sort_values()
    best = total.index[:RANGE_TOP_N]
    pr = params.loc[[c for c in best if c in params.index]]

    fig, (a1, a2) = plt.subplots(1, 2, figsize=(17, 5.6), constrained_layout=True,
                                 gridspec_kw={"width_ratios": [1, 1.5]})
    y = np.arange(len(counts))[::-1]
    a1.barh(y, counts, color="#6baed6", edgecolor="k")
    for yy, c in zip(y, counts):
        a1.text(c + 0.5, yy, f"{c} cases", va="center", fontsize=11, weight="bold")
    a1.set_yticks(y); a1.set_yticklabels(labels_f, fontsize=10)
    a1.set_xlim(0, max(counts) * 1.25)
    a1.set_xlabel(f"number of cases (out of {len(ms)})")
    a1.set_title(f"Funnel: 'good' = best {MATCH_FRACTION:.0%} of cases on that stream.\n"
                 "Every extra stream you must match removes more cases", fontsize=11)
    a1.grid(alpha=0.3, axis="x")

    rows = []
    names = list(params.columns)
    for i, p in enumerate(names[::-1]):
        lo, hi = params[p].min(), params[p].max()
        span = max(hi - lo, 1e-12)
        b0, b1 = pr[p].min(), pr[p].max()
        w = (b1 - b0) / span
        verdict = "PINNED DOWN" if w < 0.5 else ("partly constrained" if w < 0.8 else "NOT constrained")
        col = {"PINNED DOWN": "#1a7f37", "partly constrained": "#b8860b",
               "NOT constrained": "#b22222"}[verdict]
        a2.barh(i, 1, color="#e0e0e0", edgecolor="k")
        a2.barh(i, w, left=(b0 - lo) / span, color=col, edgecolor="k")
        a2.plot((pr[p].iloc[0] - lo) / span, i, "k*", ms=16, zorder=5)
        rows.append({"setting": p, "tested_min": lo, "tested_max": hi,
                     f"best{RANGE_TOP_N}_min": b0, f"best{RANGE_TOP_N}_max": b1,
                     "share_of_range_used": w, "verdict": verdict})
    a2.set_yticks(range(len(names)))
    a2.set_yticklabels([
        f"{r['setting']}  [{r['verdict']}]\nbest {RANGE_TOP_N}: {r[f'best{RANGE_TOP_N}_min']:.3g} to "
        f"{r[f'best{RANGE_TOP_N}_max']:.3g}   (tested {r['tested_min']:.3g} to {r['tested_max']:.3g})"
        for r in rows], fontsize=9)
    a2.set_xlim(0, 1); a2.set_xlabel("position inside the tested range (0 = lowest tested, 1 = highest)")
    a2.set_title(f"What history says about each setting: range used by the best {RANGE_TOP_N} cases\n"
                 "(grey = everything tested, colour = best cases, star = single best case)", fontsize=11)
    fig.suptitle("Plot 8 - What does matching history actually tell us?   "
                 "(rough guide only: just ~100 runs)", fontsize=13, weight="bold")
    save(fig, "08_history_match_summary.png")
    pd.DataFrame(rows[::-1]).to_csv(OUTPUT_DIR / "history_match_ranges.csv", index=False)
    print("\nWhat the history says about each setting (best %d cases):" % RANGE_TOP_N)
    for r in rows[::-1]:
        print(f"  {r['verdict']:18s} {r['setting']}: best cases use {r[f'best{RANGE_TOP_N}_min']:.3g} "
              f"to {r[f'best{RANGE_TOP_N}_max']:.3g} (tested {r['tested_min']:.3g} to {r['tested_max']:.3g})")


def plot_best_matches(data, labels, scores, params):
    """Plot 5: the TOP_N cases closest to history, drawn over all the others."""
    best = list(scores.index[:TOP_N])
    colors = plt.get_cmap("tab10")(np.arange(len(best)))
    fig, axes = plt.subplots(1, len(RATE_KEYS), figsize=(18, 5), constrained_layout=True)

    for ax, key in zip(np.atleast_1d(axes), RATE_KEYS):
        df = data[key]
        for c in case_cols(df):
            s_c = df[c].dropna()
            ax.plot(s_c.index, s_c.values, color="lightgrey", lw=0.6)
        for c, col in zip(best, colors):
            s_c = df[c].dropna()
            ax.plot(s_c.index, s_c.values, color=col, lw=1.8, label=c)
        ax.plot(df.index, df[HISTORY_COL], "ko", ms=3, label="History")
        ax.set_title(labels[key])
        ax.tick_params(axis="x", rotation=30)
        ax.grid(alpha=0.3)
    np.atleast_1d(axes)[0].legend(fontsize=8)
    fig.suptitle(f"Top {TOP_N} cases closest to history")
    save(fig, "05_best_matching_cases.png")

    table = params.reindex(best).assign(misfit_score=scores.loc[best],
                                        split=[split_of(c) for c in best])   # V4: which split
    scores.rename("misfit_score").to_csv(OUTPUT_DIR / "misfit_ranking.csv")
    print(f"\nTop {TOP_N} cases by misfit (lower = closer to history):")
    print(table.round(4).to_string())


def plot_field_history():
    """Plot 6 (optional): the real field's rates, incl. WATER INJECTION.

    Injection is a control the operators set, not a simulator result, so it
    matters for the forecast period.
    """
    try:
        df = read_sheet(FILES["field"], FIELD_SHEET, header=0)
    except SheetNotFound:
        print(f"(skipping field history plot: sheet '{FIELD_SHEET}' not found)")
        return
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df = df.dropna(subset=["Date"]).set_index("Date")
    df = df.drop(columns=["Identifier"], errors="ignore").apply(pd.to_numeric, errors="coerce")

    fig, ax = plt.subplots(figsize=(10, 5), constrained_layout=True)
    df.plot(ax=ax, marker="o", ms=3)
    ax.set_title("Real field production and water injection")
    ax.grid(alpha=0.3)
    save(fig, "06_field_history.png")


# =============================================================================
# STEP 5 (NEW in V4) - PROXY MODEL + CONTINUOUS HISTORY MATCH
# =============================================================================
# What the code above does: rank the ~100 simulations that already exist.
# What this section adds:
#   a) a PROXY (fast stand-in for the simulator): 4 settings -> curves
#   b) a fairer score (smooth cumulatives + log water rate, no double-counted gas)
#   c) a SEARCH over all possible settings (not only the 100 runs)
#   d) the best setting AND every setting that fits "about as well" (so you can
#      see which settings the history really pins down)
#   e) plots of the match and of the proxy's accuracy
# Train cases fit the proxy, validation cases score it, test cases are only
# reported (never used for fitting or choosing anything).

def case_id(name):
    return int(str(name).split("_")[-1])


def split_of(name):
    i = case_id(name)
    if i in TRAIN_IDS:
        return "train"
    if i in VAL_IDS:
        return "validation"
    if i in TEST_IDS:
        return "test"
    return "unassigned"


def to_t(x, kind):
    """Transform a curve before scoring/modelling ('log' keeps small water visible)."""
    x = np.asarray(x, float)
    return np.log10(np.clip(x, 0, None) + 1.0) if kind == "log" else x


def from_t(y, kind):
    y = np.asarray(y, float)
    return 10.0 ** y - 1.0 if kind == "log" else y


def eval_case(series, dates, is_cum, t0):
    """Value of ONE case's curve at the history dates (fixes the old linear-for-all).

    cumulative : linear between reported points, starting from 0 at SIM_START
    rate       : the NEAREST reported value (rates are step-like, so no smoothing)
    Dates outside what the case reports become NaN (no extrapolation).
    """
    s = series.dropna()
    if len(s) < 2:
        return np.full(len(dates), np.nan)
    ns = lambda idx: pd.DatetimeIndex(idx).values.astype("datetime64[ns]").astype("int64").astype(float)
    x, y, xi = ns(s.index), s.values.astype(float), ns(dates)
    if is_cum:
        if x[0] > t0:                                   # cumulative is 0 at the start
            x, y = np.r_[t0, x], np.r_[0.0, y]
        return np.interp(xi, x, y, left=np.nan, right=np.nan)
    hi = np.clip(np.searchsorted(x, xi), 0, len(x) - 1)
    lo = np.clip(hi - 1, 0, len(x) - 1)
    pick = np.where(np.abs(xi - x[lo]) <= np.abs(x[hi] - xi), lo, hi)
    out = y[pick]
    out[(xi < x[0]) | (xi > x[-1])] = np.nan
    return out


def build_targets(data, cases, t0):
    """For each scored stream: the history, and every case's curve on the SAME
    dates (the dates RAMP_History has). Returns (targets, usable_cases)."""
    T = {}
    for key, kind in MATCH_STREAMS.items():
        df = data[key]
        hist = df[HISTORY_COL].dropna()
        hist = hist[~hist.index.duplicated()]
        Y = np.array([to_t(eval_case(df[c], hist.index, key.endswith("_cum"), t0), kind)
                      for c in cases])
        T[key] = dict(kind=kind, dates=hist.index, obs=to_t(hist.values, kind), Y=Y)
    for key, t in T.items():                    # keep dates nearly every case covers
        keep = np.isfinite(t["Y"]).mean(axis=0) >= 0.95
        t["Y"], t["obs"], t["dates"] = t["Y"][:, keep], t["obs"][keep], t["dates"][keep]
    ok = np.ones(len(cases), bool)
    for t in T.values():
        ok &= np.isfinite(t["Y"]).all(axis=1)   # drop cases with a hole in a kept date
    for t in T.values():
        t["Y"] = t["Y"][ok]
        t["scale"] = max(float(np.abs(t["obs"]).max()), 1e-9)   # same idea as the old
    dropped = [c for c, o in zip(cases, ok) if not o]            # "divide by max history"
    print("scored streams: " + ", ".join(
        f"{k} ({MATCH_STREAMS[k]}, {len(t['obs'])} history dates)" for k, t in T.items()))
    if dropped:
        print(f"   {len(dropped)} cases dropped (curve has holes on the history dates): "
              f"{dropped[:8]}{' ...' if len(dropped) > 8 else ''}")
    return T, [c for c, o in zip(cases, ok) if o]


def stream_rmse(pred, t):
    """Normalised RMSE of predicted curves (N x G) against history, per row."""
    return np.sqrt(np.mean((pred - t["obs"][None, :]) ** 2, axis=1)) / t["scale"]


def total_misfit(preds, T):
    """Sum over the scored streams (oil cum + water cum + water rate). Lower = better."""
    return sum(stream_rmse(preds[k], T[k]) for k in T)


class CurveProxy:
    """The proxy: 4 settings in -> the scored curves out.

    Per stream: compress the curves with PCA (a few shape numbers per case), then
    one Gaussian process per shape number. A GP is a good choice for ~70 runs
    because it works with few samples and says how unsure it is.
    """

    def __init__(self, param_cols, bounds):
        self.cols, self.bounds = list(param_cols), bounds

    def to_unit(self, P):
        """Physical settings -> [0,1] box (log scale for LOG_PARAMS)."""
        P = np.array(P, float, ndmin=2)
        U = np.empty_like(P)
        for j, c in enumerate(self.cols):
            lo, hi = self.bounds[c]
            if c in LOG_PARAMS:
                U[:, j] = (np.log(P[:, j]) - np.log(lo)) / (np.log(hi) - np.log(lo))
            else:
                U[:, j] = (P[:, j] - lo) / (hi - lo)
        return U

    def from_unit(self, U):
        U = np.array(U, float, ndmin=2)
        P = np.empty_like(U)
        for j, c in enumerate(self.cols):
            lo, hi = self.bounds[c]
            if c in LOG_PARAMS:
                P[:, j] = np.exp(np.log(lo) + U[:, j] * (np.log(hi) - np.log(lo)))
            else:
                P[:, j] = lo + U[:, j] * (hi - lo)
        return P

    def fit(self, P, Y):
        U = self.to_unit(P)
        self.models = {}
        with warnings.catch_warnings():
            if FIT_PROXY_QUIET:
                warnings.simplefilter("ignore")
            for key, y in Y.items():
                mu = y.mean(axis=0)
                npc = int(min(PCA_MAX, y.shape[0] - 1, y.shape[1]))
                pca = PCA(n_components=npc).fit(y - mu)
                cum = np.nan_to_num(np.cumsum(pca.explained_variance_ratio_))
                k = int(min(npc, np.searchsorted(cum, PCA_VAR) + 1))
                Z = pca.transform(y - mu)[:, :k]
                gps = []
                for j in range(k):
                    kern = (ConstantKernel(1.0, (1e-2, 1e2))
                            * Matern(np.ones(U.shape[1]), (0.05, 50.0), nu=2.5)
                            + WhiteKernel(1e-4, (1e-8, 1e-1)))
                    gps.append(GaussianProcessRegressor(
                        kern, normalize_y=True, n_restarts_optimizer=2,
                        random_state=RANDOM_SEED).fit(U, Z[:, j]))
                self.models[key] = dict(mu=mu, comp=pca.components_[:k], gps=gps)
        return self

    def predict(self, P):
        U = self.to_unit(P)
        out = {}
        for key, m in self.models.items():
            Z = np.column_stack([gp.predict(U) for gp in m["gps"]])
            out[key] = m["mu"] + Z @ m["comp"]
        return out


def cv_predict(P, Y, param_cols, bounds, k=5):
    """Out-of-fold predictions: every case is predicted by a proxy that never saw it."""
    n = len(P)
    folds = np.arange(n) % k
    np.random.RandomState(RANDOM_SEED).shuffle(folds)
    out = {key: np.zeros_like(y) for key, y in Y.items()}
    for f in range(k):
        tr, te = folds != f, folds == f
        pr = CurveProxy(param_cols, bounds).fit(P[tr], {key: y[tr] for key, y in Y.items()})
        pred = pr.predict(P[te])
        for key in Y:
            out[key][te] = pred[key]
    return out


def proxy_report(pred, Y, T, label):
    """Accuracy of predicted curves vs the real simulated curves for these cases."""
    rows = []
    for key in T:
        err = np.sqrt(np.mean((pred[key] - Y[key]) ** 2, axis=1)) / T[key]["scale"]
        last_true, last_pred = Y[key][:, -1], pred[key][:, -1]
        sst = np.sum((last_true - last_true.mean()) ** 2)
        r2 = 1 - np.sum((last_true - last_pred) ** 2) / sst if sst > 0 else np.nan
        rows.append({"set": label, "stream": key, "n_cases": len(err),
                     "mean_err_pct_of_history_max": 100 * err.mean(),
                     "worst_err_pct": 100 * err.max(), "R2_last_history_date": r2})
    return pd.DataFrame(rows)


def run_proxy_history_match(data, params):
    print("\n=== PROXY + CONTINUOUS HISTORY MATCH (V4) ===")
    pcols = list(params.columns)
    first_key = next(iter(MATCH_STREAMS))
    cases = [c for c in case_cols(data[first_key])
             if c in params.index and params.loc[c].notna().all()]
    t0 = SIM_START or min(data[k][HISTORY_COL].dropna().index.min() for k in MATCH_STREAMS)
    t0 = pd.Timestamp(t0).value
    T, cases = build_targets(data, cases, t0)
    P = params.loc[cases, pcols].values.astype(float)
    split = np.array([split_of(c) for c in cases])
    bounds = {c: (params[c].min(), params[c].max()) for c in pcols}   # tested ranges
    Yall = {k: t["Y"] for k, t in T.items()}
    print("cases usable: " + ", ".join(f"{s}={int((split == s).sum())}"
          for s in ["train", "validation", "test", "unassigned"]))

    tr = split == "train"
    if tr.sum() < 15:
        print("   WARNING: fewer than 15 train cases under TRAIN_IDS - treating ALL "
              "cases as train and relying on cross-validation. Check TRAIN_IDS/VAL_IDS/TEST_IDS.")
        tr = np.ones(len(cases), bool)
    sub = lambda m: {k: y[m] for k, y in Yall.items()}

    # ---- 1) how good is the proxy? (honest numbers) -------------------------
    cvp = cv_predict(P[tr], sub(tr), pcols, bounds)
    reports = [proxy_report(cvp, sub(tr), T, "train (5-fold CV)")]
    proxy = CurveProxy(pcols, bounds).fit(P[tr], sub(tr))
    held = {}
    for name in ["validation", "test"]:
        m = split == name
        if m.any():
            held[name] = (m, proxy.predict(P[m]))
            reports.append(proxy_report(held[name][1], sub(m), T, name))
    rep = pd.concat(reports, ignore_index=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    rep.to_csv(OUTPUT_DIR / "proxy_validation.csv", index=False)
    print("\nProxy accuracy (error as % of the largest history value; lower = better):")
    print(rep.round(3).to_string(index=False))

    # Error of the proxy IN THE UNITS OF THE MISFIT: how wrong is its score for a case?
    ref_m, ref_pred = held["validation"] if "validation" in held else (tr, cvp)
    ref_true = {k: y[ref_m] for k, y in Yall.items()} if "validation" in held else sub(tr)
    mis_true = total_misfit(ref_true, T)
    mis_pred = total_misfit(ref_pred, T)
    proxy_err = float(np.mean(np.abs(mis_true - mis_pred)))
    print(f"\nproxy's own error in misfit units: {proxy_err:.4f} "
          f"(from {'validation cases' if 'validation' in held else 'cross-validation'})")

    # ---- 2) final proxy for the search ---------------------------------------
    use = tr | (split == "validation") if FINAL_REFIT_ON_VAL else tr
    if FINAL_REFIT_ON_VAL and (split == "validation").any():
        proxy = CurveProxy(pcols, bounds).fit(P[use], sub(use))
        print(f"final proxy refit on train+validation ({int(use.sum())} cases); test untouched")

    # ---- 3) search ALL settings, not just the existing runs ------------------
    d = len(pcols)
    U = qmc.Sobol(d=d, scramble=True, seed=RANDOM_SEED).random_base2(N_SOBOL_POW2)
    samples = proxy.from_unit(U)
    mis = total_misfit(proxy.predict(samples), T)

    def objective(x):
        x = np.asarray(x, float)
        one = x.ndim == 1
        m = total_misfit(proxy.predict(proxy.from_unit(x.reshape(d, -1).T)), T)
        return float(m[0]) if one else m

    res = differential_evolution(objective, [(0.0, 1.0)] * d, seed=RANDOM_SEED, maxiter=150,
                                 popsize=20, tol=1e-8, vectorized=True,
                                 updating="deferred", polish=True)
    best_u = np.clip(res.x, 0, 1)
    best_p = proxy.from_unit(best_u)[0]
    best_m = float(objective(best_u))
    if mis.min() < best_m:                           # the random sweep beat the optimiser
        best_p, best_m = samples[int(mis.argmin())], float(mis.min())
    best_pred = proxy.predict(best_p[None, :])

    # ---- 4) every setting that fits "about as well as the best" --------------
    margin = max(ACCEPT_MARGIN_MIN, ACCEPT_MARGIN_REL * best_m, 2 * proxy_err)
    acc = mis <= best_m + margin
    A = samples[acc]
    print(f"\nBEST settings (misfit {best_m:.4f}):")
    for c, v in zip(pcols, best_p):
        print(f"   {c:26s} {v:.4g}   (tested {bounds[c][0]:.3g} to {bounds[c][1]:.3g})")
    print(f"\nACCEPTABLE settings: {int(acc.sum())} of {len(mis)} trial sets are within "
          f"{margin:.4f} of the best (largest of: {ACCEPT_MARGIN_MIN} fixed, "
          f"{ACCEPT_MARGIN_REL:.0%} of best, 2 x proxy error)")

    # existing cases scored with the SAME metric, for comparison
    ex = pd.Series(total_misfit(Yall, T), index=cases).sort_values()
    print(f"best EXISTING case under this metric: {ex.index[0]} "
          f"(misfit {ex.iloc[0]:.4f}, {split_of(ex.index[0])}); "
          f"proxy optimum is {'BETTER' if best_m < ex.iloc[0] else 'not better'} by "
          f"{ex.iloc[0] - best_m:+.4f}")

    # which settings does history actually pin down?
    Ua = proxy.to_unit(A)
    ident = []
    for j, c in enumerate(pcols):
        w = float(Ua[:, j].max() - Ua[:, j].min())
        verdict = "PINNED DOWN" if w < 0.5 else ("partly constrained" if w < 0.8 else "NOT constrained")
        ident.append({"setting": c, "best": best_p[j], "acceptable_min": A[:, j].min(),
                      "acceptable_max": A[:, j].max(), "tested_min": bounds[c][0],
                      "tested_max": bounds[c][1], "share_of_range_acceptable": w,
                      "verdict": verdict})
    ident = pd.DataFrame(ident)
    ident.to_csv(OUTPUT_DIR / "best_parameters.csv", index=False)
    pd.DataFrame(A, columns=pcols).assign(misfit=mis[acc]).sort_values("misfit") \
        .to_csv(OUTPUT_DIR / "acceptable_parameter_sets.csv", index=False)
    print("\nWhat the history pins down (range of acceptable sets):")
    for _, r in ident.iterrows():
        print(f"  {r['verdict']:18s} {r['setting']}: {r['acceptable_min']:.3g} to "
              f"{r['acceptable_max']:.3g} (tested {r['tested_min']:.3g} to {r['tested_max']:.3g})")

    # ---- 5) is the match actually good? --------------------------------------
    print("\nResiduals at the best settings (proxy curve minus history, as % of history max):")
    flag = False
    for k, t in T.items():
        resid = (best_pred[k][0] - t["obs"]) / t["scale"]
        rm, bias = float(np.sqrt(np.mean(resid ** 2))), float(resid.mean())
        note = ""
        if rm > 0.10:
            note, flag = "  <- large", True
        elif abs(bias) > 0.5 * rm and rm > 0.03:
            note, flag = "  <- one-sided (systematic)", True
        print(f"   {k:11s} rmse {100 * rm:5.1f}%   mean bias {100 * bias:+5.1f}%{note}")
    if flag:
        print("   CHECK: the proxy cannot reproduce history in at least one stream. This can "
              "mean the 4 settings do not span the truth (model inadequacy). Do not read the "
              "acceptable ranges as 'the truth is in here'.")

    # ---- 6) plots ------------------------------------------------------------
    plot_proxy_quality(T, Yall, cvp, held, tr, sub, split)
    plot_history_match(T, Yall, proxy, best_pred, A)
    plot_posterior(ident, A, best_p, pcols)
    return dict(proxy=proxy, best_params=best_p, acceptable=A, acc_mis=mis[acc], targets=T,
                proxy_err=proxy_err, t0=t0, pcols=pcols, bounds=bounds,
                samples=samples, mis=mis)


def plot_proxy_quality(T, Yall, cvp, held, tr, sub, split):
    """Plot 9: predicted vs real simulated value at the last history date."""
    fig, axes = plt.subplots(1, len(T), figsize=(5.2 * len(T), 4.8), constrained_layout=True)
    for ax, (key, t) in zip(np.atleast_1d(axes), T.items()):
        kind = t["kind"]
        pairs = [("train (CV)", sub(tr)[key][:, -1], cvp[key][:, -1], "#999999")]
        for name, col in [("validation", "#1f77b4"), ("test", "#ff7f0e")]:
            if name in held:
                m, pred = held[name]
                pairs.append((name, Yall[key][m][:, -1], pred[key][:, -1], col))
        allv = np.concatenate([from_t(a, kind) for _, a, _, _ in pairs])
        for name, tru, prd, col in pairs:
            ax.scatter(from_t(tru, kind), from_t(prd, kind), s=22, alpha=0.8, color=col, label=name)
        lim = [allv.min(), allv.max()]
        ax.plot(lim, lim, "k--", lw=1)
        ax.set_xlabel("real simulator value"); ax.set_ylabel("proxy prediction")
        ax.set_title(f"{NICE[key]} at last history date", fontsize=10)
        ax.grid(alpha=0.3)
    np.atleast_1d(axes)[0].legend(fontsize=8)
    fig.suptitle("Plot 9 - Proxy accuracy: points on the dashed line = perfect "
                 "(validation/test never used for fitting)")
    save(fig, "09_proxy_accuracy.png")


def plot_history_match(T, Yall, proxy, best_pred, A):
    """Plot 10: the history match. Black = real field, red = proxy at the best
    settings, band = all acceptable settings, grey = the existing simulations."""
    rng = np.random.RandomState(RANDOM_SEED)
    sel = A[rng.choice(len(A), min(len(A), 300), replace=False)]
    ens = proxy.predict(sel)
    fig, axes = plt.subplots(1, len(T), figsize=(5.8 * len(T), 4.8), constrained_layout=True)
    for ax, (key, t) in zip(np.atleast_1d(axes), T.items()):
        kind, x = t["kind"], t["dates"]
        for y in t["Y"]:
            ax.plot(x, from_t(y, kind), color="lightgrey", lw=0.5)
        lo, hi = np.percentile(ens[key], [10, 90], axis=0)
        ax.fill_between(x, from_t(lo, kind), from_t(hi, kind), color="#d62728", alpha=0.25,
                        label="acceptable sets (P10-P90)")
        ax.plot(x, from_t(best_pred[key][0], kind), color="#d62728", lw=2, label="best match")
        ax.plot(x, from_t(t["obs"], kind), "ko", ms=3, label="History")
        ax.set_title(NICE[key], fontsize=10)
        ax.tick_params(axis="x", rotation=30)
        ax.grid(alpha=0.3)
    np.atleast_1d(axes)[0].legend(fontsize=8, loc="upper left")
    fig.suptitle("Plot 10 - History match from the proxy search")
    save(fig, "10_history_match.png")


def plot_posterior(ident, A, best_p, pcols):
    """Plot 11: where the acceptable settings sit inside the tested range."""
    fig, axes = plt.subplots(1, len(pcols), figsize=(4.2 * len(pcols), 3.6), constrained_layout=True)
    col = {"PINNED DOWN": "#1a7f37", "partly constrained": "#b8860b", "NOT constrained": "#b22222"}
    for ax, (j, c) in zip(np.atleast_1d(axes), enumerate(pcols)):
        r = ident.iloc[j]
        ax.hist(A[:, j], bins=25, range=(r["tested_min"], r["tested_max"]),
                color=col[r["verdict"]], alpha=0.8, edgecolor="k", lw=0.3)
        ax.axvline(best_p[j], color="k", lw=2, label="best")
        ax.set_title(f"{c}\n[{r['verdict']}]", fontsize=9, color=col[r["verdict"]], weight="bold")
        ax.set_yticks([])
    np.atleast_1d(axes)[0].legend(fontsize=8)
    fig.suptitle("Plot 11 - Acceptable settings across the tested range "
                 "(flat across the range = history cannot tell)")
    save(fig, "11_acceptable_settings.png")


# =============================================================================
# STEP 6 (NEW in V6) - FORECAST 2008 -> 2028, CHOOSING CASES, SUBMISSION TEMPLATE
# =============================================================================
# The history match above is validated INSIDE the data (1998-2008). Everything
# below is an EXTRAPOLATION beyond 2008, where the simulator curves and the real
# field have no data. Its only honest test is the HINDCAST (plot 12).
#   1. fit a decline curve to the end of EVERY case's history
#        oil, gas : Arps hyperbolic decline of the RATE
#        water    : water-oil ratio (WOR) trend against cumulative oil, capped
#   2. project each case to 2028 -> 100 pseudo-forecasts, one per case
#   3. teach a second proxy (same CurveProxy class) to forecast for ANY settings
#   4. run it on every ACCEPTABLE setting from the history match -> a forecast band
#   5. choose N_SELECT cases and fill in the organisers' template
STREAMS = ["gas", "oil", "water"]                    # order of the template columns


def _num(name):
    nums = re.findall(r"\d+", str(name))
    return int(nums[-1]) if nums else -1


def at(series, dates, key, t0):
    """Value of one case's curve at `dates` (cumulative: linear from 0 at t0;
    rate: nearest reported value). Same rule as the history match."""
    return eval_case(series, pd.DatetimeIndex(dates), key.endswith("_cum"), t0)


def arps(t, qi, Di, b):
    """Hyperbolic decline: the rate falls fast at first, then more slowly."""
    return qi / np.power(1.0 + b * Di * t, 1.0 / b)


def fit_arps(t, q):
    """Fit (qi, Di, b) to a rate history; t = years from the start of the window."""
    m = np.isfinite(t) & np.isfinite(q) & (q > 0)
    if m.sum() < 5:
        return None
    t, q = t[m], q[m]
    lq = np.log(q)

    def f(tt, lqi, lDi, b):
        return np.log(arps(tt, np.exp(lqi), np.exp(lDi), b))

    best = None
    for b0 in (0.1, 0.5, 0.9):
        try:
            popt, _ = curve_fit(f, t, lq, p0=[lq[0], np.log(0.15), b0],
                                bounds=([lq.min() - 3, -9, 0.01], [lq.max() + 3, 2.0, 1.0]),
                                maxfev=4000)
            sse = np.sum((f(t, *popt) - lq) ** 2)
            if best is None or sse < best[0]:
                best = (sse, popt)
        except Exception:
            continue
    if best is None:                                   # fallback: plain exponential
        slope, icpt = np.polyfit(t, lq, 1)
        return (float(np.exp(icpt)), float(max(-slope, 1e-6)), 0.01)
    return (float(np.exp(best[1][0])), float(np.exp(best[1][1])), float(best[1][2]))


def fit_shapes(data, case, cut, window_years, t0):
    """What the forecast needs from one case, using history up to `cut` only:
    its state at `cut` and the decline shapes fitted on the last `window_years`."""
    w0 = cut - pd.Timedelta(days=int(365.25 * window_years))
    out = {"w0": w0, "cut": cut}
    for s in ("oil", "gas"):
        ser = data[f"{s}_rate"][case].dropna()
        ser = ser[(ser.index >= w0) & (ser.index <= cut)]
        out[s] = fit_arps((ser.index - w0).days.values / 365.25, ser.values)
        if out[s] is None:
            return None
    st = {k: float(at(data[k][case], [cut], k, t0)[0])
          for k in ["oil_cum", "gas_cum", "water_cum", "oil_rate", "gas_rate", "water_rate"]}
    if not np.all(np.isfinite(list(st.values()))):
        return None
    out["state"] = st
    grid = pd.date_range(w0, cut, periods=24)
    qo = at(data["oil_rate"][case], grid, "oil_rate", t0)
    qw = at(data["water_rate"][case], grid, "water_rate", t0)
    npv = at(data["oil_cum"][case], grid, "oil_cum", t0)
    with np.errstate(divide="ignore", invalid="ignore"):
        wor = np.where(qo > 0, qw / qo, np.nan)
    ok = np.isfinite(wor) & np.isfinite(npv) & (wor > 1e-4)
    bw = 0.0
    if ok.sum() >= 5:
        slope, _ = np.polyfit(npv[ok] / max(st["oil_cum"], 1.0), np.log(wor[ok]), 1)
        bw = max(float(slope), 0.0)                    # water cut does not fall with production
    out["wor_slope"] = bw
    return out


def project(shp, out_dates):
    """Cumulative gas / oil / water at `out_dates` (all on/after shp['cut'])."""
    cut, st = shp["cut"], shp["state"]
    out_dates = pd.DatetimeIndex(out_dates)
    d = np.arange(0, (out_dates.max() - cut).days + 8, 7.0)      # weekly steps
    off = (cut - shp["w0"]).days / 365.25
    ty, dd = off + d / 365.25, np.diff(d)

    def integrate(q):
        return np.concatenate([[0.0], np.cumsum(0.5 * (q[1:] + q[:-1]) * dd)])

    q = {}
    for s in ("oil", "gas"):                            # rescaled to join the last real rate
        qi, Di, b = shp[s]
        q[s] = arps(ty, qi, Di, b) * (st[f"{s}_rate"] / max(arps(off, qi, Di, b), 1e-12))
    cum_o = st["oil_cum"] + integrate(q["oil"])
    wor_end = max(st["water_rate"] / max(st["oil_rate"], 1e-9), 1e-4)
    wor = np.minimum(wor_end * np.exp(shp["wor_slope"] * (cum_o - st["oil_cum"])
                                      / max(st["oil_cum"], 1.0)), max(WOR_MAX, wor_end))
    x = (out_dates - cut).days.values.astype(float)
    return {"gas": np.interp(x, d, st["gas_cum"] + integrate(q["gas"])),
            "oil": np.interp(x, d, cum_o),
            "water": np.interp(x, d, st["water_cum"] + integrate(wor * q["oil"]))}


def hindcast(data, cases, window_years, hist_end, t0):
    """Pretend history ended HINDCAST_YEARS earlier, forecast, and compare the
    predicted production INCREASE with what the simulator really did."""
    cut = hist_end - pd.Timedelta(days=int(365.25 * HINDCAST_YEARS))
    rows = {}
    for c in cases:
        shp = fit_shapes(data, c, cut, window_years, t0)
        if shp is None:
            continue
        pred, r = project(shp, [hist_end]), {}
        for s in STREAMS:
            true_end = float(at(data[f"{s}_cum"][c], [hist_end], f"{s}_cum", t0)[0])
            inc = true_end - shp["state"][f"{s}_cum"]
            r[s] = (pred[s][0] - shp["state"][f"{s}_cum"]) / inc - 1 if inc > 0 else np.nan
        rows[c] = r
    return pd.DataFrame(rows).T


def select_cases(U, oil_end, n, mode):
    """Indices into the acceptable sets (sorted best-misfit first)."""
    n = min(n, len(U))
    if mode == "best":
        return list(range(n))
    if mode == "span_forecast":                         # spread of the final oil forecast
        order = np.argsort(oil_end)
        return sorted(order[np.unique(np.round(np.linspace(0.03, 0.97, n)
                                               * (len(order) - 1)).astype(int))].tolist())
    chosen = [0]                                        # best match, then most different
    while len(chosen) < n:
        dist = np.min(np.linalg.norm(U[:, None, :] - U[chosen][None, :, :], axis=2), axis=1)
        dist[chosen] = -1
        chosen.append(int(np.argmax(dist)))
    return chosen


# ---- the submission template ------------------------------------------------
def find_template():
    path = find_file(FILES["template"])
    if path is None:
        for p in sorted(DATA_DIR.glob("*.xlsx")):
            if "template" in p.stem.lower() and not p.name.startswith("~$"):
                return p
    return path


def read_template_layout(path):
    """In every 'Case N' sheet find: header row, date column, the gas / oil / water
    cumulative columns and the (row, date) of each data row. Nothing is hard-coded."""
    wb = load_workbook(path)
    layout = {}
    for name in wb.sheetnames:
        if not re.search(r"case\s*_?\s*\d+", name, re.I):
            continue
        ws, hdr = wb[name], None
        for r in range(1, min(ws.max_row, 15) + 1):
            cells = [str(c.value).lower() if c.value is not None else "" for c in ws[r]]
            if sum("cumulative" in x for x in cells) >= 2:
                hdr = r
                break
        if hdr is None:
            continue
        cols = {}
        for c in ws[hdr]:
            v = str(c.value).lower() if c.value is not None else ""
            if "cumulative" in v:
                for s in STREAMS:
                    if s in v:
                        cols[s] = c.column
                        break
        counts = {col: sum(hasattr(ws.cell(row=r, column=col).value, "year")
                           for r in range(hdr + 1, ws.max_row + 1))
                  for col in range(1, ws.max_column + 1)}
        date_col = max(counts, key=counts.get)
        rows = [(r, pd.Timestamp(ws.cell(row=r, column=date_col).value))
                for r in range(hdr + 1, ws.max_row + 1)
                if hasattr(ws.cell(row=r, column=date_col).value, "year")]
        layout[name] = {"hdr": hdr, "cols": cols, "date_col": date_col, "rows": rows}
    return wb, dict(sorted(layout.items(), key=lambda kv: _num(kv[0])))


def make_own_template(dates):
    """Fallback if no template file exists: same layout, our own dates."""
    path = OUTPUT_DIR / "own_template.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    for k in range(1, N_SELECT + 1):
        ws = wb.create_sheet(f"Case {k}")
        ws.append(["Date", "Gas production cumulative [MSCF]",
                   "Oil production cumulative [STB]", "Water production cumulative [STB]"])
        for d in dates:
            ws.append([d.to_pydatetime()])
    wb.save(path)
    return path


def check_deliverable(path, obs_end):
    """Read the FILLED template back from disk and verify the deliverable's rules."""
    wb, layout = read_template_layout(path)
    print("\n=== DELIVERABLE CHECK (reading the saved file back) ===")
    all_ok = True
    for name, lay in layout.items():
        ws, problems = wb[name], []
        for s in STREAMS:
            vals = np.array([ws.cell(row=r, column=lay["cols"][s]).value
                             if ws.cell(row=r, column=lay["cols"][s]).value is not None else np.nan
                             for r, _ in lay["rows"]], float)
            if np.isnan(vals).any():
                problems.append(f"{s}: blank cells")
                continue
            if (vals < 0).any():
                problems.append(f"{s}: negative values")
            if (np.diff(vals) < -1e-6 * max(abs(vals).max(), 1)).any():
                problems.append(f"{s}: cumulative decreases")
            if np.isfinite(obs_end.get(s, np.nan)) and vals[0] < obs_end[s] * (1 - 1e-9):
                problems.append(f"{s}: starts below the 2008 history value")
        print(f"  {name:8s} rows={len(lay['rows'])}  "
              + ("PASS" if not problems else "FAIL: " + "; ".join(problems)))
        all_ok &= not problems
    print("  ALL SHEETS PASS" if all_ok else "  FIX THE SHEETS MARKED FAIL BEFORE SUBMITTING")


# ---- plots 12 and 13 --------------------------------------------------------
def plot_hindcast(errs, chosen):
    """Plot 12: how good is the decline-curve forecast, tested on the past?"""
    fig, (a, b) = plt.subplots(1, 2, figsize=(16, 5.4), constrained_layout=True)
    xs = np.arange(len(STREAMS))
    for k, wy in enumerate(errs):
        med = [errs[wy][s].abs().median() * 100 for s in STREAMS]
        a.bar(xs + k * 0.25, med, 0.25, edgecolor="k", alpha=1.0 if wy == chosen else 0.55,
              label=f"fit on last {wy:g} yr" + ("  <- chosen" if wy == chosen else ""))
    a.set_xticks(xs + 0.25)
    a.set_xticklabels([s.capitalize() for s in STREAMS])
    a.set_ylabel(f"typical error of the {HINDCAST_YEARS:g}-year production increase [%]")
    a.set_title("Which fitting window forecasts best?  (lower is better)", weight="bold")
    a.legend()
    a.grid(alpha=0.3, axis="y")
    b.boxplot([errs[chosen][s].dropna().values * 100 for s in STREAMS], showfliers=False)
    b.set_xticklabels([s.capitalize() for s in STREAMS])
    b.axhline(0, color="k", lw=1)
    b.set_ylabel("forecast error [%]   (above 0 = over-predicted)")
    b.set_title(f"Chosen window ({chosen:g} yr): spread of errors across cases", weight="bold")
    b.grid(alpha=0.3, axis="y")
    fig.suptitle(f"Plot 12 - Hindcast: history cut {HINDCAST_YEARS:g} years early, forecast, compare "
                 "with what the simulator really did.\nA 20-year forecast will be LESS accurate "
                 "than this.", fontsize=12, weight="bold")
    save(fig, "12_hindcast_check.png")


def plot_forecast(data, fdates, bands, selected, hist_end, obs_end):
    """Plot 13: the 2008-2028 forecast, cumulative (top) and implied rate (bottom)."""
    fig, axes = plt.subplots(2, 3, figsize=(19, 10), constrained_layout=True, sharex=True)
    cmap = plt.get_cmap("viridis")
    days = np.diff(np.concatenate([[hist_end.value],
                                   fdates.values.astype("datetime64[ns]").astype("int64")])) / 8.64e13
    for j, s in enumerate(STREAMS):
        ax, ax2 = axes[0, j], axes[1, j]
        lo, mid, hi = bands[s]
        h = data[f"{s}_cum"][HISTORY_COL].dropna()
        ax.plot(h.index, h.values, "ko", ms=3, label="real history")
        ax.fill_between(fdates, lo, hi, color="#1f77b4", alpha=0.3, label="acceptable settings P10-P90")
        ax.plot(fdates, mid, color="#1f77b4", lw=2, label="P50")
        hr = data[f"{s}_rate"][HISTORY_COL].dropna()
        ax2.plot(hr.index, hr.values, "ko", ms=3, label="real history")
        for arr, ls, lw in [(lo, "--", 1), (mid, "-", 2), (hi, "--", 1)]:
            prev = np.concatenate([[obs_end[s]], arr[:-1]])
            ax2.plot(fdates, (arr - prev) / days, color="#1f77b4", ls=ls, lw=lw)
        for rank, c in enumerate(selected[s]):
            col = cmap(rank / max(len(selected[s]) - 1, 1))
            ax.plot(fdates, c, color=col, lw=1.2, alpha=0.9)
            prev = np.concatenate([[obs_end[s]], c[:-1]])
            ax2.plot(fdates, (c - prev) / days, color=col, lw=1.0, alpha=0.8)
        for a_ in (ax, ax2):
            a_.axvline(hist_end, color="k", ls=":")
            a_.grid(alpha=0.3)
        ax.set_title(f"{s.capitalize()} cumulative", weight="bold")
        ax2.set_title(f"{s.capitalize()} rate (implied by the cumulative)", weight="bold")
    axes[0, 0].legend(fontsize=9, loc="upper left")
    fig.suptitle("Plot 13 - Forecast 2008 to 2028.  Coloured lines = the cases submitted (dark = best "
                 "history match).  Right of the dotted line is EXTRAPOLATION: the band shows only the "
                 "spread between equally good matches, not all the uncertainty.",
                 fontsize=12, weight="bold")
    save(fig, "13_forecast.png")


def run_forecast_and_deliverable(data, params, hm):
    print("\n=== FORECAST 2008 -> 2028 + SUBMISSION TEMPLATE (NEW in V6) ===")
    proxy, t0, pcols, bounds = hm["proxy"], hm["t0"], hm["pcols"], hm["bounds"]
    hist_end = data["oil_cum"][HISTORY_COL].dropna().index.max()
    cases = [c for c in case_cols(data["oil_cum"]) if c in params.index and params.loc[c].notna().all()]
    print(f"history ends {hist_end.date()}; the simulator curves stop here too, so everything "
          "after is extrapolation")

    # ---- the template defines the dates to forecast ---------------------------
    tpath = find_template()
    layout = None
    if tpath is not None:
        wb, layout = read_template_layout(tpath)
        first = next(iter(layout.values()))
        fdates = pd.DatetimeIndex([d for _, d in first["rows"]])
        print(f"template: {tpath.name} ({len(layout)} case sheets), "
              f"{len(fdates)} dates {fdates.min().date()} -> {fdates.max().date()}")
        if (fdates <= hist_end).any():
            print("  WARNING: template has dates on/before the history end; those rows are skipped.")
            fdates = fdates[fdates > hist_end]
    else:
        print("no template file found -> will write our own 10-sheet workbook")
        fdates = pd.date_range(FORECAST_START_IF_NO_TEMPLATE, FORECAST_END_IF_NO_TEMPLATE, freq="QS")

    # ---- 1) hindcast: test the forecast method on the past, choose the window ------
    errs = {}
    for wy in FIT_WINDOW_CANDIDATES:
        errs[wy] = hindcast(data, cases, wy, hist_end, t0)
        print(f"  hindcast, fit on last {wy:g} yr -> typical error of the {HINDCAST_YEARS:g}-yr increase: "
              + "  ".join(f"{s} {errs[wy][s].abs().median():.1%}" for s in STREAMS))
    chosen = FIT_WINDOW_YEARS or min(
        errs, key=lambda w: np.nanmean([errs[w][s].abs().median() for s in STREAMS]))
    if chosen not in errs:
        errs[chosen] = hindcast(data, cases, chosen, hist_end, t0)
    print(f"  chosen fitting window: last {chosen:g} years")
    pd.concat({f"{w:g}yr": e for w, e in errs.items()}, axis=1).to_csv(OUTPUT_DIR / "hindcast_errors.csv")
    plot_hindcast(errs, chosen)

    # ---- 2) one pseudo-forecast per case, then teach a proxy to forecast ----------
    fd_ext = pd.DatetimeIndex([hist_end]).append(fdates)    # column 0 = state at history end
    F, ok_cases = {s: [] for s in STREAMS}, []
    for c in cases:
        shp = fit_shapes(data, c, hist_end, chosen, t0)
        if shp is None:
            continue
        pr = project(shp, fd_ext)
        ok_cases.append(c)
        for s in STREAMS:
            F[s].append(pr[s])
    print(f"  cases with a usable forecast fit: {len(ok_cases)} of {len(cases)}")
    P = params.loc[ok_cases, pcols].values.astype(float)
    split = np.array([split_of(c) for c in ok_cases])
    use = np.ones(len(ok_cases), bool) if FORECAST_REFIT_ON_ALL else np.isin(split, ["train", "validation"])
    if use.sum() < 15:
        print("  WARNING: fewer than 15 train/validation cases - using ALL cases for the forecast proxy.")
        use = np.ones(len(ok_cases), bool)
    kinds = {f"fc_{s}": ("log" if s == "water" else "linear") for s in STREAMS}
    Yf = {f"fc_{s}": to_t(np.array(F[s]), kinds[f"fc_{s}"]) for s in STREAMS}
    cvp = cv_predict(P[use], {k: y[use] for k, y in Yf.items()}, pcols, bounds)
    for s in STREAMS:
        k = f"fc_{s}"
        tru, prd = from_t(Yf[k][use], kinds[k]), from_t(cvp[k], kinds[k])
        err = np.sqrt(np.mean((tru - prd) ** 2, axis=1)) / (np.abs(tru).max(axis=1) + 1e-12)
        print(f"  forecast proxy {s:5s}: reproduces the per-case extrapolations to "
              f"{np.median(err):.2%} (5-fold CV, % of each curve's peak)")
    print("  (this only measures the PROXY'S fidelity to the decline-curve forecasts, not whether "
          "the decline curves are right - that is the hindcast above)")
    fproxy = CurveProxy(pcols, bounds).fit(P[use], {k: y[use] for k, y in Yf.items()})

    # ---- 3) forecast for every ACCEPTABLE setting from the history match -------------
    A, am = hm["acceptable"], hm["acc_mis"]
    if len(A) < max(MIN_FORECAST_SETS, N_SELECT):
        k = max(MIN_FORECAST_SETS, N_SELECT * 5)
        idx = np.argsort(hm["mis"])[:k]
        print(f"  only {len(A)} acceptable sets; using the {k} best-fitting trial sets instead "
              "for the forecast band")
        A, am = hm["samples"][idx], hm["mis"][idx]
    if len(A) > MAX_FORECAST_SETS:
        idx = np.random.RandomState(RANDOM_SEED).choice(len(A), MAX_FORECAST_SETS, replace=False)
        A, am = A[idx], am[idx]
    order = np.argsort(am)
    A, am = A[order], am[order]
    pred = fproxy.predict(A)
    obs_end = {s: last_valid(data[f"{s}_cum"][HISTORY_COL]) for s in STREAMS}
    ens = {}
    for s in STREAMS:
        arr = from_t(pred[f"fc_{s}"], kinds[f"fc_{s}"])
        fc = arr[:, 1:]
        if ANCHOR_TO_HISTORY and np.isfinite(obs_end[s]):
            fc = fc + (obs_end[s] - arr[:, 0])[:, None]    # join the real history exactly
        fc = np.maximum.accumulate(np.maximum(fc, 0), axis=1)
        if np.isfinite(obs_end[s]):
            fc = np.maximum(fc, obs_end[s])
        ens[s] = fc
    bands = {s: tuple(np.percentile(ens[s], q, axis=0) for q in (10, 50, 90)) for s in STREAMS}
    if ANCHOR_TO_HISTORY:
        shift = {s: np.median(obs_end[s] - from_t(pred[f"fc_{s}"], kinds[f"fc_{s}"])[:, 0]) / obs_end[s]
                 for s in STREAMS if np.isfinite(obs_end[s]) and obs_end[s] > 0}
        print("  anchoring to the real history shifted the forecasts by (median, % of 2008 cumulative): "
              + ", ".join(f"{s} {v:+.1%}" for s, v in shift.items()))

    # ---- 4) choose the cases to submit -----------------------------------------------
    pick = select_cases(proxy.to_unit(A), ens["oil"][:, -1], N_SELECT, SELECTION_MODE)
    pick = sorted(pick, key=lambda i: am[i])
    print(f"\nSELECTED CASES (mode '{SELECTION_MODE}', {len(pick)} of {len(A)} acceptable sets):")
    rows, long_rows = [], []
    for rank, i in enumerate(pick, start=1):
        row = {"case": rank, "misfit": am[i], **dict(zip(pcols, A[i]))}
        for s in STREAMS:
            row[f"{s}_cum_{fdates[-1].year}"] = ens[s][i, -1]
        rows.append(row)
        for k, d in enumerate(fdates):
            long_rows.append({"case": rank, "date": d, **{f"{s}_cum": ens[s][i, k] for s in STREAMS}})
    sel = pd.DataFrame(rows)
    sel.to_csv(OUTPUT_DIR / "selected_cases.csv", index=False)
    pd.DataFrame(long_rows).to_csv(OUTPUT_DIR / "forecast_selected_cases.csv", index=False)
    print(sel.round(4).to_string(index=False))
    plot_forecast(data, fdates, bands, {s: [ens[s][i] for i in pick] for s in STREAMS}, hist_end, obs_end)

    # ---- 5) fill a COPY of the template and check it ------------------------------------
    if tpath is None:
        tpath = make_own_template(fdates)
        wb, layout = read_template_layout(tpath)
    out_path = OUTPUT_DIR / f"{tpath.stem}_FILLED.xlsx"
    sheets = list(layout.items())
    if len(sheets) < len(pick):
        print(f"WARNING: template has {len(sheets)} case sheets but {len(pick)} cases were selected.")
    date_k = {d: k for k, d in enumerate(fdates)}
    for (name, lay), i in zip(sheets, pick):
        for r, d in lay["rows"]:
            if d in date_k:
                for s in STREAMS:
                    wb[name].cell(row=r, column=lay["cols"][s]).value = float(ens[s][i, date_k[d]])
    wb.save(out_path)
    print(f"\nsaved filled template: {out_path}   (original template untouched)")
    check_deliverable(out_path, obs_end)
    print("\nREMINDERS: (1) 2008-2028 is an extrapolation; the organisers' injection / well plan after "
          "2008 is NOT in the data.\n(2) Confirm what 'Case 1..10' mean and how they are scored; "
          "change SELECTION_MODE if needed.")


# =============================================================================
# MAIN
# =============================================================================
def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 1) load the 4 settings (01 Introduction) and the 6 production sheets
    #    (02 Field Production)
    print(f"Reading from: {DATA_DIR.resolve()}")
    params = load_parameters()
    data, labels = {}, {}
    for key, sheet in PRODUCTION_SHEETS.items():
        data[key], labels[key] = load_production(sheet)

    # 2) inventory of the heavy case files (names only, nothing heavy is read)
    n_cases = len(case_cols(next(iter(data.values()))))
    inventory_case_files(n_cases)

    # 3) check the data before trusting it
    run_sanity_checks(data, params)

    # 4) plots and tables
    plot_ensemble(data, labels, params)
    plot_parameter_coverage(params)

    summary = build_summary(data, params)
    param_cols = list(params.columns)
    plot_top_links(data, labels, summary, param_cols)
    plot_param_vs_outcome(data, labels, summary, param_cols)
    plot_correlation(summary, param_cols)
    plot_sensitivity_over_time(data, params)
    plot_history_match_summary(data, params)

    scores = rank_cases_by_misfit(data)
    plot_best_matches(data, labels, scores, params)
    plot_field_history()

    # NEW (V4): proxy model + continuous history match
    hm = run_proxy_history_match(data, params)

    # NEW (V6): forecast 2008 -> 2028, choose cases, fill the template
    if RUN_FORECAST:
        run_forecast_and_deliverable(data, params, hm)

    if SHOW_PLOTS:
        plt.show()


if __name__ == "__main__":
    main()