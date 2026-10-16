"""
ML proxy, history match and 20-year forecast (oil & gas hackathon)
==================================================================

WHAT THIS SCRIPT DOES  (it is the step AFTER the V4 exploration script)
    1. Learns a PROXY (surrogate): 4 reservoir settings  ->  production curves
    2. HISTORY MATCH: searches the 4 settings for the ones whose proxy curves
       best reproduce the real field (RAMP_History), and keeps ALL settings that
       fit about equally well (the history cannot tell them apart).
    3. FORECAST 2008 -> 2028: extends every case with a decline-curve model
       and teaches a second proxy to forecast for ANY settings.
    4. Picks the best N_SELECT cases and FILLS IN the organisers' template
       ('09 Template Deliverable') with their cumulative forecasts.

HOW THE PIECES FIT  (read this once, it explains every plot)
    The simulated curves, and the real field, STOP at 2008 (history end).
    Nothing in the data shows what happens after. So:
      * Plots P1-P4 (proxy + history match) are INSIDE the data: they can be
        validated honestly on held-out cases (71-85).
      * The forecast (P5-P6) is an EXTRAPOLATION. Its only honest test is a
        'hindcast': pretend history ended 3 years earlier, forecast, and
        compare with what the simulator really did (plot P5).
        Expect the 20-year result to be much less certain than P5 suggests.

FILES (found automatically in DATA_DIR, matched by name; see FILES below)
    01 Introduction        USED   sheet 'Uncertainty Values' = the 4 settings
    02 Field Production    USED   6 production sheets: Case_N + RAMP_History
    09 Template Deliverable USED  filled in at the end (a copy; the original
                                  is never modified). If it is missing, the
                                  script writes its own 10-sheet workbook.
    03-08 Train/Validation/Test grid files   NOT used here (see note below).
    If the files are not separate, any single .xlsx in DATA_DIR that holds
    the needed sheet names is used instead (e.g. one big trainingData.xlsx).

    NOTE ON THE GRID FILES: the 4 settings are the only inputs the history
    match can search over, so grid summaries cannot be proxy inputs. They are
    useful for diagnostics (what the grid snapshot is), which is a separate
    script.

HOW TO RUN
    pip install numpy pandas matplotlib openpyxl scipy scikit-learn
    python HistoryMatchingProxy.py path/to/folder      (or edit DATA_DIR)
    Plots, CSVs and the filled template go to OUTPUT_DIR.

OUTPUTS
    P1_proxy_validation.png        is the proxy accurate on cases it never saw?
    P2_proxy_example_curves.png    4 held-out cases: truth vs proxy (+ band)
    P3_history_match_parameters.png what the history says about each setting
    P4_history_match_fit.png       best match and spread vs the real field
    P5_hindcast_check.png          how good is the decline-curve forecast?
    P6_forecast.png                2008-2028 forecast bands + selected cases
    <template>_FILLED.xlsx         the deliverable
    + CSVs: proxy_validation, hindcast_errors, history_match_accepted,
            selected_cases, forecast_selected_cases
"""
import re
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from scipy.optimize import curve_fit, minimize
from scipy.stats import qmc
from sklearn.decomposition import PCA
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, Matern, WhiteKernel
from sklearn.linear_model import Ridge
from sklearn.preprocessing import PolynomialFeatures

warnings.filterwarnings("ignore")

# =============================================================================
# CONFIG - everything you may want to change is in this block
# =============================================================================
DATA_DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("data")
OUTPUT_DIR = Path("proxy_output")
SHOW_PLOTS = True            # False = only save PNGs, do not open windows
SEED = 0

FILES = {                    # matched by file name (case-insensitive, 'starts with')
    "intro":    "01 Introduction",         # Uncertainty Values
    "field":    "02 Field Production",     # production sheets
    "template": "09 Template Deliverable",
}
PARAM_SHEET = "Uncertainty Values"
PRODUCTION_SHEETS = {
    "gas_cum":    "Gas Production Cumulative",
    "gas_rate":   "Gas Production Rate",
    "oil_cum":    "Oil production cumulative",
    "oil_rate":   "Oil Production Rate",
    "water_cum":  "Water Production Cumulative",
    "water_rate": "Water Production Rate",
}
HISTORY_COL = "RAMP_History"
CASE_PREFIX = "Case_"
STREAMS = ["gas", "oil", "water"]                 # order of the template columns
RATE_KEYS = ["oil_rate", "gas_rate", "water_rate"]  # streams scored in history match
NICE = {"gas_cum": "Gas cumulative", "gas_rate": "Gas rate",
        "oil_cum": "Oil cumulative", "oil_rate": "Oil rate",
        "water_cum": "Water cumulative", "water_rate": "Water rate"}

# --- which cases are used for what (README: train trains, validation scores)
TRAIN_CASES = set(range(1, 71))
VALIDATION_CASES = set(range(71, 86))
TEST_CASES = set(range(86, 101))
USE_TEST_FOR_TRAINING = True    # README: test cases help train robustness
REFIT_ON_ALL = True             # final proxy uses ALL cases after validation

# --- curve preparation
RATE_RESAMPLE = "linear"        # "linear" or "previous" (step) for RATE curves
MIN_DATE_COVERAGE = 0.9         # keep history dates where >=90% of cases have values

# --- proxy
PCA_VARIANCE = 0.999            # keep enough components to explain this much
PCA_MAX_COMPONENTS = 5
GP_RESTARTS = 2
CV_FOLDS = 5

# --- history match
N_SAMPLES = 2 ** 15             # settings tried through the proxy (32768)
N_POLISH = 5                    # best samples refined with a local optimiser
ACCEPT_TOP_FRACTION = 0.01      # 'fits about as well as the best' = the best 1% of settings tried
MIN_ACCEPTED = 50               # but never fewer than this many (and >= 5 x N_SELECT)
MAX_ACCEPTED = 500
LOG_WATER = True                # model water on a log scale (it spans orders of magnitude)

# --- forecast
FORECAST_START_IF_NO_TEMPLATE = "2008-04-01"
FORECAST_END_IF_NO_TEMPLATE = "2028-01-01"
FIT_WINDOW_YEARS = None         # None = choose automatically with the hindcast
FIT_WINDOW_CANDIDATES = [2.0, 3.0, 5.0]
HINDCAST_YEARS = 3.0            # pretend history ended this long before the real end
WOR_MAX = 50.0                  # cap on water-oil ratio (50 = 98% water cut)
ANCHOR_TO_HISTORY = True        # shift forecasts so they join the real history exactly

# --- choosing the cases to submit
N_SELECT = 10
# "diverse_params": best match first, then the most different settings that still fit
# "best":           simply the lowest-misfit settings (often near duplicates)
# "span_forecast":  accepted settings spread from low to high 2028 oil forecast
SELECTION_MODE = "diverse_params"

STRONG_GOOD, STRONG_OK, STRONG_FAIR = 0.02, 0.05, 0.10   # error verdict thresholds


# =============================================================================
# STEP 1 - FIND THE EXCEL FILES AND READ SHEETS (same logic as the V4 script)
# =============================================================================
class SheetNotFound(Exception):
    pass


_NAMES_CACHE = {}


def all_xlsx():
    return [p for p in sorted(DATA_DIR.glob("*.xlsx")) if not p.name.startswith("~$")]


def find_file(stem):
    """The .xlsx in DATA_DIR whose name matches `stem`, else None."""
    target = stem.strip().lower()
    for p in all_xlsx():
        if p.stem.strip().lower() == target:
            return p
    for p in all_xlsx():
        if p.stem.strip().lower().startswith(target):
            return p
    return None


def sheet_names(path):
    key = str(path)
    if key not in _NAMES_CACHE:
        with pd.ExcelFile(path) as xl:
            _NAMES_CACHE[key] = xl.sheet_names
    return _NAMES_CACHE[key]


def find_workbook_with_sheet(sheet_name):
    target = sheet_name.strip().lower()
    for p in all_xlsx():
        if target in [n.strip().lower() for n in sheet_names(p)]:
            return p
    return None


def read_sheet(file_stem, sheet_name, header=None):
    path = find_file(file_stem)
    if path is None or sheet_name.strip().lower() not in [
            n.strip().lower() for n in sheet_names(path)]:
        path = find_workbook_with_sheet(sheet_name)     # one combined workbook?
    if path is None:
        raise FileNotFoundError(
            f"Could not find sheet '{sheet_name}' (expected in a file like "
            f"'{file_stem}.xlsx') in {DATA_DIR.resolve()}.\n"
            f"Excel files I can see there: {[p.name for p in all_xlsx()]}")
    lookup = {n.strip().lower(): n for n in sheet_names(path)}
    return pd.read_excel(path, sheet_name=lookup[sheet_name.strip().lower()], header=header)


def load_parameters():
    """Uncertainty Values -> one row per case ('Case_1'...), one column per setting."""
    df = read_sheet(FILES["intro"], PARAM_SHEET, header=0)
    df = df.rename(columns={df.columns[0]: "Item"})
    df.columns = [str(c).strip() for c in df.columns]
    df["Item"] = df["Item"].astype(str).str.strip().str.replace(" ", "_", regex=False)
    return df.set_index("Item").apply(pd.to_numeric, errors="coerce").dropna(how="all")


def load_production(sheet_name):
    """One production sheet -> clean (dates x cases) table (layout-tolerant)."""
    raw = read_sheet(FILES["field"], sheet_name, header=None)
    header_row = raw.index[raw.iloc[:, 0].astype(str).str.strip().eq("Date")][0]
    body = raw.iloc[header_row + 2:, :].copy()
    body.columns = ["Date"] + [str(c).strip() for c in raw.iloc[header_row, 1:]]
    body["Date"] = pd.to_datetime(body["Date"], errors="coerce")
    body = body.dropna(subset=["Date"]).set_index("Date").sort_index()
    body = body.apply(pd.to_numeric, errors="coerce")
    return body.loc[:, ~body.columns.isin(["nan", ""])]


def case_cols(df):
    return [c for c in df.columns if str(c).startswith(CASE_PREFIX)]


def case_number(name):
    nums = re.findall(r"\d+", str(name))
    return int(nums[-1]) if nums else -1


def split_of(name):
    n = case_number(name)
    if n in TRAIN_CASES:
        return "train"
    if n in VALIDATION_CASES:
        return "validation"
    if n in TEST_CASES:
        return "test"
    return "other"


def last_valid(series):
    s = series.dropna()
    return s.iloc[-1] if len(s) else np.nan


def resample(series, dates, kind="linear"):
    """Evaluate one case's curve at `dates` (cases report on their OWN dates).

    'linear'   : straight line between the case's points
    'previous' : step - keep the last reported value (for rates held over a step)
    Dates outside the range the case covers become NaN (no extrapolation).
    """
    s = series.dropna()
    dates = pd.DatetimeIndex(dates)
    if len(s) < 2:
        return pd.Series(np.nan, index=dates)
    x = s.index.values.astype("datetime64[ns]").astype("int64")
    xi = dates.values.astype("datetime64[ns]").astype("int64")
    if kind == "previous":
        pos = np.searchsorted(x, xi, side="right") - 1
        vals = np.where((pos >= 0) & (xi <= x[-1]), s.values[np.clip(pos, 0, None)], np.nan)
    else:
        vals = np.interp(xi, x, s.values, left=np.nan, right=np.nan)
    return pd.Series(vals, index=dates)


def kind_for(key):
    return RATE_RESAMPLE if key.endswith("_rate") else "linear"


def save(fig, name):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    fig.savefig(OUTPUT_DIR / name, dpi=130, bbox_inches="tight")
    print(f"saved {OUTPUT_DIR / name}")


# =============================================================================
# STEP 2 - TEMPLATE (the deliverable): read its layout, fill it, check it
# =============================================================================
def find_template():
    path = find_file(FILES["template"])
    if path is None:
        for p in all_xlsx():
            if "template" in p.stem.lower():
                return p
    return path


def read_template_layout(path):
    """Locate, in every 'Case N' sheet: header row, date column, the gas / oil /
    water cumulative columns, and the (row, date) of every data row."""
    wb = load_workbook(path)
    layout = {}
    for name in wb.sheetnames:
        if not re.search(r"case\s*_?\s*\d+", name, re.I):
            continue
        ws = wb[name]
        hdr = None
        for r in range(1, min(ws.max_row, 15) + 1):
            cells = [str(c.value).lower() if c.value is not None else "" for c in ws[r]]
            if sum("cumulative" in x for x in cells) >= 2:
                hdr = r
                break
        if hdr is None:
            continue
        cols = {}
        for c in ws[hdr]:
            v = str(c.value).lower() if c.value is not None else ""
            if "cumulative" in v:
                for s in STREAMS:
                    if s in v:
                        cols[s] = c.column
                        break
        counts = {}
        for col in range(1, ws.max_column + 1):
            n = sum(hasattr(ws.cell(row=r, column=col).value, "year")
                    for r in range(hdr + 1, ws.max_row + 1))
            counts[col] = n
        date_col = max(counts, key=counts.get)
        rows = [(r, pd.Timestamp(ws.cell(row=r, column=date_col).value))
                for r in range(hdr + 1, ws.max_row + 1)
                if hasattr(ws.cell(row=r, column=date_col).value, "year")]
        layout[name] = {"hdr": hdr, "cols": cols, "date_col": date_col, "rows": rows}
    ordered = dict(sorted(layout.items(), key=lambda kv: case_number(kv[0])))
    return wb, ordered


def make_own_template(dates):
    """Fallback when no template file is found: same layout, our own dates."""
    path = OUTPUT_DIR / "own_template.xlsx"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for k in range(1, N_SELECT + 1):
        ws = wb.create_sheet(f"Case {k}")
        ws.append(["Date", "Gas production cumulative [MSCF]",
                   "Oil production cumulative [STB]", "Water production cumulative [STB]"])
        for d in dates:
            ws.append([d.to_pydatetime()])
    wb.save(path)
    return path


# =============================================================================
# STEP 3 - THE PROXY: PCA (compress each curve) + Gaussian process (settings -> curve)
# =============================================================================
class Scaler:
    """Maps the 4 settings to the unit cube [0,1]^4 using the tested ranges."""

    def __init__(self, params):
        self.cols = list(params.columns)
        self.lo = params.min().values.astype(float)
        self.hi = params.max().values.astype(float)

    def to_unit(self, X):
        return (np.asarray(X, float) - self.lo) / (self.hi - self.lo)

    def from_unit(self, U):
        return self.lo + np.asarray(U, float) * (self.hi - self.lo)


class CurveProxy:
    """settings -> a whole curve.

    1. PCA squeezes each curve (e.g. 37 dates) into a few numbers.
    2. One model per number: 'gp' = Gaussian process (smooth, gives an
       uncertainty), 'poly' = quadratic ridge regression (simple baseline).
    3. The predicted numbers are expanded back into a curve.
    """

    def __init__(self, kind="gp", log=False):
        self.kind = kind
        self.log = log

    def fit(self, X, Y):
        X, Y = np.asarray(X, float), np.asarray(Y, float)
        if self.log:                       # fit on a log scale, then transform back
            self.c = 1e-3 * max(float(Y.max()), 1e-12)
            Y = np.log(np.maximum(Y, 0) + self.c)
        n_max = min(PCA_MAX_COMPONENTS, X.shape[0] - 1, Y.shape[1])
        pca = PCA(n_components=n_max).fit(Y)
        k = int(np.searchsorted(np.cumsum(pca.explained_variance_ratio_), PCA_VARIANCE) + 1)
        self.k = max(1, min(k, n_max))
        self.mean = pca.mean_
        self.comp = pca.components_[:self.k]
        Z = (Y - self.mean) @ self.comp.T
        self.mu, self.sd = Z.mean(0), Z.std(0) + 1e-12
        Zs = (Z - self.mu) / self.sd
        d = X.shape[1]
        if self.kind == "gp":
            self.models = []
            for j in range(self.k):
                kern = (ConstantKernel(1.0, (1e-2, 1e2))
                        * Matern(length_scale=np.ones(d), length_scale_bounds=(0.05, 50.0), nu=2.5)
                        + WhiteKernel(1e-3, (1e-8, 1e-1)))
                gp = GaussianProcessRegressor(kern, n_restarts_optimizer=GP_RESTARTS,
                                              random_state=SEED)
                self.models.append(gp.fit(X, Zs[:, j]))
        else:
            self.poly = PolynomialFeatures(2)
            self.ridge = Ridge(alpha=1e-3).fit(self.poly.fit_transform(X), Zs)
        return self

    def predict(self, X, return_std=False):
        X = np.asarray(X, float)
        if self.kind == "gp":
            out = [m.predict(X, return_std=True) for m in self.models]
            Zs = np.column_stack([o[0] for o in out])
            Zsd = np.column_stack([o[1] for o in out])
        else:
            Zs = self.ridge.predict(self.poly.transform(X)).reshape(len(X), -1)
            Zsd = np.zeros_like(Zs)
        Yhat = self.mean + (Zs * self.sd + self.mu) @ self.comp
        contrib = (Zsd * self.sd)[:, :, None] * self.comp[None, :, :]
        std = np.sqrt((contrib ** 2).sum(axis=1))
        if self.log:
            Yhat = np.exp(Yhat)
            std = std * Yhat
            Yhat = Yhat - self.c
        return (Yhat, std) if return_std else Yhat


def curve_errors(Yt, Yp):
    """Per case: curve RMSE as a fraction of that case's own peak value."""
    return np.sqrt(np.mean((Yt - Yp) ** 2, axis=1)) / (np.abs(Yt).max(axis=1) + 1e-12)


def final_r2(Yt, Yp):
    yt, yp = Yt[:, -1], Yp[:, -1]
    ss = ((yt - yt.mean()) ** 2).sum()
    return float(1 - ((yt - yp) ** 2).sum() / ss) if ss > 0 else np.nan


def verdict_word(frac):
    if frac < STRONG_GOOD:
        return "excellent"
    if frac < STRONG_OK:
        return "good"
    if frac < STRONG_FAIR:
        return "fair"
    return "POOR"


VERDICT_COLOR = {"excellent": "#1a7f37", "good": "#1a7f37", "fair": "#b8860b", "POOR": "#b22222"}


def use_log(name):
    return LOG_WATER and "water" in name


def kfold_predict(X, Y, kind, k, log=False):
    """Out-of-fold predictions: every case is predicted by a model that never saw it."""
    idx = np.random.default_rng(SEED).permutation(len(X))
    Yp = np.zeros_like(Y)
    for f in range(k):
        te = idx[f::k]
        tr = np.setdiff1d(idx, te)
        Yp[te] = CurveProxy(kind, log).fit(X[tr], Y[tr]).predict(X[te])
    return Yp


def validate_proxy(name, X, Y, ids, kinds=("gp", "poly")):
    """Score proxies on held-out cases. Prefers the README's validation cases
    (71-85); falls back to k-fold cross-validation if they are not available."""
    split = np.array([split_of(c) for c in ids])
    tr = np.isin(split, ["train", "test"] if USE_TEST_FOR_TRAINING else ["train"])
    va = split == "validation"
    res = {"name": name, "kinds": {}}
    if va.sum() >= 3 and tr.sum() >= 10:
        res["how"] = f"validation cases ({int(va.sum())} held out, never used in fitting)"
        res["ids"], res["Yt"] = np.array(ids)[va], Y[va]
        for kind in kinds:
            mdl = CurveProxy(kind, use_log(name)).fit(X[tr], Y[tr])
            if kind == "gp":
                res["kinds"][kind], res["std"] = mdl.predict(X[va], return_std=True)
            else:
                res["kinds"][kind] = mdl.predict(X[va])
    else:
        res["how"] = f"{CV_FOLDS}-fold cross-validation (validation cases not found)"
        res["ids"], res["Yt"] = np.array(ids), Y
        for kind in kinds:
            res["kinds"][kind] = kfold_predict(X, Y, kind, CV_FOLDS, use_log(name))
    res["rows"] = []
    for kind, Yp in res["kinds"].items():
        err = curve_errors(res["Yt"], Yp)
        res["rows"].append({"stream": name, "model": kind, "median_error": float(np.median(err)),
                            "worst_error": float(err.max()), "final_R2": final_r2(res["Yt"], Yp),
                            "verdict": verdict_word(float(np.median(err))), "how": res["how"]})
    return res


# =============================================================================
# STEP 4 - CURVE MATRICES (every case on a common set of history dates)
# =============================================================================
def build_curve_matrix(df, key):
    """cases x dates, evaluated at the dates where the real field has history."""
    hist_dates = df[HISTORY_COL].dropna().index
    M = pd.DataFrame({c: resample(df[c], hist_dates, kind_for(key)) for c in case_cols(df)})
    M = M.loc[M.notna().mean(axis=1) >= MIN_DATE_COVERAGE].dropna(axis=1)
    return M.T


# =============================================================================
# STEP 5 - HISTORY MATCH
# =============================================================================
def misfit(Upool, proxies, obs, scale):
    """Total mismatch of the proxy against the real field.

    For each rate (oil, gas, water): RMSE over the history dates divided by the
    largest history value, then summed. Same recipe as the V4 script, so the
    numbers are comparable with the V4 misfit scores. Lower = better.
    """
    total = np.zeros(len(Upool))
    for key in RATE_KEYS:
        pred = proxies[key].predict(Upool)
        total += np.sqrt(np.mean(((pred - obs[key]) / scale[key]) ** 2, axis=1))
    return total


def run_history_match(proxies, obs, scale, d):
    sob = qmc.Sobol(d=d, scramble=True, seed=SEED)
    pool = sob.random(N_SAMPLES)
    m = misfit(pool, proxies, obs, scale)
    starts = pool[np.argsort(m)[:N_POLISH]]
    polished = []
    for s0 in starts:
        r = minimize(lambda u: misfit(np.clip(u, 0, 1)[None, :], proxies, obs, scale)[0],
                     s0, bounds=[(0, 1)] * d, method="L-BFGS-B")
        polished.append(np.clip(r.x, 0, 1))
    pool = np.vstack([pool, np.array(polished)])
    m = misfit(pool, proxies, obs, scale)
    best = float(m.min())
    n_keep = max(int(len(m) * ACCEPT_TOP_FRACTION), MIN_ACCEPTED, N_SELECT * 5)
    acc = np.argsort(m)[:min(n_keep, MAX_ACCEPTED)]      # best first
    return pool, m, acc, best


# =============================================================================
# STEP 6 - FORECAST: decline curves (oil, gas) + water-oil ratio trend (water)
# =============================================================================
def arps(t, qi, Di, b):
    """Hyperbolic decline: rate falls fast at first, then more slowly."""
    return qi / np.power(1.0 + b * Di * t, 1.0 / b)


def fit_arps(t, q):
    """Fit qi, Di, b to a rate history (t in years from the window start)."""
    m = np.isfinite(t) & np.isfinite(q) & (q > 0)
    if m.sum() < 5:
        return None
    t, q = t[m], q[m]
    lq = np.log(q)

    def f(tt, lqi, lDi, b):
        return np.log(arps(tt, np.exp(lqi), np.exp(lDi), b))

    best = None
    for b0 in (0.1, 0.5, 0.9):
        try:
            popt, _ = curve_fit(f, t, lq, p0=[lq[0], np.log(0.15), b0],
                                bounds=([lq.min() - 3, -9, 0.01], [lq.max() + 3, 2.0, 1.0]),
                                maxfev=4000)
            sse = np.sum((f(t, *popt) - lq) ** 2)
            if best is None or sse < best[0]:
                best = (sse, popt)
        except Exception:
            continue
    if best is None:                                   # fallback: simple exponential
        slope, icpt = np.polyfit(t, lq, 1)
        return (float(np.exp(icpt)), float(max(-slope, 1e-6)), 0.01)
    return (float(np.exp(best[1][0])), float(np.exp(best[1][1])), float(best[1][2]))


def fit_shapes(data, case, cut, window_years):
    """Everything the forecast needs from one case, using history up to `cut`:
    the state at `cut` (cumulatives, rates) and the decline shapes fitted on
    the last `window_years` of history."""
    w0 = cut - pd.Timedelta(days=int(365.25 * window_years))
    out = {"w0": w0, "cut": cut}
    for s in ("oil", "gas"):
        ser = data[f"{s}_rate"][case].dropna()
        ser = ser[(ser.index >= w0) & (ser.index <= cut)]
        t = (ser.index - w0).days.values / 365.25
        out[s] = fit_arps(t, ser.values)
        if out[s] is None:
            return None
    cutv = pd.DatetimeIndex([cut])
    st = {}
    for key in ["oil_cum", "gas_cum", "water_cum", "oil_rate", "gas_rate", "water_rate"]:
        st[key] = float(resample(data[key][case], cutv, kind_for(key)).iloc[0])
    if not np.all(np.isfinite(list(st.values()))):
        return None
    out["state"] = st
    grid = pd.date_range(w0, cut, periods=24)
    qo = resample(data["oil_rate"][case], grid, "linear").values
    qw = resample(data["water_rate"][case], grid, "linear").values
    npv = resample(data["oil_cum"][case], grid, "linear").values
    with np.errstate(divide="ignore", invalid="ignore"):
        wor = np.where(qo > 0, qw / qo, np.nan)
    ok = np.isfinite(wor) & np.isfinite(npv) & (wor > 1e-4)
    scale = max(st["oil_cum"], 1.0)
    bw = 0.0
    if ok.sum() >= 5:
        slope, _ = np.polyfit(npv[ok] / scale, np.log(wor[ok]), 1)
        bw = max(float(slope), 0.0)                    # water cut does not fall with production
    out["wor_slope"] = bw
    return out


def project(shp, out_dates):
    """Cumulative gas / oil / water at `out_dates` (all after shp['cut'])."""
    cut, st = shp["cut"], shp["state"]
    out_dates = pd.DatetimeIndex(out_dates)
    dmax = (out_dates.max() - cut).days
    d = np.arange(0, dmax + 8, 7.0)
    off = (cut - shp["w0"]).days / 365.25
    ty = off + d / 365.25
    dd = np.diff(d)

    def integrate(q):
        return np.concatenate([[0.0], np.cumsum(0.5 * (q[1:] + q[:-1]) * dd)])

    q = {}
    for s in ("oil", "gas"):
        qi, Di, b = shp[s]
        q[s] = arps(ty, qi, Di, b) * (st[f"{s}_rate"] / max(arps(off, qi, Di, b), 1e-12))
    cum_o = st["oil_cum"] + integrate(q["oil"])
    cum_g = st["gas_cum"] + integrate(q["gas"])
    wor_end = max(st["water_rate"] / max(st["oil_rate"], 1e-9), 1e-4)
    cap = max(WOR_MAX, wor_end)
    wor = np.minimum(wor_end * np.exp(shp["wor_slope"] * (cum_o - st["oil_cum"])
                                      / max(st["oil_cum"], 1.0)), cap)
    q_w = wor * q["oil"]
    cum_w = st["water_cum"] + integrate(q_w)
    x = (out_dates - cut).days.values.astype(float)
    return {"gas": np.interp(x, d, cum_g), "oil": np.interp(x, d, cum_o),
            "water": np.interp(x, d, cum_w)}


def hindcast(data, cases, window_years, hist_end):
    """Pretend history ended HINDCAST_YEARS earlier. Forecast, then compare the
    predicted production increase with what the simulator really did."""
    cut = hist_end - pd.Timedelta(days=int(365.25 * HINDCAST_YEARS))
    rows = {}
    for c in cases:
        shp = fit_shapes(data, c, cut, window_years)
        if shp is None:
            continue
        pred = project(shp, [hist_end])
        r = {}
        for s in STREAMS:
            true_end = float(resample(data[f"{s}_cum"][c], pd.DatetimeIndex([hist_end])).iloc[0])
            true_inc = true_end - shp["state"][f"{s}_cum"]
            r[s] = (pred[s][0] - shp["state"][f"{s}_cum"]) / true_inc - 1 if true_inc > 0 else np.nan
        rows[c] = r
    return pd.DataFrame(rows).T


# =============================================================================
# STEP 7 - CHOOSING THE CASES TO SUBMIT
# =============================================================================
def select_cases(U, m, oil2028, n, mode):
    """Indices into the accepted set (already sorted best-misfit first)."""
    n = min(n, len(U))
    if mode == "best":
        return list(range(n))
    if mode == "span_forecast":
        order = np.argsort(oil2028)
        pick = order[np.unique(np.round(np.linspace(0.03, 0.97, n) * (len(order) - 1)).astype(int))]
        return sorted(pick.tolist())
    chosen = [0]                                       # best match first
    while len(chosen) < n:
        dist = np.min(np.linalg.norm(U[:, None, :] - U[chosen][None, :, :], axis=2), axis=1)
        dist[chosen] = -1
        chosen.append(int(np.argmax(dist)))
    return chosen


def check_deliverable(path, obs_end):
    """Read the FILLED template back and verify the rules of the deliverable."""
    wb, layout = read_template_layout(path)
    print("\n=== DELIVERABLE CHECK (reading the saved file back) ===")
    all_ok = True
    for name, lay in layout.items():
        ws = wb[name]
        problems = []
        for s in STREAMS:
            vals = np.array([ws.cell(row=r, column=lay["cols"][s]).value
                             if ws.cell(row=r, column=lay["cols"][s]).value is not None else np.nan
                             for r, _ in lay["rows"]], float)
            if np.isnan(vals).any():
                problems.append(f"{s}: blank cells")
                continue
            if (vals < 0).any():
                problems.append(f"{s}: negative values")
            if (np.diff(vals) < -1e-6 * max(abs(vals).max(), 1)).any():
                problems.append(f"{s}: cumulative decreases")
            if s in obs_end and np.isfinite(obs_end[s]) and vals[0] < obs_end[s] * (1 - 1e-9):
                problems.append(f"{s}: starts below the 2008 history value")
        status = "PASS" if not problems else "FAIL: " + "; ".join(problems)
        all_ok &= not problems
        print(f"  {name:8s} rows={len(lay['rows'])}  {status}")
    print("  ALL SHEETS PASS" if all_ok else "  FIX THE SHEETS MARKED FAIL BEFORE SUBMITTING")


# =============================================================================
# PLOTS
# =============================================================================
def plot_proxy_validation(results):
    order = ["oil_cum", "gas_cum", "water_cum", "oil_rate", "gas_rate", "water_rate"]
    fig, axes = plt.subplots(2, 3, figsize=(17, 10), constrained_layout=True)
    how = ""
    for ax, key in zip(axes.ravel(), order):
        r = results[key]
        how = r["how"]
        yt = r["Yt"][:, -1]
        row = [x for x in r["rows"] if x["model"] == "gp"][0]
        ax.scatter(yt, r["kinds"]["gp"][:, -1], s=45, color="#1f77b4", edgecolor="k", label="Gaussian process")
        ax.scatter(yt, r["kinds"]["poly"][:, -1], s=45, marker="x", color="#ff7f0e", label="simple polynomial")
        lim = [min(yt.min(), r["kinds"]["gp"][:, -1].min()), max(yt.max(), r["kinds"]["gp"][:, -1].max())]
        ax.plot(lim, lim, "k--", lw=1, label="perfect")
        ax.set_xlabel("simulator (truth)")
        ax.set_ylabel("proxy prediction")
        ax.set_title(f"{NICE[key]} at end of history\ntypical curve error {row['median_error']:.1%} "
                     f"({row['verdict']}), R\u00b2={row['final_R2']:.3f}",
                     color=VERDICT_COLOR[row["verdict"]], weight="bold", fontsize=11)
        ax.grid(alpha=0.3)
    axes[0, 0].legend(fontsize=9)
    fig.suptitle(f"P1 - Is the proxy accurate on cases it never saw?  ({how})\n"
                 "Dots on the dashed line = proxy equals simulator.  Error = curve RMSE as a % of "
                 "each case's peak.", fontsize=13, weight="bold")
    save(fig, "P1_proxy_validation.png")


def plot_proxy_examples(results, scaler):
    keys = ["oil_cum", "gas_cum", "water_cum"]
    r0 = results["oil_cum"]
    err = curve_errors(r0["Yt"], r0["kinds"]["gp"])
    order = np.argsort(err)
    pick = [order[0], order[len(order) // 3], order[2 * len(order) // 3], order[-1]]
    labels = ["best", "typical", "typical", "WORST"]
    fig, axes = plt.subplots(3, 4, figsize=(18, 10), constrained_layout=True, sharex="col")
    for j, (pi, lab) in enumerate(zip(pick, labels)):
        cid = r0["ids"][pi]
        for i, key in enumerate(keys):
            r = results[key]
            if cid not in list(r["ids"]):
                continue
            ix = list(r["ids"]).index(cid)
            ax = axes[i, j]
            dates = r["dates"]
            ax.plot(dates, r["Yt"][ix], "k-", lw=2.5, label="simulator")
            ax.plot(dates, r["kinds"]["gp"][ix], "r--", lw=2, label="proxy (GP)")
            ax.plot(dates, r["kinds"]["poly"][ix], color="#ff7f0e", ls=":", lw=2, label="polynomial")
            if "std" in r:
                s = r["std"][ix]
                ax.fill_between(dates, r["kinds"]["gp"][ix] - 2 * s, r["kinds"]["gp"][ix] + 2 * s,
                                color="red", alpha=0.15, label="proxy \u00b12\u03c3")
            if i == 0:
                ax.set_title(f"{cid} ({lab} oil match)", weight="bold")
            if j == 0:
                ax.set_ylabel(NICE[key])
            ax.grid(alpha=0.3)
    axes[0, 0].legend(fontsize=8)
    fig.suptitle("P2 - Held-out cases: simulator truth vs what the proxy predicts "
                 "from the 4 settings alone", fontsize=13, weight="bold")
    save(fig, "P2_proxy_example_curves.png")


def plot_history_match_parameters(params, scaler, U, m, acc, best_i):
    names = list(params.columns)
    fig, axes = plt.subplots(1, len(names), figsize=(4.6 * len(names), 4.8), constrained_layout=True)
    A = scaler.from_unit(U[acc])
    bestv = scaler.from_unit(U[best_i])
    rows = []
    for ax, j, nm in zip(np.atleast_1d(axes), range(len(names)), names):
        lo, hi = scaler.lo[j], scaler.hi[j]
        share = (A[:, j].max() - A[:, j].min()) / (hi - lo)
        if share < 0.5:
            v, col = "PINNED DOWN", "#1a7f37"
        elif share < 0.8:
            v, col = "partly constrained", "#b8860b"
        else:
            v, col = "NOT constrained", "#b22222"
        ax.hist(A[:, j], bins=15, range=(lo, hi), color=col, alpha=0.7, edgecolor="k")
        ax.axvline(bestv[j], color="k", lw=2.5)
        ax.set_xlim(lo, hi)
        ax.set_title(f"{nm}\n{v}", color=col, weight="bold")
        ax.set_xlabel(f"tested {lo:.3g} to {hi:.3g}   (black line = best)")
        ax.grid(alpha=0.3)
        rows.append((nm, v, A[:, j].min(), A[:, j].max(), bestv[j], lo, hi))
    axes_list = np.atleast_1d(axes)
    axes_list[0].set_ylabel("number of accepted settings")
    fig.suptitle(f"P3 - What the history says about each setting  (the {len(acc)} best-fitting settings out of "
                 f"{N_SAMPLES} tried)\n"
                 "Narrow bar = history determines it.  Spread across the whole range = it cannot.",
                 fontsize=12, weight="bold")
    save(fig, "P3_history_match_parameters.png")
    print("\nWhat the history says about each setting (all settings that fit about as well as the best):")
    for nm, v, a0, a1, b, lo, hi in rows:
        print(f"  {v:18s} {nm}: best {b:.4g}; accepted {a0:.4g} to {a1:.4g} (tested {lo:.4g} to {hi:.4g})")
    return rows


def plot_history_match_fit(data, mats, proxies, scaler, U, acc, best_i):
    fig, axes = plt.subplots(1, 3, figsize=(18, 5.2), constrained_layout=True)
    sub = acc[:300]
    for ax, key in zip(axes, RATE_KEYS):
        M = mats[key]
        dates = M.columns
        for i in range(len(M)):
            ax.plot(dates, M.iloc[i].values, color="#bbbbbb", lw=0.6)
        band = proxies[key].predict(U[sub])
        ax.fill_between(dates, np.percentile(band, 10, axis=0), np.percentile(band, 90, axis=0),
                        color="#1f77b4", alpha=0.35, label="accepted settings (P10-P90)")
        ax.plot(dates, proxies[key].predict(U[[best_i]])[0], color="#d62728", lw=2.5, label="best match (proxy)")
        hist = data[key][HISTORY_COL].dropna()
        ax.plot(hist.index, hist.values, "ko", ms=4, label="real field")
        ax.set_title(NICE[key], weight="bold")
        ax.grid(alpha=0.3)
    axes[0].legend(fontsize=9)
    fig.suptitle("P4 - History match: grey = all simulated cases, blue band = settings that fit about as "
                 "well as the best, red = best, dots = real field", fontsize=12, weight="bold")
    save(fig, "P4_history_match_fit.png")


def plot_hindcast(errs_by_window, chosen):
    fig, (a, b) = plt.subplots(1, 2, figsize=(16, 5.4), constrained_layout=True)
    w = list(errs_by_window)
    xs = np.arange(len(STREAMS))
    for k, wy in enumerate(w):
        med = [errs_by_window[wy][s].abs().median() * 100 for s in STREAMS]
        a.bar(xs + k * 0.25, med, 0.25, label=f"fit on last {wy:g} yr" + ("  <- chosen" if wy == chosen else ""),
              edgecolor="k", alpha=1.0 if wy == chosen else 0.55)
    a.set_xticks(xs + 0.25)
    a.set_xticklabels([s.capitalize() for s in STREAMS])
    a.set_ylabel(f"typical error of the {HINDCAST_YEARS:g}-year production increase [%]")
    a.set_title("Which fitting window forecasts best?  (lower is better)", weight="bold")
    a.legend()
    a.grid(alpha=0.3, axis="y")
    E = errs_by_window[chosen]
    b.boxplot([E[s].dropna().values * 100 for s in STREAMS], showfliers=False)
    b.set_xticklabels([s.capitalize() for s in STREAMS])
    b.axhline(0, color="k", lw=1)
    b.set_ylabel("forecast error [%]   (above 0 = over-predicted)")
    b.set_title(f"Chosen window ({chosen:g} yr): spread of errors across cases", weight="bold")
    b.grid(alpha=0.3, axis="y")
    fig.suptitle(f"P5 - Hindcast: history cut {HINDCAST_YEARS:g} years early, forecast, compare with what the "
                 "simulator really did.\nA 20-year forecast will be LESS accurate than this.",
                 fontsize=12, weight="bold")
    save(fig, "P5_hindcast_check.png")


def plot_forecast(data, fdates, bands, selected, hist_end, obs_cum_hist):
    fig, axes = plt.subplots(2, 3, figsize=(19, 10), constrained_layout=True, sharex=True)
    cmap = plt.get_cmap("viridis")
    for j, s in enumerate(STREAMS):
        ax = axes[0, j]
        h = data[f"{s}_cum"][HISTORY_COL].dropna()
        ax.plot(h.index, h.values, "ko", ms=3, label="real history")
        lo, mid, hi = bands[s]
        ax.fill_between(fdates, lo, hi, color="#1f77b4", alpha=0.3, label="accepted settings P10-P90")
        ax.plot(fdates, mid, color="#1f77b4", lw=2, label="P50")
        for rank, c in enumerate(selected[s]):
            ax.plot(fdates, c, color=cmap(rank / max(len(selected[s]) - 1, 1)), lw=1.2, alpha=0.9)
        ax.axvline(hist_end, color="k", ls=":")
        ax.set_title(f"{s.capitalize()} cumulative", weight="bold")
        ax.grid(alpha=0.3)
        ax2 = axes[1, j]
        hr = data[f"{s}_rate"][HISTORY_COL].dropna()
        ax2.plot(hr.index, hr.values, "ko", ms=3, label="real history")
        days = np.diff(np.concatenate([[hist_end.value], fdates.values.astype("datetime64[ns]").astype("int64")])) / 8.64e13
        for name, arr, col, lw in [("P10", lo, "#1f77b4", 1), ("P50", mid, "#1f77b4", 2), ("P90", hi, "#1f77b4", 1)]:
            prev = np.concatenate([[obs_cum_hist[s]], arr[:-1]])
            ax2.plot(fdates, (arr - prev) / days, color=col, lw=lw, ls="-" if name == "P50" else "--")
        for rank, c in enumerate(selected[s]):
            prev = np.concatenate([[obs_cum_hist[s]], c[:-1]])
            ax2.plot(fdates, (c - prev) / days, color=cmap(rank / max(len(selected[s]) - 1, 1)), lw=1.0, alpha=0.8)
        ax2.axvline(hist_end, color="k", ls=":")
        ax2.set_title(f"{s.capitalize()} rate (implied by the cumulative)", weight="bold")
        ax2.grid(alpha=0.3)
    axes[0, 0].legend(fontsize=9, loc="upper left")
    fig.suptitle("P6 - Forecast 2008 to 2028.  Coloured lines = the cases submitted (dark = best history match).  "
                 "Right of the dotted line is EXTRAPOLATION: the band shows only the spread between "
                 "equally good matches, not all the uncertainty.", fontsize=12, weight="bold")
    save(fig, "P6_forecast.png")


# =============================================================================
# MAIN
# =============================================================================
def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"Looking for Excel files in: {DATA_DIR.resolve()}")
    print(f"  found: {[p.name for p in all_xlsx()]}")

    # ---- load -------------------------------------------------------------
    params = load_parameters()
    data = {k: load_production(sh) for k, sh in PRODUCTION_SHEETS.items()}
    hist_end = data["oil_cum"][HISTORY_COL].dropna().index.max()
    print(f"\n=== DATA ===\nsettings: {params.shape[0]} cases x {list(params.columns)}")
    print(f"history ends: {hist_end.date()}   (simulated curves stop here too)")
    for k, df in data.items():
        cc = case_cols(df)
        print(f"  {k:10s} cases={len(cc)} history points={df[HISTORY_COL].notna().sum()} "
              f"last date={df.index.max().date()}")
    if any(df.index.max() > hist_end + pd.Timedelta(days=1) for df in data.values()):
        print("  NOTE: some sheets contain dates AFTER the history end - check that the simulator "
              "curves do not already hold future data.")
    ids_all = [c for c in params.index if c in set(case_cols(data["oil_cum"]))]
    counts = pd.Series([split_of(c) for c in ids_all]).value_counts().to_dict()
    print(f"  cases with settings and curves: {len(ids_all)}  splits: {counts}")

    # ---- template ---------------------------------------------------------
    tpath = find_template()
    layout = None
    if tpath is not None:
        wb, layout = read_template_layout(tpath)
        print(f"\ntemplate found: {tpath.name}  ({len(layout)} case sheets)")
        first = next(iter(layout.values()))
        fdates = pd.DatetimeIndex([d for _, d in first["rows"]])
        print(f"  dates: {fdates.min().date()} -> {fdates.max().date()} ({len(fdates)} rows)")
        if (fdates <= hist_end).any():
            print("  WARNING: template has dates on/before the history end; those rows are skipped.")
            fdates = fdates[fdates > hist_end]
    else:
        print("\nno template file found - will write our own 10-sheet workbook.")
        fdates = pd.date_range(FORECAST_START_IF_NO_TEMPLATE, FORECAST_END_IF_NO_TEMPLATE, freq="QS")

    scaler = Scaler(params.loc[ids_all])
    Xall = scaler.to_unit(params.loc[ids_all].values)

    # ---- history proxy: validate -----------------------------------------
    print("\n=== PROXY VALIDATION (settings -> history curves) ===")
    mats, results, rows = {}, {}, []
    for key in PRODUCTION_SHEETS:
        M = build_curve_matrix(data[key], key)
        M = M.loc[[c for c in ids_all if c in M.index]]
        mats[key] = M
        X = scaler.to_unit(params.loc[M.index].values)
        res = validate_proxy(key, X, M.values, list(M.index))
        res["dates"] = M.columns
        results[key] = res
        rows += res["rows"]
        for r in res["rows"]:
            print(f"  {NICE[key]:17s} {r['model']:4s} typical error {r['median_error']:6.2%} "
                  f"worst {r['worst_error']:6.2%}  R2(final) {r['final_R2']:.3f}  -> {r['verdict']}")
    print(f"  (scored on: {results['oil_cum']['how']})")
    pd.DataFrame(rows).to_csv(OUTPUT_DIR / "proxy_validation.csv", index=False)
    plot_proxy_validation(results)
    plot_proxy_examples(results, scaler)

    # ---- history proxy: final fit ----------------------------------------
    fit_mask = {}
    proxies = {}
    for key, M in mats.items():
        X = scaler.to_unit(params.loc[M.index].values)
        if REFIT_ON_ALL:
            sel = np.ones(len(M), bool)
        else:
            sel = np.isin([split_of(c) for c in M.index], ["train", "test"])
        proxies[key] = CurveProxy("gp", use_log(key)).fit(X[sel], M.values[sel])
    print(f"\nfinal proxies fitted on {'ALL cases' if REFIT_ON_ALL else 'train + test cases'}")

    # ---- history match ---------------------------------------------------
    print("\n=== HISTORY MATCH ===")
    obs, scale = {}, {}
    for key in RATE_KEYS:
        dts = mats[key].columns
        obs[key] = data[key][HISTORY_COL].reindex(dts).values.astype(float)
        scale[key] = max(np.nanmax(np.abs(obs[key])), 1e-9)
    # reference: best of the 100 simulated cases (same score, computed directly)
    sim_scores = np.zeros(len(ids_all))
    for key in RATE_KEYS:
        M = mats[key].reindex(ids_all)
        sim_scores += np.sqrt(np.nanmean(((M.values - obs[key]) / scale[key]) ** 2, axis=1))
    print(f"  best of the {len(ids_all)} simulated cases: misfit {np.nanmin(sim_scores):.4f} "
          f"({ids_all[int(np.nanargmin(sim_scores))]})")
    pool, m, acc, best = run_history_match(proxies, obs, scale, Xall.shape[1])
    best_i = int(np.argmin(m))
    print(f"  best settings found through the proxy: predicted misfit {best:.4f}")
    print("  (proxy-predicted: it has not been re-run in the simulator, so treat it as a "
          "candidate, not a proven match)")
    print(f"  accepted settings = the best {len(acc)} of {len(m)} tried (misfit up to {m[acc].max():.4f} "
          f"vs best {best:.4f}); these 'fit about equally well'")
    bestv = scaler.from_unit(pool[best_i])
    for nm, v in zip(params.columns, bestv):
        print(f"    {nm}: {v:.4g}")
    plot_history_match_parameters(params, scaler, pool, m, acc, best_i)
    plot_history_match_fit(data, mats, proxies, scaler, pool, acc, best_i)
    acc_df = pd.DataFrame(scaler.from_unit(pool[acc]), columns=params.columns)
    acc_df["misfit"] = m[acc]
    acc_df.to_csv(OUTPUT_DIR / "history_match_accepted.csv", index=False)

    # ---- forecast layer ----------------------------------------------------
    print("\n=== FORECAST (extrapolation beyond the simulated data) ===")
    errs = {}
    for wy in FIT_WINDOW_CANDIDATES:
        e = hindcast(data, ids_all, wy, hist_end)
        errs[wy] = e
        print(f"  hindcast, fit on last {wy:g} yr: typical error of the {HINDCAST_YEARS:g}-yr increase  "
              + "  ".join(f"{s} {e[s].abs().median():.1%}" for s in STREAMS))
    if FIT_WINDOW_YEARS is None:
        chosen = min(errs, key=lambda w: np.nanmean([errs[w][s].abs().median() for s in STREAMS]))
    else:
        chosen = FIT_WINDOW_YEARS
        errs.setdefault(chosen, hindcast(data, ids_all, chosen, hist_end))
    print(f"  chosen fitting window: last {chosen:g} years")
    pd.concat({f"{w:g}yr": e for w, e in errs.items()}, axis=1).to_csv(OUTPUT_DIR / "hindcast_errors.csv")
    plot_hindcast({w: errs[w] for w in FIT_WINDOW_CANDIDATES if w in errs} | {chosen: errs[chosen]}, chosen)

    F = {s: {} for s in STREAMS}
    for c in ids_all:
        shp = fit_shapes(data, c, hist_end, chosen)
        if shp is None:
            continue
        pr = project(shp, fdates)
        for s in STREAMS:
            F[s][c] = pr[s]
    ok_cases = [c for c in ids_all if c in F["oil"]]
    print(f"  cases with a usable forecast fit: {len(ok_cases)} of {len(ids_all)}")
    Xf = scaler.to_unit(params.loc[ok_cases].values)
    fproxy = {}
    for s in STREAMS:
        Y = np.array([F[s][c] for c in ok_cases])
        r = validate_proxy(f"forecast_{s}", Xf, Y, ok_cases, kinds=("gp",))
        print(f"  forecast proxy {s:5s}: reproduces the per-case extrapolations to "
              f"{r['rows'][0]['median_error']:.2%} ({r['rows'][0]['verdict']}) on held-out cases")
        fproxy[s] = CurveProxy("gp", use_log(s)).fit(Xf, Y)

    # ---- ensemble + selection ---------------------------------------------
    obs_end = {}
    for s in STREAMS:
        v = data[f"{s}_cum"][HISTORY_COL].dropna()
        obs_end[s] = float(v.iloc[-1]) if len(v) else np.nan
    cum_idx = {s: mats[f"{s}_cum"] for s in STREAMS}
    Uacc = pool[acc]
    ens = {}
    for s in STREAMS:
        pred = fproxy[s].predict(Uacc)
        if ANCHOR_TO_HISTORY and np.isfinite(obs_end[s]):
            hist_pred_end = proxies[f"{s}_cum"].predict(Uacc)[:, -1]
            pred = pred + (obs_end[s] - hist_pred_end)[:, None]
        pred = np.maximum.accumulate(np.maximum(pred, 0), axis=1)
        if np.isfinite(obs_end[s]):
            pred = np.maximum(pred, obs_end[s])
        ens[s] = pred
    bands = {s: (np.percentile(ens[s], 10, axis=0), np.percentile(ens[s], 50, axis=0),
                 np.percentile(ens[s], 90, axis=0)) for s in STREAMS}
    chosen_idx = select_cases(Uacc, m[acc], ens["oil"][:, -1], N_SELECT, SELECTION_MODE)
    chosen_idx = sorted(chosen_idx, key=lambda i: m[acc][i])
    print(f"\n=== SELECTED CASES (mode: {SELECTION_MODE}) ===")
    sel_rows, long_rows = [], []
    sel_curves = {s: [ens[s][i] for i in chosen_idx] for s in STREAMS}
    for rank, i in enumerate(chosen_idx, start=1):
        th = scaler.from_unit(Uacc[i])
        row = {"case": rank, "misfit": m[acc][i]}
        row.update({nm: v for nm, v in zip(params.columns, th)})
        for s in STREAMS:
            row[f"{s}_cum_{fdates[-1].year}"] = ens[s][i, -1]
        sel_rows.append(row)
        for k, d in enumerate(fdates):
            long_rows.append({"case": rank, "date": d, **{f"{s}_cum": ens[s][i, k] for s in STREAMS}})
    sel_df = pd.DataFrame(sel_rows)
    sel_df.to_csv(OUTPUT_DIR / "selected_cases.csv", index=False)
    pd.DataFrame(long_rows).to_csv(OUTPUT_DIR / "forecast_selected_cases.csv", index=False)
    print(sel_df.round(4).to_string(index=False))
    obs_cum_hist = obs_end
    plot_forecast(data, fdates, bands, sel_curves, hist_end, obs_cum_hist)

    # ---- fill the template -------------------------------------------------
    if tpath is None:
        tpath = make_own_template(fdates)
        wb, layout = read_template_layout(tpath)
    out_path = OUTPUT_DIR / f"{tpath.stem}_FILLED.xlsx"
    sheets = list(layout.items())
    if len(sheets) < len(chosen_idx):
        print(f"WARNING: template has {len(sheets)} case sheets but {len(chosen_idx)} cases were selected.")
    date_to_k = {d: k for k, d in enumerate(fdates)}
    for (name, lay), i in zip(sheets, chosen_idx):
        ws = wb[name]
        for r, d in lay["rows"]:
            if d in date_to_k:
                for s in STREAMS:
                    ws.cell(row=r, column=lay["cols"][s]).value = float(ens[s][i, date_to_k[d]])
    wb.save(out_path)
    print(f"\nsaved filled template: {out_path}   (original template untouched)")
    check_deliverable(out_path, obs_end)
    print("\nREMINDERS: (1) 2008-2028 is an extrapolation - the organisers' injection / well plan after "
          "2008 is NOT in the data.\n(2) Confirm with the organisers what 'Case 1..10' mean and how they "
          "are scored; change SELECTION_MODE if needed.")
    if SHOW_PLOTS:
        plt.show()


if __name__ == "__main__":
    main()